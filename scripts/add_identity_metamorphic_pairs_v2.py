# -*- coding: utf-8 -*-
"""Append table-structure metamorphic fixtures/pairs to Benchmark v2 (no holdout rewrite)."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

REPO = Path(__file__).resolve().parents[1]
V2 = REPO / "data" / "eval" / "document_set_benchmark_v2"
NOW = datetime.now(timezone.utc).isoformat()


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _write_json(path: Path, obj: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _read_jsonl(path: Path) -> list[dict]:
    if not path.is_file():
        return []
    return [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows) + "\n", encoding="utf-8")


def _append_jsonl(path: Path, rows: list[dict], key: str = "case_id") -> None:
    existing = _read_jsonl(path)
    have = {r.get(key) for r in existing}
    for r in rows:
        if r.get(key) not in have:
            existing.append(r)
            have.add(r.get(key))
    _write_jsonl(path, existing)


def _mdtm_cols(path: Path, rows: list[tuple[str, str, str]], *, col_order: list[str], title: str, note_col: bool = False) -> None:
    doc = Document()
    doc.add_heading(title, level=1)
    n_cols = len(col_order) + (1 if note_col else 0)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    for i, name in enumerate(col_order):
        table.rows[0].cells[i].text = name
    if note_col:
        table.rows[0].cells[len(col_order)].text = "비고"
    for i, (req, des, test) in enumerate(rows, start=1):
        mapping = {
            "Requirement": req,
            "Design": des,
            "Test": test,
            "Req": req,
            "Design ID": des,
            "Test ID": test,
        }
        for j, name in enumerate(col_order):
            table.rows[i].cells[j].text = mapping.get(name, "")
        if note_col:
            table.rows[i].cells[len(col_order)].text = "note"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def main() -> None:
    # cleanup mistaken files from prior addon attempt
    for bad in (
        V2 / "development" / "cases.jsonl",
        V2 / "development" / "document_labels.jsonl",
        V2 / "development" / "node_labels.jsonl",
        V2 / "development" / "node_eligibility.jsonl",
        V2 / "development" / "writer_expectations.jsonl",
    ):
        if bad.is_file():
            bad.unlink()

    fx = V2 / "fixtures" / "ec_sw"
    gr = V2 / "fixtures" / "general_report"
    rows = [
        ("Req. 10", "5.1.1", "TC-10"),
        ("Req. 11", "5.1.2", "TC-11"),
        ("Req. 12", "5.1.3", "TC-12"),
        ("Req. 101", "6.0.1", "TC-101"),
    ]
    base = fx / "mdtm_base.docx"
    if not base.is_file():
        raise SystemExit("mdtm_base.docx missing")

    _mdtm_cols(
        fx / "mdtm_cols_dtr.docx",
        rows,
        col_order=["Design", "Test", "Requirement"],
        title="Traceability Matrix ColReorder",
    )
    _mdtm_cols(
        fx / "mdtm_note_col.docx",
        rows,
        col_order=["Requirement", "Design", "Test"],
        title="Traceability Matrix Notes",
        note_col=True,
    )
    doc = Document()
    table = doc.add_table(rows=1 + len(rows), cols=3)
    for i, name in enumerate(["Requirement", "Design", "Test"]):
        table.rows[0].cells[i].text = name
    for i, (req, des, test) in enumerate(rows, start=1):
        table.rows[i].cells[0].text = req
        table.rows[i].cells[1].text = des
        table.rows[i].cells[2].text = test
    doc.save(str(fx / "mdtm_no_caption.docx"))
    _mdtm_cols(
        fx / "mdtm_row_shuffle.docx",
        list(reversed(rows)),
        col_order=["Requirement", "Design", "Test"],
        title="Traceability Matrix RowShuffle",
    )

    doc = Document()
    doc.add_heading("배경", level=1)
    doc.add_heading("추진 계획", level=1)
    t = doc.add_table(rows=3, cols=2)
    t.rows[0].cells[0].text = "단계"
    t.rows[0].cells[1].text = "기간"
    t.rows[1].cells[0].text = "1"
    t.rows[1].cells[1].text = "2026-Q4"
    t.rows[2].cells[0].text = "2"
    t.rows[2].cells[1].text = "2027-Q1"
    doc.save(str(gr / "report_schedule_synonym.docx"))

    from document_ai.domain_packs.ec_sw.mdtm_indexer import index_mdtm_document

    indexed = index_mdtm_document(source_path=base, document_id="MDTM")
    by_req = {}
    for n in indexed.get("nodes") or []:
        for r in n.source_identifiers.get("requirement_ids") or []:
            by_req[r] = n.node_id
    req11 = by_req.get("Req. 11")

    new_cases = [
        {
            "case_id": "v2_dev_ec_tbl_reorder",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 11 추적성 행 갱신",
            "split": "development",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/mdtm_cols_dtr.docx",
                    "filename": "mdtm_cols_dtr.docx",
                    "role": "traceability",
                }
            ],
            "enabled_documents": ["MDTM_COLS_DTR"],
            "expected_status": "SUCCESS",
            "tags": ["table_structure", "column_reorder"],
            "difficulty": "medium",
            "notes": "table column reorder metamorphic",
            "source_type": "GENERATED",
            "transformation": "table_column_reorder",
            "base_case_id": "v2_dev_ec_req11_exact",
            "fixture_fingerprint": _sha(fx / "mdtm_cols_dtr.docx"),
            "sanitization_status": "N/A",
        },
        {
            "case_id": "v2_dev_ec_tbl_empty",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 11 추적성 행 갱신",
            "split": "development",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/mdtm_with_gap.docx",
                    "filename": "mdtm_with_gap.docx",
                    "role": "traceability",
                }
            ],
            "enabled_documents": ["MDTM_WITH_GAP"],
            "expected_status": "SUCCESS",
            "tags": ["table_structure", "empty_row"],
            "difficulty": "medium",
            "notes": "empty row metamorphic",
            "source_type": "GENERATED",
            "transformation": "table_empty_row",
            "base_case_id": "v2_dev_ec_req11_exact",
            "fixture_fingerprint": _sha(fx / "mdtm_with_gap.docx"),
            "sanitization_status": "N/A",
        },
        {
            "case_id": "v2_dev_ec_tbl_note",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 11 추적성 행 갱신",
            "split": "development",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/mdtm_note_col.docx",
                    "filename": "mdtm_note_col.docx",
                    "role": "traceability",
                }
            ],
            "enabled_documents": ["MDTM_NOTE_COL"],
            "expected_status": "SUCCESS",
            "tags": ["table_structure", "note_column"],
            "difficulty": "medium",
            "notes": "note column metamorphic",
            "source_type": "GENERATED",
            "transformation": "table_note_column",
            "base_case_id": "v2_dev_ec_req11_exact",
            "fixture_fingerprint": _sha(fx / "mdtm_note_col.docx"),
            "sanitization_status": "N/A",
        },
        {
            "case_id": "v2_dev_ec_tbl_nocap",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 11 추적성 행 갱신",
            "split": "development",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/mdtm_no_caption.docx",
                    "filename": "mdtm_no_caption.docx",
                    "role": "traceability",
                }
            ],
            "enabled_documents": ["MDTM_NO_CAPTION"],
            "expected_status": "SUCCESS",
            "tags": ["table_structure", "caption_removal"],
            "difficulty": "medium",
            "notes": "caption removal metamorphic",
            "source_type": "GENERATED",
            "transformation": "table_caption_removal",
            "base_case_id": "v2_dev_ec_req11_exact",
            "fixture_fingerprint": _sha(fx / "mdtm_no_caption.docx"),
            "sanitization_status": "N/A",
        },
        {
            "case_id": "v2_dev_ec_tbl_rowpos",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 11 추적성 행 갱신",
            "split": "development",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/mdtm_row_shuffle.docx",
                    "filename": "mdtm_row_shuffle.docx",
                    "role": "traceability",
                }
            ],
            "enabled_documents": ["MDTM_ROW_SHUFFLE"],
            "expected_status": "SUCCESS",
            "tags": ["table_structure", "row_position"],
            "difficulty": "medium",
            "notes": "row position metamorphic",
            "source_type": "GENERATED",
            "transformation": "table_row_position",
            "base_case_id": "v2_dev_ec_req11_exact",
            "fixture_fingerprint": _sha(fx / "mdtm_row_shuffle.docx"),
            "sanitization_status": "N/A",
        },
        {
            "case_id": "v2_dev_gr_sched_syn",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "일정 표 2026-Q4 수정",
            "split": "development",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_schedule_synonym.docx",
                    "filename": "report_schedule_synonym.docx",
                    "role": "general_report",
                }
            ],
            "enabled_documents": ["REPORT_SCHEDULE_SYNONYM"],
            "expected_status": "SUCCESS",
            "tags": ["schedule", "table", "table_structure", "header_synonym"],
            "difficulty": "medium",
            "notes": "schedule header synonym",
            "source_type": "GENERATED",
            "transformation": "table_header_synonym",
            "base_case_id": "v2_dev_gr_schedule_table",
            "fixture_fingerprint": _sha(gr / "report_schedule_synonym.docx"),
            "sanitization_status": "N/A",
        },
    ]

    identity_meta = {
        "expected_short_id": "MDTM",
        "expected_document_role": "traceability",
        "expected_domain_pack_id": "ec_sw_v1",
    }
    new_pairs = [
        {
            "pair_id": "mp_tbl_reorder",
            "base_case_id": "v2_dev_ec_req11_exact",
            "variant_case_id": "v2_dev_ec_tbl_reorder",
            "transformation": "table_column_reorder",
            "expected_relation": "same_document_decision",
            **identity_meta,
        },
        {
            "pair_id": "mp_tbl_empty",
            "base_case_id": "v2_dev_ec_req11_exact",
            "variant_case_id": "v2_dev_ec_tbl_empty",
            "transformation": "table_empty_row",
            "expected_relation": "same_document_decision",
            **identity_meta,
        },
        {
            "pair_id": "mp_tbl_note",
            "base_case_id": "v2_dev_ec_req11_exact",
            "variant_case_id": "v2_dev_ec_tbl_note",
            "transformation": "table_note_column",
            "expected_relation": "same_document_decision",
            **identity_meta,
        },
        {
            "pair_id": "mp_tbl_nocap",
            "base_case_id": "v2_dev_ec_req11_exact",
            "variant_case_id": "v2_dev_ec_tbl_nocap",
            "transformation": "table_caption_removal",
            "expected_relation": "same_document_decision",
            **identity_meta,
        },
        {
            "pair_id": "mp_tbl_rowpos",
            "base_case_id": "v2_dev_ec_req11_exact",
            "variant_case_id": "v2_dev_ec_tbl_rowpos",
            "transformation": "table_row_position",
            "expected_relation": "same_document_decision",
            **identity_meta,
        },
        {
            "pair_id": "mp_tbl_hdr_syn",
            "base_case_id": "v2_dev_gr_schedule_table",
            "variant_case_id": "v2_dev_gr_sched_syn",
            "transformation": "table_header_synonym",
            "expected_relation": "same_document_decision",
            "expected_document_role": "general_report",
            "expected_domain_pack_id": "generic_document_v1",
        },
        {
            "pair_id": "mp_tbl_sched_list",
            "base_case_id": "v2_dev_gr_schedule_table",
            "variant_case_id": "v2_dev_gr_schedule_list",
            "transformation": "table_schedule_list",
            "expected_relation": "same_document_decision",
            "expected_document_role": "general_report",
        },
        {
            "pair_id": "mp_tbl_sched_para",
            "base_case_id": "v2_dev_gr_schedule_table",
            "variant_case_id": "v2_dev_gr_schedule_para",
            "transformation": "table_schedule_paragraph",
            "expected_relation": "same_document_decision",
            "expected_document_role": "general_report",
        },
    ]

    # Write case JSON files + update manifest list
    man_path = V2 / "manifest.json"
    man = json.loads(man_path.read_text(encoding="utf-8"))
    dev_list = list(man.get("development_cases") or [])
    for c in new_cases:
        domain = c["domain"]
        rel = f"development/cases/{domain}/{c['case_id']}.json"
        _write_json(V2 / rel, c)
        if rel not in dev_list:
            dev_list.append(rel)
    man["development_cases"] = dev_list
    man["meta"] = man.get("meta") or {}
    man["meta"]["development_count"] = len(dev_list)
    man["counts"] = {
        "development_cases": len(dev_list),
        "metamorphic_pairs": None,  # filled below
    }
    man["identity_metamorphic_addon_at"] = NOW

    # Labels append (existing gold untouched; only add new case_ids)
    doc_rows = []
    node_rows = []
    elig_rows = []
    writer_rows = []
    for c in new_cases:
        cid = c["case_id"]
        doc_id = c["enabled_documents"][0]
        if c["domain"] == "ec_sw":
            doc_rows.append(
                {
                    "case_id": cid,
                    "document_id": doc_id,
                    "gold_status": "IMPACTED",
                    "review_notes": "structure_based_pre_prediction_label",
                }
            )
            if req11:
                node_rows.append(
                    {
                        "case_id": cid,
                        "document_id": doc_id,
                        "node_id": req11,
                        "gold_status": "PATCH_CANDIDATE",
                        "node_evaluation_mode": "REQUIRED",
                        "primary_node_id": req11,
                        "acceptable_node_ids": [],
                        "acceptable_node_groups": [],
                        "label_rationale": "pre_prediction_fixture_structure",
                        "labeled_by": "identity_metamorphic_addon",
                        "labeled_at": NOW,
                        "confidence": 0.85,
                    }
                )
                elig_rows.append(
                    {
                        "case_id": cid,
                        "domain": "ec_sw",
                        "document_id": doc_id,
                        "node_evaluation_mode": "REQUIRED",
                        "primary_node_id": req11,
                        "acceptable_node_ids": [],
                        "acceptable_node_groups": [],
                        "label_rationale": "pre_prediction_fixture_structure",
                        "labeled_by": "identity_metamorphic_addon",
                        "labeled_at": NOW,
                        "confidence": 0.85,
                    }
                )
        else:
            nid = "general_report_v1.schedule"
            doc_rows.append(
                {
                    "case_id": cid,
                    "document_id": doc_id,
                    "gold_status": "REVIEW_REQUIRED",
                    "review_notes": "structure_based_pre_prediction_label",
                }
            )
            node_rows.append(
                {
                    "case_id": cid,
                    "document_id": doc_id,
                    "node_id": nid,
                    "gold_status": "REVIEW_REQUIRED",
                    "node_evaluation_mode": "AMBIGUOUS",
                    "primary_node_id": nid,
                    "acceptable_node_ids": [],
                    "acceptable_node_groups": [],
                    "label_rationale": "pre_prediction_fixture_structure",
                    "labeled_by": "identity_metamorphic_addon",
                    "labeled_at": NOW,
                    "confidence": 0.85,
                }
            )
            elig_rows.append(
                {
                    "case_id": cid,
                    "domain": "general_report",
                    "document_id": doc_id,
                    "node_evaluation_mode": "AMBIGUOUS",
                    "primary_node_id": nid,
                    "acceptable_node_ids": [nid],
                    "acceptable_node_groups": [[nid]],
                    "label_rationale": "pre_prediction_fixture_structure",
                    "labeled_by": "identity_metamorphic_addon",
                    "labeled_at": NOW,
                    "confidence": 0.85,
                }
            )
        writer_rows.append(
            {
                "case_id": cid,
                "should_write": False,
                "expected_result_status": "GATED",
                "original_must_remain_unchanged": True,
                "expected_changed_document_ids": [],
                "format_preservation_required": True,
            }
        )

    labels = V2 / "development" / "labels"
    _append_jsonl(labels / "document_impacts.jsonl", doc_rows)
    _append_jsonl(labels / "node_impacts.jsonl", node_rows)
    _append_jsonl(labels / "node_evaluation_eligibility.jsonl", elig_rows)
    _append_jsonl(labels / "writer_expectations.jsonl", writer_rows)

    pairs = _read_jsonl(V2 / "metamorphic_pairs.jsonl")
    have_p = {p["pair_id"] for p in pairs}
    for p in new_pairs:
        if p["pair_id"] not in have_p:
            pairs.append(p)
    _write_jsonl(V2 / "metamorphic_pairs.jsonl", pairs)
    man["counts"]["metamorphic_pairs"] = len(pairs)
    man_path.write_text(json.dumps(man, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"development_cases={len(dev_list)} metamorphic_pairs={len(pairs)}")


if __name__ == "__main__":
    main()
