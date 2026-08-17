"""Copy blank EC-SW templates and diff vs filled examples."""
import json
import shutil
import unicodedata
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOWNLOADS = Path(r"c:\Users\hsh71\Downloads")
PROJECT = Path(__file__).resolve().parents[1]
TEMPLATES = PROJECT / "data" / "templates" / "ec_sw"
EXAMPLES = PROJECT / "data" / "examples" / "ec_sw"
OUT = PROJECT / "data" / "schemas" / "ec_sw"

TEMPLATES.mkdir(parents=True, exist_ok=True)
EXAMPLES.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)

FILLED = {
    "spec_design": EXAMPLES / "spec_mddr_EC-SW-MDDR(XA) 소프트웨어 설계 명세서.docx",
    "spec_requirements": EXAMPLES / "spec_mdsr_EC-SW-MDSR(XA)_소프트웨어_요구사항명세서.docx",
}


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def classify_blank(filename: str) -> str | None:
    name = nfc(filename)
    if "XX-XX-XX(0)" not in name:
        return None
    if "양식" not in name:
        return None
    if "설계" in name and "명세" in name:
        return "spec_design"
    if "요구" in name and "명세" in name:
        return "spec_requirements"
    return None


def iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def cell_texts(table: Table) -> list[list[str]]:
    return [[c.text.strip() for c in row.cells] for row in table.rows]


def extract_structure(path: Path) -> dict:
    doc = Document(str(path))
    paragraphs = []
    tables = []
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            t = block.text.strip()
            if t:
                paragraphs.append(t[:200])
        elif isinstance(block, Table):
            rows = cell_texts(block)
            non_empty = sum(1 for r in rows for c in r if c)
            tables.append({"rows": len(rows), "cols": len(rows[0]) if rows else 0, "non_empty_cells": non_empty, "sample": rows[:3]})
    return {
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "paragraphs_head": paragraphs[:25],
        "tables_summary": tables,
        "size_kb": round(path.stat().st_size / 1024, 1),
    }


def find_blank_templates() -> dict[str, Path]:
    found: dict[str, Path] = {}
    for p in DOWNLOADS.glob("*.docx"):
        template_id = classify_blank(p.name)
        if template_id and template_id not in found:
            found[template_id] = p
    return found


def table_diff(blank: Path, filled: Path) -> dict:
    b_doc = Document(str(blank))
    f_doc = Document(str(filled))
    b_tables = [b for b in iter_blocks(b_doc) if isinstance(b, Table)]
    f_tables = [b for b in iter_blocks(f_doc) if isinstance(b, Table)]

    diffs = []
    for i, (bt, ft) in enumerate(zip(b_tables, f_tables)):
        br = cell_texts(bt)
        fr = cell_texts(ft)
        if br == fr:
            continue
        changed_cells = []
        for ri, (brow, frow) in enumerate(zip(br, fr)):
            for ci, (bc, fc) in enumerate(zip(brow, frow)):
                if bc != fc and (bc or fc):
                    changed_cells.append({"row": ri, "col": ci, "blank": bc[:80], "filled": fc[:80]})
        if changed_cells:
            diffs.append({"table_index": i, "changed_count": len(changed_cells), "samples": changed_cells[:15]})

    return {
        "blank_tables": len(b_tables),
        "filled_tables": len(f_tables),
        "table_count_match": len(b_tables) == len(f_tables),
        "diff_tables": len(diffs),
        "diffs": diffs[:20],
    }


def main():
    blanks = find_blank_templates()
    if not blanks:
        raise SystemExit("No blank templates found in Downloads (XX-XX-XX(0)*[양식].docx)")

    report = {
        "blanks": {},
        "diffs": {},
        "xxcs": {
            "blank_template": False,
            "strategy": "Use filled XXCS as structural reference; clone table blocks per IA/UC/SI requirement; or create simplified internal template from filled doc skeleton",
        },
    }

    dest_names = {
        "spec_design": "template_mddr.docx",
        "spec_requirements": "template_mdsr.docx",
    }

    for template_id, src in blanks.items():
        dest = TEMPLATES / dest_names[template_id]
        shutil.copy2(src, dest)
        report["blanks"][template_id] = {
            "source": nfc(src.name),
            "dest": str(dest.relative_to(PROJECT)).replace("\\", "/"),
            **extract_structure(dest),
        }

        filled_path = FILLED.get(template_id)
        if filled_path and filled_path.exists():
            report["diffs"][template_id] = table_diff(dest, filled_path)
        else:
            report["diffs"][template_id] = {"error": f"filled not found: {filled_path}"}

    out_path = OUT / "template_diff_analysis.json"
    out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {out_path}")
    for tid in report["blanks"]:
        b = report["blanks"][tid]
        d = report["diffs"].get(tid, {})
        print(f"{tid}: tables={b['table_count']} diff_tables={d.get('diff_tables', '?')}")


if __name__ == "__main__":
    main()
