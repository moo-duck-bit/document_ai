# -*- coding: utf-8 -*-
"""Document AI presentation — AI Engineering framing.

Slide order (speaking): B1~B5 → metrics → UI/artifacts.
Regenerate screenshots: scripts/capture_pilot_demo_screenshots.py
Writer-ON Req.11 artifacts: docs/presentation/artifacts/req11_writer_on/
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
OUT_DIR = REPO / "docs" / "presentation"
SHOT_DIR = OUT_DIR / "screenshots"
PPTX_PATH = OUT_DIR / "Document_AI_v2_Presentation.pptx"
PPTX_PATH_ALT = OUT_DIR / "Document_AI_v2_Presentation_updated.pptx"
MD_PATH = OUT_DIR / "Document_AI_v2_Presentation.md"
SCRIPT_PATH = OUT_DIR / "SPEAKER_SCRIPT.md"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
NAVY_DARK = RGBColor(0x0F, 0x2A, 0x43)
GRAY = RGBColor(0x5C, 0x65, 0x70)
BLACK = RGBColor(0x1C, 0x1F, 0x24)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xEE, 0xF2, 0xF6)
ACCENT = RGBColor(0x2D, 0x6A, 0x9F)
CODE_BG = RGBColor(0x0F, 0x17, 0x20)
CODE_FG = RGBColor(0xE8, 0xEE, 0xF5)


def _run(p, text, *, size=20, bold=False, color=BLACK, font="Malgun Gothic"):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font


def _bar(slide, prs):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.1))
    s.fill.solid()
    s.fill.fore_color.rgb = NAVY
    s.line.fill.background()


def _blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bar(slide, prs)
    return slide


def _page_number(slide, prs, n: int, total: int):
    tb = slide.shapes.add_textbox(
        Inches(11.6), Inches(7.05), Inches(1.5), Inches(0.35)
    )
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, f"{n} / {total}", size=10, color=GRAY)


def add_page_numbers(prs):
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, start=1):
        _page_number(slide, prs, i, total)


def title_slide(prs, title, subtitle, notes=""):
    slide = _blank(prs)
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(12), Inches(1.6))
    _run(tb.text_frame.paragraphs[0], title, size=32, bold=True, color=NAVY_DARK)
    sb = slide.shapes.add_textbox(Inches(0.7), Inches(3.9), Inches(12), Inches(1.6))
    sb.text_frame.word_wrap = True
    _run(sb.text_frame.paragraphs[0], subtitle, size=16, color=GRAY)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def bullets(prs, title, lines, notes=""):
    slide = _blank(prs)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.32), Inches(12.2), Inches(0.7))
    _run(tb.text_frame.paragraphs[0], title, size=22, bold=True, color=NAVY_DARK)
    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(12), Inches(5.6))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(7)
        _run(p, f"•  {line}", size=17, color=BLACK)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def two_col(prs, title, left_t, left, right_t, right, notes=""):
    slide = _blank(prs)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12.2), Inches(0.6))
    _run(tb.text_frame.paragraphs[0], title, size=20, bold=True, color=NAVY_DARK)
    for x, ht, lines, border in (
        (0.55, left_t, left, NAVY),
        (6.95, right_t, right, ACCENT),
    ):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.05), Inches(5.8), Inches(5.5)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.color.rgb = border
        h = slide.shapes.add_textbox(Inches(x + 0.2), Inches(1.2), Inches(5.4), Inches(0.4))
        _run(h.text_frame.paragraphs[0], ht, size=15, bold=True, color=border)
        b = slide.shapes.add_textbox(Inches(x + 0.2), Inches(1.75), Inches(5.4), Inches(4.5))
        tf = b.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            _run(p, f"•  {line}", size=13, color=BLACK)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def flow(prs, title, steps, caption="", notes=""):
    slide = _blank(prs)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.2), Inches(0.6))
    _run(tb.text_frame.paragraphs[0], title, size=20, bold=True, color=NAVY_DARK)
    n = len(steps)
    gap = 0.08
    w = (12.1 - gap * (n - 1)) / n
    for i, label in enumerate(steps):
        x = 0.6 + i * (w + gap)
        sh = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.4), Inches(w), Inches(1.5)
        )
        sh.fill.solid()
        sh.fill.fore_color.rgb = NAVY if i % 2 == 0 else ACCENT
        sh.line.fill.background()
        tf = sh.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        _run(tf.paragraphs[0], label, size=10, bold=True, color=WHITE)
    if caption:
        c = slide.shapes.add_textbox(Inches(0.7), Inches(4.3), Inches(12), Inches(1.8))
        c.text_frame.word_wrap = True
        _run(c.text_frame.paragraphs[0], caption, size=14, color=GRAY)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def code_slide(prs, title, path_label, code, meaning, notes=""):
    slide = _blank(prs)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.45))
    _run(tb.text_frame.paragraphs[0], title, size=18, bold=True, color=NAVY_DARK)
    pb = slide.shapes.add_textbox(Inches(0.5), Inches(0.7), Inches(12.3), Inches(0.3))
    _run(pb.text_frame.paragraphs[0], path_label, size=10, color=GRAY)
    bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(1.05), Inches(12.4), Inches(4.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = CODE_BG
    bg.line.fill.background()
    cb = slide.shapes.add_textbox(Inches(0.65), Inches(1.2), Inches(12.1), Inches(4.2))
    tf = cb.text_frame
    for i, line in enumerate(code.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(1)
        r = p.add_run()
        r.text = line if line else " "
        r.font.size = Pt(10)
        r.font.name = "Consolas"
        r.font.color.rgb = CODE_FG
    mb = slide.shapes.add_textbox(Inches(0.5), Inches(5.75), Inches(12.3), Inches(1.2))
    mb.text_frame.word_wrap = True
    _run(mb.text_frame.paragraphs[0], meaning, size=13, color=BLACK)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def table_slide(prs, title, headers, rows, notes=""):
    slide = _blank(prs)
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.55))
    _run(tb.text_frame.paragraphs[0], title, size=20, bold=True, color=NAVY_DARK)
    shape = slide.shapes.add_table(
        len(rows) + 1,
        len(headers),
        Inches(0.5),
        Inches(1.0),
        Inches(12.3),
        Inches(min(0.48 * (len(rows) + 1), 5.6)),
    )
    table = shape.table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(11)
                r.font.color.rgb = WHITE
                r.font.name = "Malgun Gothic"
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(11)
                    r.font.name = "Malgun Gothic"
            if i % 2:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def screenshot_slide(prs, title, image_name, captions, notes=""):
    slide = _blank(prs)
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.42))
    _run(tb.text_frame.paragraphs[0], title, size=18, bold=True, color=NAVY_DARK)
    cap = slide.shapes.add_textbox(Inches(0.45), Inches(0.62), Inches(12.4), Inches(0.85))
    tf = cap.text_frame
    tf.word_wrap = True
    for i, line in enumerate(captions):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        _run(p, f"•  {line}", size=12, color=BLACK)
    img = SHOT_DIR / image_name
    if not img.exists():
        miss = slide.shapes.add_textbox(Inches(0.6), Inches(2.5), Inches(12), Inches(1))
        _run(
            miss.text_frame.paragraphs[0],
            f"(이미지 없음: {image_name})",
            size=14,
            color=GRAY,
        )
    else:
        max_w, max_h = Inches(12.2), Inches(5.6)
        pic = slide.shapes.add_picture(str(img), Inches(0.55), Inches(1.55))
        scale = min(max_w / pic.width, max_h / pic.height)
        pic.width = int(pic.width * scale)
        pic.height = int(pic.height * scale)
        pic.left = int((prs.slide_width - pic.width) / 2)
        pic.top = Inches(1.55)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ========== 0. Framing ==========
    title_slide(
        prs,
        "Document AI × AI Engineering",
        "마인드리움·탐색임상 문서 자동화 사례로 본\n"
        "에이전트 하네스 · B1~B5 · 평가 · 안전 · 제품 루프",
        "순서: 한 문장 정의 → 비교 → 원리(B1~B5·코드) → 숫자 → UI/산출물.",
    )
    bullets(
        prs,
        "오늘 발표 순서",
        [
            "0) Document AI가 무엇인지 — 최종 목표 vs 오늘 초점 + 비교",
            "1) 배경·AI Engineering",
            "2) 원리 — B1→B5 (화면 캡처) + 코드 흐름",
            "3) 벤치마크·Pilot 수치",
            "4) 산출물·Writer · 보고서/제안서 확장",
            "5) 향후 · Q&A",
        ],
        "라이브 데모 없음. B1~B5는 Pilot 화면 캡처로 설명.",
    )
    # --- Definition: North Star + today's focus ---
    title_slide(
        prs,
        "Document AI란 (한 문장)",
        "최종 목표: 완성본으로 학습해, 빈 양식과 새 케이스로 문서를 자동 작성한다 (Form Fill).\n\n"
        "오늘 초점: 그 전제 층 — 변경이 어디에 닿는지 찾아 보여 주고,\n"
        "사람이 승인한 뒤에만 복사본에 안전하게 반영한다 (Change Impact).",
        "최종 목적(자동 작성)과 오늘 데모(변경 영향·안전)를 구분해 기억.",
    )
    flow(
        prs,
        "Document AI 동작 흐름 (오늘: Change Impact)",
        ["변경요청", "찾기", "보여주기", "사람 승인", "복사본만"],
        "원본은 건드리지 않는다. 오늘 메인 경로에 생성형 LLM 문장 작성을 두지 않는다.",
    )
    table_slide(
        prs,
        "무엇이 다른가 — Document AI의 자리",
        ["구분", "LLM에 문서 맡기기", "일반 RAG QA", "Document AI (오늘)"],
        [
            ["목표", "문장·초안 생성", "질문에 답", "변경 위치 탐색·안전 반영"],
            ["쓰는가", "자유롭게 생성", "거의 안 씀", "승인 후에만 복사본"],
            ["근거", "프롬프트", "검색 청크", "Req/표·규칙·REVIEW 사유"],
            ["실패 비용", "그럴듯한 오답", "틀린 답변", "오패치·원본 손상 차단"],
            ["사람 역할", "후편집", "해석", "후보 승인·거절·보류"],
        ],
        "오늘 층 = 변경 영향+안전 게이트. 최종 Form Fill(자동 초안)과는 층을 구분.",
    )
    two_col(
        prs,
        "AI Engineering이란 (이 프로젝트)",
        "하지 않는 것",
        [
            "프롬프트만으로 무근거 문서 생성",
            "벤치 점수만 올리는 튜닝",
            "원본 파일에 바로 쓰기",
            "설명 없는 블랙박스 추천",
        ],
        "하는 것",
        [
            "하네스·스킬·티켓으로 작업 통제",
            "도메인 규칙 + 구조적 추론 (B1~B5)",
            "Holdout·회귀·안전 지표",
            "사람-in-the-loop 게이트",
            "Pilot으로 제품 검증",
        ],
    )
    bullets(
        prs,
        "도메인 배경: 마인드리움 · 탐색임상",
        [
            "마인드리움 앱·탐색임상 과정에서 EC-SW 문서가 지속 갱신",
            "한 변경이 MDSR→MDDR→추적성/시험으로 전파",
            "병목은 ‘문장 창작’이 아니라 위치 찾기·일치 확인·누락 방지",
            "그래서: 찾기 → 보여주기 → 승인 → 복사본만 반영",
        ],
    )
    two_col(
        prs,
        "North Star vs 현재",
        "North Star (Form Fill)",
        [
            "완성본 학습 → 빈 양식+케이스로 초안 자동 작성",
            "구조화 필드 LLM 금지 · free_text만 제한적",
        ],
        "현재 (Change Impact · 오늘)",
        [
            "B1~B5로 영향 위치·게이트·패치",
            "Pilot UI · Domain Pack 확장",
        ],
    )

    # ========== 1. B1~B5 (screen captures) ==========
    bullets(
        prs,
        "Document Change Agent = B1→B5",
        [
            "오늘 구현의 메인 로직: 인덱싱 → 검색 → 판정 → 게이트 → 패치",
            "문장을 ‘알아서 써 넣는’ 파이프라인이 아님 — 단계 계약",
            "이제부터 각 단계를 Pilot 화면 캡처와 함께 본다",
        ],
    )
    flow(
        prs,
        "B1 ~ B5",
        ["B1 Index", "B2 Retrieve", "B3 Judge", "B4 Gate", "B5 Patch"],
        "앞단 Recall(놓치지 않기) · 뒷단 Precision/Safety(잘못 쓰지 않기)",
    )
    table_slide(
        prs,
        "B1~B5 — 단계 · 질문 · 산출",
        ["단계", "질문", "산출"],
        [
            ["B1 Index", "어떤 단위로 자를까?", "Req 블록 인덱스 (변경 없음)"],
            ["B2 Retrieve", "후보가 어디인가?", "Top-K (lexical+TF-IDF)"],
            ["B3 Judge", "정말 영향인가?", "IMPACTED / UNCERTAIN / NOT"],
            ["B4 Gate", "써도 안전한가?", "CONSISTENT만 통과 · 사람 승인"],
            ["B5 Patch", "어디에 남길까?", "복사본 · Diff · 원본 보존"],
        ],
    )
    screenshot_slide(
        prs,
        "B1 Index — 화면: 업로드 · 문서 유형",
        "01_upload_zoom.png",
        [
            "질문: 문서를 어떤 단위로 자를까? → Req 블록 인덱스 (이 단계에서는 수정 없음)",
            "화면: 시나리오·문서를 올리고 세션 복사본만 사용 — 인덱싱의 입력 준비",
            "요점: 이후 검색·판정의 ‘근거 단위’를 여기서 고정한다",
        ],
    )
    screenshot_slide(
        prs,
        "B2 Retrieve — 화면: 변경 요청",
        "04_change_request_zoom.png",
        [
            "질문: 변경 요청과 관련 후보는 어디인가? → Top-K (lexical+TF-IDF, LLM 임베딩 아님)",
            "화면: ‘Req.11 갱신 + 추적성 검토’처럼 사람이 바꾸고 싶은 내용을 적는다",
            "요점: 무엇을 바꿀지는 사람, 어디가 관련인지는 시스템이 후보로 줄인다",
        ],
    )
    screenshot_slide(
        prs,
        "B3 Judge — 화면: 영향 분석",
        "05_analyze_zoom.png",
        [
            "질문: 그 후보가 정말 영향인가? → IMPACTED / UNCERTAIN / NOT",
            "화면: 읽기 전용 분석 · 패치 후보·검토 항목 스키마 (아직 파일 미수정)",
            "요점: ID만 같다고 IMPACTED로 올리지 않음 · 애매하면 UNCERTAIN→사람 검토",
        ],
    )
    screenshot_slide(
        prs,
        "B4 Gate — 화면: 검토·승인",
        "06_review_zoom.png",
        [
            "질문: 지금 써도 안전한가? → CONSISTENT만 통과, 명시 승인만 Writer 대상",
            "화면: 사유·근거와 함께 승인 / 거절 / 보류",
            "요점: ‘관련 있다’와 ‘지금 써도 된다’를 여기서 분리한다",
        ],
    )
    screenshot_slide(
        prs,
        "B5 Patch — 화면: Writer · 복사본",
        "12_req11_writer_on_copy.png",
        [
            "질문: 어디에 남길까? → 복사본만 · fingerprint로 원본 보존 · Diff 기록",
            "화면: WRITTEN_COPY_ONLY (MVP는 안전 복사; 셀 문장 rewrite는 로드맵)",
            "요점: 오늘 메인 경로에 생성형 LLM 문장 작성을 두지 않음 · 조건 미충족 시 BLOCKED",
        ],
    )

    # ========== 1b. Code = process ==========
    bullets(
        prs,
        "코드로 보는 프로세스",
        [
            "아래 코드는 ‘똑똑한 문장 생성’이 아니라 단계 계약이다",
            "함수 이름 ≈ B1~B5 · UI ‘분석/승인/Writer’와 1:1로 대응",
            "읽기 포인트: 입력이 무엇인지, 무엇을 넘기는지, 어디서 막는가",
        ],
    )
    code_slide(
        prs,
        "코드 ① 오케스트레이션 — 전체 흐름",
        "pilot_v2 / change agent",
        """# UI ‘분석 실행’ ≈ 이 파이프라인
session = create_session(docs, change_request)
session = resolve_identity(...)      # Pack 추천·확인
session = analyze(session)           # B1→B3 (+후보 스키마)
# UI ‘승인’ 
session = apply_decisions(APPROVE|REJECT|HOLD)
# UI ‘Writer’
session = run_writer_if_allowed(     # B4·B5 게이트
    enable_write=..., env_flag=...
)""",
        "프로세스: 세션 만들기 → 유형 확인 → 분석 → 사람 결정 → 조건부 쓰기.",
    )
    code_slide(
        prs,
        "코드 ② B1·B2 — 인덱스 · 검색",
        "index / retrieve",
        """# B1: 문서를 Req 블록으로 자름 (미변경)
index = build_req_index(mdsr, mddr)  # id, title, body, locator

# B2: CR과 비슷한 후보만 Top-K
scores = 0.4*lexical(CR, index) + 0.6*tfidf_cosine(CR, index)
candidates = top_k(scores, k=15)     # gold 누수 금지""",
        "왜: 전체 문서를 바로 판정하지 않고, 설명 가능한 점수로 후보를 좁힌다.",
    )
    code_slide(
        prs,
        "코드 ③ B3·B4 — 판정 · 게이트",
        "judge / gate",
        """# B3: 후보마다 영향 라벨
for c in candidates:
    label = judge(c, rules)  # IMPACTED | UNCERTAIN | NOT
    # rule: ID만으로 IMPACTED 금지 · false-friend 차단

# B4: 쓰기 자격
if label != IMPACTED: skip
if consistency != CONSISTENT:  # CONFLICT/REVIEW
    block_write(); ask_human()
else:
    allow_patch_path()""",
        "왜: ‘관련 있음’과 ‘지금 써도 됨’을 코드에서 분리한다.",
    )
    code_slide(
        prs,
        "코드 ④ B5 — Writer 게이트 · 원본 보존",
        "controlled_writer",
        """gates = {
  "routing_confirmed": True,
  "explicit_approved_items": True,  # 사람 APPROVE
  "enable_write_flag": True,        # UI 체크
  "controlled_writer_env": True,    # 서버 플래그
}
if all(gates.values()):
    write_copy_only(dst)            # 원본 path 금지
    assert fingerprint(src) == before_sha
else:
    return BLOCKED""",
        "capability ≠ activation. 하나라도 false면 쓰지 않는다.",
    )

    # ========== 2. Metrics ==========
    bullets(
        prs,
        "성능 지표 — 무엇을 측정하는가",
        [
            "생성 문장의 자연스러움이 아니라, 위치 탐색·판정·안전성·사용성을 측정",
            "Document F1 / E2E: 영향 문서·파이프라인 성공",
            "Node Top-1 / Recall: 올바른 위치 탐색",
            "False Patch / Unsafe / Original Preservation: 안전",
            "Pilot: Completion · 사람 동의 · Trust (벤치와 분리)",
        ],
        "출처: Benchmark v2 run 20260801T192542Z_5a685d77 · Pilot Run 01",
    )
    table_slide(
        prs,
        "Benchmark · Pilot 수치 요약",
        ["층", "지표", "결과", "의미"],
        [
            ["Holdout", "Document Macro F1", "0.956", "봉인 셋 일반화"],
            ["Holdout", "E2E Success", "0.957", "파이프라인 성공"],
            ["Holdout", "Required Node Top-1", "0.800", "필수 위치 1순위"],
            ["Holdout", "False Patch / Unsafe", "0 / 0", "오패치·위험 실패 0"],
            ["Safety", "Original Preservation", "1.0", "원본 무변경"],
            ["Pilot 01", "Completion / Safety", "1.0 / PASS", "9세션·무단쓰기 0"],
            ["Pilot 01", "Doc·Node 동의", "0.889", "사람 평가 일치"],
            ["Pilot 01", "Trust / Usability", "3.78 / 4.0", "1–5 척도"],
        ],
        "Holdout n=23 · Pilot Writer OFF dry-run · 벤치≠제품 점수 혼동 금지",
    )
    two_col(
        prs,
        "수치 해석 · 한계 (정직하게)",
        "이렇게 읽으면 됩니다",
        [
            "안전 지표가 핵심 주장",
            "Holdout F1/E2E는 일반화 증거",
            "Pilot은 ‘써도 되는지’ 별도 층",
        ],
        "한계 · 아직 주장하지 않는 것",
        [
            "의미만으로 된 요청: 후보 0건 (bug로 기록)",
            "Pilot 01: 정해진 시나리오·Writer OFF",
            "Pilot Writer MVP: 텍스트 재작성 전 복사본 단계",
            "벤치 점수로 Pilot을 튜닝하지 않음",
        ],
    )

    # ========== 3. Artifacts (UI already shown in B1~B5) ==========
    bullets(
        prs,
        "산출물 — B1~B5가 남긴 증거",
        [
            "앞에서 B1~B5를 화면으로 보았음 · 여기서는 파일·Diff 증거",
            "시나리오: EC-SW 단일 Req.11 변경",
            "이어서 보고서·제안서 Pack 확장",
        ],
    )
    screenshot_slide(
        prs,
        "산출물 A — Before: Req.11 위치",
        "11_req11_before_locate.png",
        [
            "원본 추적성 표에서 Req.11 행을 시스템이 지목 (생성보다 위치 탐색)",
            "artifact: docs/presentation/artifacts/req11_writer_on/before_*.docx",
        ],
    )
    screenshot_slide(
        prs,
        "산출물 B — Writer ON: 복사본 · 원본 보존",
        "12_req11_writer_on_copy.png",
        [
            "status WRITTEN_COPY_ONLY · original_preservation ok",
            "정직 고지: MVP는 안전 복사까지 (셀 텍스트 rewrite는 로드맵)",
            "copy: artifacts/req11_writer_on/after_MDTM_PILOT_BASE.docx",
        ],
    )
    screenshot_slide(
        prs,
        "산출물 C — B5 텍스트 Diff 예시 (scenario-001)",
        "13_b5_text_before_after.png",
        [
            "Document Change Agent B5에서 실제 before/after 문안 Diff",
            "Pilot Writer 텍스트 apply와는 단계가 다름 — 혼동하지 말 것",
        ],
    )
    screenshot_slide(
        prs,
        "Writer OFF 경로 (Pilot dry-run)",
        "07_writer_blocked_zoom.png",
        [
            "기능이 구현되어 있어도, 실행 조건이 충족되지 않으면 동작하지 않음",
            "보조: capability ≠ activation · status BLOCKED",
        ],
    )

    # Domain expansion (compact)
    bullets(
        prs,
        "도메인 확장: 보고서 · 제안서",
        [
            "같은 B1~B5·게이트 · Domain Pack만 추가",
            "general_report: 방법론 등 섹션 단위",
            "business_proposal: 예산·일정·위험 (금액은 REVIEW)",
        ],
    )
    screenshot_slide(
        prs,
        "확장 — 일반보고서 분석",
        "09_gr_analyze.png",
        ["방법론 CR · REVIEW 항목으로 동일 UI 계약"],
    )
    screenshot_slide(
        prs,
        "확장 — 사업제안서 검토",
        "10_bp_review.png",
        ["예산 후보 REVIEW_REQUIRED · 사람 승인 전제"],
    )

    bullets(
        prs,
        "향후 방향",
        [
            "단기: Pilot Run 02(실사용자) · Writer ON 소규모 안전 실험",
            "중기: 의미-only Locator · 승인 유지 텍스트 apply",
            "장기: Form Fill North Star · Pack 추가 = 새 문서 종류",
            "원칙: holdout 봉인 · gold 분리 · 구조화 필드 LLM 금지",
        ],
    )
    bullets(
        prs,
        "요약",
        [
            "최종 목표: Form Fill (빈 양식·케이스로 문서 자동 작성)",
            "오늘 초점: Change Impact — B1~B5로 찾기·승인·안전 복사",
            "성능은 Holdout·Safety·Pilot로 층위별 측정",
            "장기: Form Fill North Star · Pack으로 문서 종류 확장",
        ],
    )
    title_slide(
        prs,
        "Q & A",
        "Document AI × AI Engineering\nB1~B5 · 평가 · 안전 · 산출물",
    )

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    add_page_numbers(prs)
    try:
        prs.save(str(PPTX_PATH))
        out = PPTX_PATH
    except PermissionError:
        prs.save(str(PPTX_PATH_ALT))
        out = PPTX_PATH_ALT
        print("NOTE: original pptx locked; saved as", out.name)
    print("slides", len(prs.slides))
    print(out)

    MD_PATH.write_text(
        """# Document AI × AI Engineering — 구성

PPT: `Document_AI_v2_Presentation.pptx`  
스크립트: `SPEAKER_SCRIPT.md` (**PPT 페이지별 p.1~p.35**)  
산출물: `artifacts/req11_writer_on/` · `screenshots/`

## 발표 순서
1. **정의** — 최종 목표(Form Fill) vs 오늘 초점(Change Impact) + 비교 표
2. 배경 · AI Engineering · North Star vs 현재
3. **B1~B5 (Pilot 화면 캡처)** + 코드 흐름
4. 벤치·Pilot 수치
5. 산출물 / Writer OFF / Pack 확장 · 향후

## 기억할 구분
- **최종 목표**: 빈 양식 + 케이스로 문서 자동 작성 (Form Fill)
- **오늘 초점**: 변경 영향 찾기 → 보여 주기 → 승인 → 복사본만 (Change Impact / B1~B5)
""",
        encoding="utf-8",
    )
    # SPEAKER_SCRIPT.md is maintained as page-by-page script (do not overwrite here)
    print("script:", SCRIPT_PATH, "(manual page-by-page)")


if __name__ == "__main__":
    build()
