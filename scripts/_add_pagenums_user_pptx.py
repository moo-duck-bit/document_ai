"""Add visible n/total page numbers to user's edited PPTX."""
import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

GRAY = RGBColor(0x5A, 0x6A, 0x7A)

DOWNLOADS = Path(os.environ.get("PPTX_SRC_DIR") or Path.home() / "Downloads")
SRC = DOWNLOADS / "Document_AI_v2_Presentation_수정본.pptx"
OUT_DL = DOWNLOADS / "Document_AI_v2_Presentation_수정본_페이지번호.pptx"
OUT_REPO = (
    Path(__file__).resolve().parents[1]
    / "docs" / "presentation" / "Document_AI_v2_Presentation_수정본_페이지번호.pptx"
)


def is_old_pagenum(shape, n: int) -> bool:
    if not shape.has_text_frame:
        return False
    t = shape.text_frame.text.strip()
    if t != str(n):
        return False
    if (shape.top or 0) < 6_200_000:
        return False
    if (shape.left or 0) < 10_000_000:
        return False
    return True


def clear_old_pagenums(slide, n: int) -> int:
    sp_tree = slide.shapes._spTree
    to_remove = [
        sh._element for sh in list(slide.shapes) if is_old_pagenum(sh, n)
    ]
    for el in to_remove:
        sp_tree.remove(el)
    return len(to_remove)


def main() -> None:
    prs = Presentation(str(SRC))
    total = len(prs.slides)
    removed = 0
    for i, slide in enumerate(prs.slides, 1):
        removed += clear_old_pagenums(slide, i)
        tb = slide.shapes.add_textbox(
            Inches(11.2), Inches(7.05), Inches(1.9), Inches(0.35)
        )
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = f"{i} / {total}"
        run.font.size = Pt(11)
        run.font.color.rgb = GRAY
        run.font.name = "Malgun Gothic"

    OUT_DL.parent.mkdir(parents=True, exist_ok=True)
    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_DL))
    prs.save(str(OUT_REPO))
    print(f"saved: {OUT_DL}")
    print(f"saved: {OUT_REPO}")
    print(f"removed old footers: {removed}")
    print(f"slides: {total}")


if __name__ == "__main__":
    main()
