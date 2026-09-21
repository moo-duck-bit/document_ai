#!/usr/bin/env python3
"""Audit generated MDSR/MDDR outputs for a case — human review checklist helper."""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document

from document_ai.learn.docx_io import paragraph_deep_text, table_matrix
from document_ai.render.mdsr import UNIQUE_MINDRIUM_PATTERN, verify_completeness
from document_ai.render.mddr import verify_mddr_completeness

REQ_HEADING = re.compile(r"^Req\.\s*\d+", re.IGNORECASE)


def audit_case(case_dir: Path) -> dict:
    case_dir = case_dir.resolve()
    result: dict = {"case": str(case_dir), "mdsr": {}, "mddr": {}, "ok": True}

    req_path = case_dir / "requirements.json"
    design_path = case_dir / "design_items.json"
    mdsr_doc = case_dir / "output_mdsr.docx"
    mddr_doc = case_dir / "output_mddr.docx"

    requirements = json.loads(req_path.read_text(encoding="utf-8")) if req_path.exists() else {"requirements": []}
    design_items = json.loads(design_path.read_text(encoding="utf-8")).get("items", []) if design_path.exists() else []

    if mdsr_doc.exists():
        doc = Document(str(mdsr_doc))
        mindrium = [
            paragraph_deep_text(p)[:80]
            for p in doc.paragraphs
            if UNIQUE_MINDRIUM_PATTERN.search(paragraph_deep_text(p))
        ]
        req_count = sum(
            1
            for t in doc.tables
            if t.rows and re.match(r"^Req\.", t.rows[0].cells[0].text.strip())
        )
        review = verify_completeness(doc, requirements.get("requirements", []))
        result["mdsr"] = {
            "path": str(mdsr_doc),
            "req_tables": req_count,
            "mindrium_hits": len(mindrium),
            "review": review,
        }
        if mindrium or review["issue_count"] or review["residual_count"]:
            result["ok"] = False
    else:
        result["mdsr"] = {"missing": str(mdsr_doc)}
        result["ok"] = False

    if mddr_doc.exists():
        doc = Document(str(mddr_doc))
        review = verify_mddr_completeness(doc, design_items)
        result["mddr"] = {
            "path": str(mddr_doc),
            "review": review,
        }
        if review["issue_count"] or review["residual_count"]:
            result["ok"] = False
    else:
        result["mddr"] = {"missing": str(mddr_doc)}
        result["ok"] = False

    return result


def format_report(report: dict) -> str:
    lines = [
        f"# Case audit - {report['case']}",
        "",
        f"**Overall:** {'PASS' if report['ok'] else 'NEEDS REVIEW'}",
        "",
        "## MDSR",
    ]
    mdsr = report.get("mdsr", {})
    if "missing" in mdsr:
        lines.append(f"- Missing: {mdsr['missing']}")
    else:
        lines.extend(
            [
                f"- Req. tables: {mdsr.get('req_tables', 0)}",
                f"- Mindrium residual: {mdsr.get('mindrium_hits', 0)}",
                f"- Auto review issues: {mdsr.get('review', {}).get('issue_count', '?')}",
                f"- Auto review residual: {mdsr.get('review', {}).get('residual_count', '?')}",
            ]
        )
    lines.extend(["", "## MDDR"])
    mddr = report.get("mddr", {})
    if "missing" in mddr:
        lines.append(f"- Missing: {mddr['missing']}")
    else:
        rev = mddr.get("review", {})
        lines.extend(
            [
                f"- Design blocks filled: {rev.get('design_blocks_filled', '?')} / {rev.get('design_items_expected', '?')}",
                f"- Auto review issues: {rev.get('issue_count', '?')}",
            ]
        )
        if rev.get("issues"):
            lines.append("- Issues:")
            for issue in rev["issues"][:10]:
                lines.append(f"  - {issue}")
    lines.extend(
        [
            "",
            "## Human review (Word)",
            "- [ ] 표지·승인자 이름/직함/날짜",
            "- [ ] 3. 제품 개요 하위 항목",
            "- [ ] 4. 고유 요구사항 (도메인 문구)",
            "- [ ] 5. Req. 설명·목적·기준",
            "- [ ] 추적성 매트릭스",
            "- [ ] MDDR 설계 항목·Figure 캡션",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit case MDSR/MDDR outputs")
    parser.add_argument("--case", default="data/cases/jm_collection", help="Case directory")
    parser.add_argument("--md", help="Write markdown report to path")
    args = parser.parse_args()

    report = audit_case(ROOT / args.case)
    text = format_report(report)
    print(text)
    if args.md:
        out = Path(args.md)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(text, encoding="utf-8")
        print(f"Wrote {out}")
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
