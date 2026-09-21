# -*- coding: utf-8 -*-
"""Generate document_set_benchmark dataset fixtures and golden labels."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document

from document_ai.domain_packs.ec_sw.mdtm_indexer import index_mdtm_document
from document_ai.template.generic_templates import build_general_report_template

REPO = Path(__file__).resolve().parents[1]
BASE = REPO / "data" / "eval" / "document_set_benchmark"
FIX = BASE / "fixtures"
CASES = BASE / "cases"
LABELS = BASE / "labels"


def write_report(path: Path, paragraphs: list[str]) -> None:
    doc = Document()
    doc.add_heading("일반 보고서", level=1)
    for p in paragraphs:
        if p.startswith("#"):
            doc.add_heading(p.lstrip("#").strip(), level=2)
        else:
            doc.add_paragraph(p)
    doc.save(str(path))


def main() -> None:
    for p in [
        FIX / "ec_sw",
        FIX / "general_report",
        CASES / "ec_sw",
        CASES / "general_report",
        LABELS,
    ]:
        p.mkdir(parents=True, exist_ok=True)

    src = next((REPO / "data" / "examples" / "ec_sw").glob("matrix_mdtm*.docx"))
    mdtm_fix = FIX / "ec_sw" / "MDTM.docx"
    shutil.copy2(src, mdtm_fix)

    stub = FIX / "ec_sw" / "MDSR_stub.docx"
    d = Document()
    d.add_paragraph("MDSR stub Req. 2")
    d.save(str(stub))

    nodes = index_mdtm_document(source_path=mdtm_fix, document_id="MDTM")["nodes"]
    req_node: dict[str, dict] = {}
    for n in nodes:
        if n.node_type != "TABLE_ROW":
            continue
        for r in (n.source_identifiers or {}).get("requirement_ids") or []:
            req_node[r] = {
                "node_id": n.node_id,
                "design_ids": list((n.source_identifiers or {}).get("design_ids") or []),
                "test_ids": list((n.source_identifiers or {}).get("test_ids") or []),
            }

    req1, req2, req3 = req_node["Req. 1"], req_node["Req. 2"], req_node["Req. 3"]
    rid101 = "Req. 101" if "Req. 101" in req_node else "Req. 10"

    write_report(
        FIX / "general_report" / "report_base.docx",
        [
            "# 방법론",
            "데이터 출처는 내부 설문이다.",
            "# 결과",
            "핵심 발견은 만족도 상승이다.",
            "# 결론",
            "권고사항은 추가 검증이다.",
            "# 일정",
            "2026-Q3 완료 예정",
        ],
    )
    write_report(
        FIX / "general_report" / "report_dup_heading.docx",
        [
            "# 결과",
            "첫번째 결과 섹션",
            "# 결과",
            "두번째 결과 섹션 중복",
            "# 결론",
            "권고사항",
        ],
    )
    td = Document()
    td.add_heading("일반 보고서", 1)
    td.add_heading("결과", 2)
    t = td.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "지표"
    t.cell(0, 1).text = "값"
    t.cell(1, 0).text = "만족도"
    t.cell(1, 1).text = "70"
    td.add_heading("결론", 2)
    td.add_paragraph("권고사항은 유지한다.")
    td.save(str(FIX / "general_report" / "report_table.docx"))
    (FIX / "general_report" / "report_base.md").write_text(
        "# 일반 보고서\n## 방법론\n데이터 출처는 내부 설문이다.\n## 결과\n핵심 발견\n## 결론\n권고사항\n",
        encoding="utf-8",
    )

    design = req2["design_ids"][0] if req2["design_ids"] else "4.2.2"
    test = req2["test_ids"][0] if req2["test_ids"] else "4.1.1.2"

    sections = [s.section_id for s in build_general_report_template().sections]
    # Prefer real template section ids when present
    methodology = "methodology" if "methodology" in sections else "background"
    results = "results" if "results" in sections else "findings"
    conclusion = "conclusion" if "conclusion" in sections else "recommendations"
    meth_node = f"general_report_v1.{methodology}"
    res_node = f"general_report_v1.{results}"
    conc_node = f"general_report_v1.{conclusion}"

    ec_cases = [
        {
            "case_id": "ec_sw_exact_req_single",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 2 행 추적성 갱신",
            "difficulty": "easy",
            "tags": ["exact_id", "writer_blocked"],
            "notes": "Exact Req. 2 single row gold PATCH",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "IMPACTED",
            "gold_nodes": [(req2["node_id"], "PATCH_CANDIDATE")],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_exact_req_multi_mention",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 2 및 동일 Req. 2 재확인",
            "difficulty": "easy",
            "tags": ["exact_id"],
            "notes": "Same Req mentioned twice still one row",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "IMPACTED",
            "gold_nodes": [(req2["node_id"], "PATCH_CANDIDATE")],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_multi_req_ids",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 1과 Req. 3을 동시에 수정",
            "difficulty": "medium",
            "tags": ["exact_id", "multi_document"],
            "notes": "Two exact req patches",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "IMPACTED",
            "gold_nodes": [
                (req1["node_id"], "PATCH_CANDIDATE"),
                (req3["node_id"], "PATCH_CANDIDATE"),
            ],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_design_id_only",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": f"설계 ID {design} 관련 추적성 검토",
            "difficulty": "hard",
            "tags": ["semantic_only"],
            "notes": "Design-id only; workflow CR may under-match",
            "expected_status": "SAFE_FAILURE",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [(req2["node_id"], "REVIEW_REQUIRED")],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_test_id_only",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": f"시험 ID {test} 결과 반영",
            "difficulty": "hard",
            "tags": ["semantic_only"],
            "notes": "Test-id only via CR text",
            "expected_status": "SAFE_FAILURE",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [(req2["node_id"], "REVIEW_REQUIRED")],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_semantic_only",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "사용자 인증 및 접근통제 보안 요구 변경",
            "difficulty": "medium",
            "tags": ["semantic_only"],
            "notes": "No Req ID; semantic review expected",
            "expected_status": "SAFE_FAILURE",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_duplicate_req_row",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 2 중복 행 가능성 검토",
            "difficulty": "medium",
            "tags": ["duplicate_reference"],
            "notes": "Corpus has unique Req.2; exact patch",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "IMPACTED",
            "gold_nodes": [(req2["node_id"], "PATCH_CANDIDATE")],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_missing_identifier",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "요구사항 식별자 없이 행 내용만 변경 요청",
            "difficulty": "hard",
            "tags": ["ambiguous"],
            "notes": "Missing identifier",
            "expected_status": "SAFE_FAILURE",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "UNRELATED",
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_empty_row_context",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "빈 식별자 행에 대한 정리",
            "difficulty": "hard",
            "tags": ["ambiguous"],
            "notes": "Empty-row style request",
            "expected_status": "SAFE_FAILURE",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "UNRELATED",
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_no_impact",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "마케팅 슬로건 XYZABCQUUX 변경",
            "difficulty": "easy",
            "tags": ["no_impact"],
            "notes": "No-impact marketing phrase",
            "expected_status": "SUCCESS",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "UNRELATED",
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_ambiguous_request",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "관련 항목을 적절히 수정해 주세요",
            "difficulty": "hard",
            "tags": ["ambiguous"],
            "notes": "Ambiguous; acceptable UNRELATED or REVIEW",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "UNRELATED",
            "gold_doc_alts": ["REVIEW_REQUIRED"],
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_writer_blocked_review",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 2 자동 반영 금지 검토 필요",
            "difficulty": "medium",
            "tags": ["writer_blocked", "writer_update"],
            "notes": "Writer must remain gated",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "IMPACTED",
            "gold_nodes": [(req2["node_id"], "PATCH_CANDIDATE")],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_bad_req_format",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "REQ2 형식 오류 식별자",
            "difficulty": "medium",
            "tags": ["ambiguous"],
            "notes": "Bad Req format should not false-patch",
            "expected_status": "SUCCESS",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "UNRELATED",
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_mdtm_only_multi_doc",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": "Req. 2 MDTM만 영향",
            "difficulty": "medium",
            "tags": ["multi_document", "exact_id"],
            "notes": "MDTM IMPACTED; MDSR stub UNRELATED",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                },
                {
                    "path": "fixtures/ec_sw/MDSR_stub.docx",
                    "filename": "MDSR_stub.docx",
                    "role": "requirements",
                },
            ],
            "gold_doc_multi": [("MDTM", "IMPACTED"), ("MDSR_STUB", "UNRELATED")],
            "gold_nodes": [(req2["node_id"], "PATCH_CANDIDATE")],
            "should_write": False,
        },
        {
            "case_id": "ec_sw_req101_exact",
            "domain": "ec_sw",
            "document_set_id": "ec_sw",
            "change_request": f"{rid101} 단일 행 갱신",
            "difficulty": "easy",
            "tags": ["exact_id"],
            "notes": f"Exact {rid101}",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/ec_sw/MDTM.docx",
                    "filename": "MDTM.docx",
                    "role": "traceability",
                }
            ],
            "gold_doc": "IMPACTED",
            "gold_nodes": [(req_node[rid101]["node_id"], "PATCH_CANDIDATE")],
            "should_write": False,
        },
    ]

    gr_cases = [
        {
            "case_id": "gr_methodology_source",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "방법론 데이터 출처를 외부 공개자료로 수정",
            "difficulty": "easy",
            "tags": ["writer_update"],
            "notes": "Methodology section token overlap",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_base.docx",
                    "filename": "report_base.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [(meth_node, "REVIEW_REQUIRED")],
            "should_write": False,
        },
        {
            "case_id": "gr_results_finding",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "결과 핵심 발견 문구 수정",
            "difficulty": "easy",
            "tags": ["writer_update"],
            "notes": "Results section",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_base.docx",
                    "filename": "report_base.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [(res_node, "REVIEW_REQUIRED")],
            "should_write": False,
        },
        {
            "case_id": "gr_conclusion_recommendation",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "결론 권고사항 추가",
            "difficulty": "easy",
            "tags": ["writer_update"],
            "notes": "Conclusion recommendation",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_base.docx",
                    "filename": "report_base.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [(conc_node, "REVIEW_REQUIRED")],
            "should_write": False,
        },
        {
            "case_id": "gr_schedule_table",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "일정 표 수정 2026-Q4",
            "difficulty": "medium",
            "tags": ["writer_update"],
            "notes": "Schedule section",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_base.docx",
                    "filename": "report_base.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "gr_duplicate_heading",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "결과 섹션 중복 heading 정리",
            "difficulty": "hard",
            "tags": ["ambiguous", "duplicate_reference"],
            "notes": "Duplicate heading fixture",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_dup_heading.docx",
                    "filename": "report_dup_heading.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [(res_node, "REVIEW_REQUIRED")],
            "should_write": False,
        },
        {
            "case_id": "gr_missing_section",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "존재하지 않는 부록Z 섹션 수정",
            "difficulty": "medium",
            "tags": ["no_impact"],
            "notes": "Nonexistent section",
            "expected_status": "SUCCESS",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_base.docx",
                    "filename": "report_base.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "UNRELATED",
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "gr_semantic_similar",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "성과 요약과 유사한 핵심 내용 보강",
            "difficulty": "hard",
            "tags": ["semantic_only"],
            "notes": "Semantic similar to summary/results",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_base.docx",
                    "filename": "report_base.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_doc_alts": ["UNRELATED"],
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "gr_no_impact",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "사무실 화분 배치 변경 XYZABC",
            "difficulty": "easy",
            "tags": ["no_impact"],
            "notes": "No impact",
            "expected_status": "SUCCESS",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_base.docx",
                    "filename": "report_base.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "UNRELATED",
            "gold_nodes": [],
            "should_write": False,
        },
        {
            "case_id": "gr_table_cell",
            "domain": "general_report",
            "document_set_id": "general_report",
            "change_request": "결과 표 만족도 값 수정",
            "difficulty": "medium",
            "tags": ["writer_update"],
            "notes": "Table cell style request",
            "expected_status": "PARTIAL",
            "input_documents": [
                {
                    "path": "fixtures/general_report/report_table.docx",
                    "filename": "report_table.docx",
                    "role": "general_report",
                }
            ],
            "gold_doc": "REVIEW_REQUIRED",
            "gold_nodes": [(res_node, "REVIEW_REQUIRED")],
            "should_write": False,
        },
    ]

    all_cases = ec_cases + gr_cases
    manifest_cases = []
    doc_impacts = []
    node_impacts = []
    patch_exp = []
    writer_exp = []

    for c in all_cases:
        case_body = {
            "case_id": c["case_id"],
            "domain": c["domain"],
            "document_set_id": c["document_set_id"],
            "change_request": c["change_request"],
            "input_documents": c["input_documents"],
            "enabled_documents": [
                d.get("document_id") or Path(d["filename"]).stem.upper()
                for d in c["input_documents"]
            ],
            "expected_status": c["expected_status"],
            "tags": c["tags"],
            "difficulty": c["difficulty"],
            "notes": c["notes"],
            "source_type": "fixture",
        }
        rel = f"cases/{c['domain']}/{c['case_id']}.json"
        (BASE / rel).write_text(
            json.dumps(case_body, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        manifest_cases.append(rel)

        if c.get("gold_doc_multi"):
            for did, st in c["gold_doc_multi"]:
                row = {
                    "case_id": c["case_id"],
                    "document_id": did,
                    "gold_status": st,
                    "review_notes": c["notes"],
                }
                if c.get("gold_doc_alts"):
                    row["acceptable_alternatives"] = c["gold_doc_alts"]
                doc_impacts.append(row)
        else:
            fname = c["input_documents"][0]["filename"]
            did = "MDTM" if c["domain"] == "ec_sw" else Path(fname).stem.upper()
            row = {
                "case_id": c["case_id"],
                "document_id": did,
                "gold_status": c["gold_doc"],
                "review_notes": c["notes"],
            }
            if c.get("gold_doc_alts"):
                row["acceptable_alternatives"] = c["gold_doc_alts"]
            doc_impacts.append(row)

        for nid, st in c.get("gold_nodes") or []:
            did = (
                "MDTM"
                if c["domain"] == "ec_sw"
                else Path(c["input_documents"][0]["filename"]).stem.upper()
            )
            node_impacts.append(
                {
                    "case_id": c["case_id"],
                    "document_id": did,
                    "node_id": nid,
                    "gold_status": st,
                }
            )
            patch_exp.append(
                {
                    "case_id": c["case_id"],
                    "document_id": did,
                    "node_id": nid,
                    "should_create_patch_candidate": st == "PATCH_CANDIDATE",
                    "expected_operation": "UPDATE",
                    "human_review_required": st != "PATCH_CANDIDATE",
                }
            )
        writer_exp.append(
            {
                "case_id": c["case_id"],
                "should_write": bool(c.get("should_write")),
                "expected_result_status": "GATED",
                "original_must_remain_unchanged": True,
                "expected_changed_document_ids": [],
                "format_preservation_required": True,
            }
        )

    manifest = {
        "meta": {
            "name": "document_set_benchmark_v1",
            "version": "1.0",
            "case_count": len(all_cases),
            "template_sections": sections,
        },
        "cases": manifest_cases,
    }
    (BASE / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    def write_jsonl(path: Path, rows: list) -> None:
        path.write_text(
            "\n".join(json.dumps(r, ensure_ascii=False) for r in rows) + "\n",
            encoding="utf-8",
        )

    write_jsonl(LABELS / "document_impacts.jsonl", doc_impacts)
    write_jsonl(LABELS / "node_impacts.jsonl", node_impacts)
    write_jsonl(LABELS / "patch_expectations.jsonl", patch_exp)
    write_jsonl(LABELS / "writer_expectations.jsonl", writer_exp)
    print(f"cases={len(all_cases)} ec={len(ec_cases)} gr={len(gr_cases)}")
    print(f"req2={req2['node_id']}")
    print(f"sections={sections}")


if __name__ == "__main__":
    main()
