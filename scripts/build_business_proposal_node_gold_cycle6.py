# -*- coding: utf-8 -*-
"""Document AI Cycle 6: Business Proposal Node Gold Labeling & Blind Evaluation.

Builds REQUIRED/AMBIGUOUS/OPTIONAL/NOT_APPLICABLE node gold for every
``business_proposal`` case (existing + newly added), from document structure
and change_request text ONLY (python-docx inspection + deterministic policy —
never reads prediction/ranking artifacts, never auto-applies
``business_proposal_proposed_label_changes.json``, never uses prediction
Top-1 as gold).

Steps
-----
1. Add new fixtures + cases (idempotent) to reach REQUIRED>=8, AMBIGUOUS>=2,
   OPTIONAL>=2, NOT_APPLICABLE>=2.
2. Audit all BP cases (old label vs. new gold, decision table).
3. Inventory every BP fixture document (headings/paragraphs/tables/sections).
4. Pass1 label (labeled_by=pass1_structure_policy).
5. Pass2 label (labeled_by=pass2_independent_review) — same deterministic
   policy, independent inventory re-derivation; record disagreements.
6. Validate + finalize agreed rows to sealed/.
7. Project sealed gold into development/labels + holdout/sealed_labels
   (merging with, not overwriting, non-BP rows).
8. Re-seal holdout via ``write_seal_manifest``.
9. Write summary + agreement metrics (raw / mode / reference agreement).

Does NOT touch ranking / structural_match / proposal_table_retrieval / query_intent
scoring logic, does NOT overwrite prior benchmark run result directories, and
does NOT modify ``data/examples`` or ``data/freeze``.
"""

from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document

REPO = Path(__file__).resolve().parents[1]
SRC = REPO / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from document_ai.evaluation.business_proposal_gold import (  # noqa: E402
    audit as audit_mod,
    labeling_pass,
    protocol,
    validation as validation_mod,
)
from document_ai.evaluation.business_proposal_gold.document_inventory import (  # noqa: E402
    inventory_document,
)
from document_ai.evaluation.business_proposal_gold.project_to_benchmark import (  # noqa: E402
    project_gold_into_benchmark,
)
from document_ai.evaluation.business_proposal_gold.schema import (  # noqa: E402
    BusinessProposalGoldRow,
)

V2 = REPO / "data" / "eval" / "document_set_benchmark_v2"
FX = V2 / "fixtures" / "business_proposal"
NOW = datetime.now(timezone.utc).isoformat()


# ---------------------------------------------------------------------------
# 1. New fixtures (python-docx, structure only)
# ---------------------------------------------------------------------------


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            table.rows[i].cells[j].text = v


def build_proposal_schedule_table(path: Path) -> None:
    doc = Document()
    doc.add_heading("사업 제안서", level=0)
    doc.add_heading("일정", level=1)
    doc.add_paragraph("본 사업의 일정은 아래 표와 같습니다.")
    _add_table(
        doc,
        ["단계", "기간", "시작", "종료"],
        [
            ["분석", "4주", "2026-01-01", "2026-01-28"],
            ["설계", "6주", "2026-01-29", "2026-03-11"],
            ["구현", "8주", "2026-03-12", "2026-05-06"],
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def build_proposal_org_table(path: Path) -> None:
    doc = Document()
    doc.add_heading("사업 제안서", level=0)
    doc.add_heading("수행 조직", level=1)
    doc.add_paragraph("본 사업의 수행 조직 구성은 다음과 같습니다.")
    _add_table(
        doc,
        ["조직", "역할", "담당"],
        [
            ["PM", "총괄", "홍길동"],
            ["개발팀", "구현", "김철수"],
            ["QA팀", "품질보증", "이영희"],
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def build_proposal_kpi(path: Path) -> None:
    doc = Document()
    doc.add_heading("사업 제안서", level=0)
    doc.add_heading("기대 효과", level=1)
    doc.add_paragraph("본 사업의 기대 효과는 다음 KPI로 측정합니다.")
    _add_table(
        doc,
        ["KPI", "목표", "지표"],
        [
            ["매출증가", "10%", "분기매출"],
            ["고객만족", "90점", "설문점수"],
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def build_proposal_deliverables(path: Path) -> None:
    doc = Document()
    doc.add_heading("사업 제안서", level=0)
    doc.add_heading("산출물", level=1)
    doc.add_paragraph(
        "요구사항 정의서, 설계 문서, 소스코드, 시험 결과 보고서를 산출물로 제출합니다."
    )
    doc.add_paragraph("- 요구사항 정의서")
    doc.add_paragraph("- 설계 문서")
    doc.add_paragraph("- 시험 결과 보고서")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


NEW_FIXTURE_BUILDERS = {
    "proposal_schedule_table.docx": build_proposal_schedule_table,
    "proposal_org_table.docx": build_proposal_org_table,
    "proposal_kpi.docx": build_proposal_kpi,
    "proposal_deliverables.docx": build_proposal_deliverables,
}


def build_new_fixtures() -> list[str]:
    written = []
    for name, builder in NEW_FIXTURE_BUILDERS.items():
        path = FX / name
        builder(path)
        written.append(str(path).replace("\\", "/"))
    return written


# ---------------------------------------------------------------------------
# 2. New case definitions
# ---------------------------------------------------------------------------

NEW_DEV_CASES = [
    dict(
        case_id="v2_dev_bp_sched_table",
        cr="수행 일정 표를 수정",
        fixture="proposal_schedule_table.docx",
        tags=["sched_table", "schedule", "table"],
    ),
    dict(
        case_id="v2_dev_bp_labor_cell",
        cr="인건비 금액을 수정",
        fixture="proposal_base.docx",
        tags=["labor_cell", "budget", "table"],
    ),
    dict(
        case_id="v2_dev_bp_kpi",
        cr="KPI 표를 수정",
        fixture="proposal_kpi.docx",
        tags=["kpi", "table"],
    ),
    dict(
        case_id="v2_dev_bp_deliverable",
        cr="산출물 목록 수정",
        fixture="proposal_deliverables.docx",
        tags=["deliverable"],
    ),
    dict(
        case_id="v2_dev_bp_market_add",
        cr="시장분석 절을 추가",
        fixture="proposal_base.docx",
        tags=["market_add"],
    ),
    dict(
        case_id="v2_dev_bp_style_overall",
        cr="문서 전반 표현 개선",
        fixture="proposal_base.docx",
        tags=["style_overall"],
    ),
    dict(
        case_id="v2_dev_bp_schedule_section",
        cr="일정 섹션 전반 수정",
        fixture="proposal_base.docx",
        tags=["schedule_section"],
    ),
]

NEW_HOLDOUT_CASES = [
    dict(
        case_id="v2_hol_bp_org_table",
        cr="수행 조직 역할표 변경",
        fixture="proposal_org_table.docx",
        tags=["org_table", "organization", "table", "holdout"],
    ),
    dict(
        case_id="v2_hol_bp_effect_para",
        cr="기대효과 문단 수정",
        fixture="proposal_base.docx",
        tags=["effect_para", "holdout"],
    ),
    dict(
        case_id="v2_hol_bp_risk_amb",
        cr="위험관리 내용 보완",
        fixture="proposal_base.docx",
        tags=["risk_amb", "holdout"],
    ),
]


def _case_json(case_id: str, cr: str, fixture: str, tags: list[str], split: str) -> dict[str, Any]:
    doc_id = Path(fixture).stem.upper()
    return {
        "case_id": case_id,
        "domain": "business_proposal",
        "document_set_id": "business_proposal",
        "change_request": cr,
        "split": split,
        "input_documents": [
            {
                "path": f"fixtures/business_proposal/{fixture}",
                "filename": fixture,
                "role": "general_report",
            }
        ],
        "enabled_documents": [doc_id],
        "expected_status": "SUCCESS",
        "tags": tags,
        "difficulty": "medium",
        "notes": "cycle6_business_proposal_node_gold",
        "source_type": "GENERATED",
        "transformation": None,
        "base_case_id": None,
        "fixture_fingerprint": _sha256(FX / fixture),
        "sanitization_status": "N/A",
    }


def write_new_cases() -> dict[str, list[str]]:
    written = {"development": [], "holdout": []}
    for spec in NEW_DEV_CASES:
        case = _case_json(spec["case_id"], spec["cr"], spec["fixture"], spec["tags"], "development")
        path = V2 / "development" / "cases" / "business_proposal" / f"{spec['case_id']}.json"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(case, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        written["development"].append(f"development/cases/business_proposal/{spec['case_id']}.json")
    for spec in NEW_HOLDOUT_CASES:
        case = _case_json(spec["case_id"], spec["cr"], spec["fixture"], spec["tags"], "holdout")
        path = V2 / "holdout" / "cases" / "business_proposal" / f"{spec['case_id']}.json"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(case, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        written["holdout"].append(f"holdout/cases/business_proposal/{spec['case_id']}.json")
    return written


def update_manifest(new_rels: dict[str, list[str]]) -> None:
    manifest_path = V2 / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    dev_list = manifest.setdefault("development_cases", [])
    hol_list = manifest.setdefault("holdout_cases", [])
    for rel in new_rels["development"]:
        if rel not in dev_list:
            dev_list.append(rel)
    for rel in new_rels["holdout"]:
        if rel not in hol_list:
            hol_list.append(rel)
    manifest.setdefault("meta", {})["development_count"] = len(dev_list)
    manifest["meta"]["holdout_count"] = len(hol_list)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# 3. document_impacts / writer_expectations for new cases
# ---------------------------------------------------------------------------


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    body = "\n".join(json.dumps(r, ensure_ascii=False) for r in rows)
    path.write_text((body + "\n") if rows else "", encoding="utf-8")


def update_document_and_writer_labels(new_case_ids_by_split: dict[str, list[str]]) -> None:
    for split, label_dir in (
        ("development", V2 / "development" / "labels"),
        ("holdout", V2 / "holdout" / "sealed_labels"),
    ):
        case_ids = set(new_case_ids_by_split[split])
        if not case_ids:
            continue
        doc_path = label_dir / "document_impacts.jsonl"
        writer_path = label_dir / "writer_expectations.jsonl"
        docs = [r for r in _read_jsonl(doc_path) if r.get("case_id") not in case_ids]
        writers = [r for r in _read_jsonl(writer_path) if r.get("case_id") not in case_ids]
        cases = NEW_DEV_CASES if split == "development" else NEW_HOLDOUT_CASES
        for spec in cases:
            doc_id = Path(spec["fixture"]).stem.upper()
            docs.append(
                {
                    "case_id": spec["case_id"],
                    "document_id": doc_id,
                    "gold_status": "REVIEW_REQUIRED",
                    "review_notes": "cycle6_structure_based_pre_prediction_label",
                }
            )
            writers.append(
                {
                    "case_id": spec["case_id"],
                    "should_write": False,
                    "expected_result_status": "GATED",
                    "original_must_remain_unchanged": True,
                    "expected_changed_document_ids": [],
                    "format_preservation_required": True,
                }
            )
        _write_jsonl(doc_path, sorted(docs, key=lambda r: r["case_id"]))
        _write_jsonl(writer_path, sorted(writers, key=lambda r: r["case_id"]))


# ---------------------------------------------------------------------------
# 4. Collect all BP cases (existing + new) from the manifest
# ---------------------------------------------------------------------------


def collect_all_bp_cases() -> list[dict[str, Any]]:
    manifest = json.loads((V2 / "manifest.json").read_text(encoding="utf-8"))
    cases = []
    for split, key in (("development", "development_cases"), ("holdout", "holdout_cases")):
        for rel in manifest.get(key) or []:
            cpath = V2 / rel
            if not cpath.is_file():
                continue
            case = json.loads(cpath.read_text(encoding="utf-8"))
            if case.get("domain") != "business_proposal":
                continue
            case = dict(case)
            case["split"] = case.get("split") or split
            cases.append(case)
    return sorted(cases, key=lambda c: c["case_id"])


# True pre-Cycle6 baseline for the 12 BP cases that existed before this script ran,
# taken verbatim from ``generate_document_set_benchmark_v2.py`` (business_proposal
# branch: every case got a placeholder OPTIONAL, or NOT_APPLICABLE for "no_impact"
# tagged cases — there was no real REQUIRED/AMBIGUOUS distinction pre-Cycle6).
# Hardcoded (rather than re-read from disk) because ``node_evaluation_eligibility.jsonl``
# is overwritten in-place by ``project_gold_into_benchmark`` on every run, so re-reading
# it after the first successful run would trivially compare Cycle6-gold against itself.
_PRE_CYCLE6_BP_BASELINE_MODE: dict[str, str] = {
    "v2_dev_bp_schedule": "OPTIONAL",
    "v2_dev_bp_budget": "OPTIONAL",
    "v2_dev_bp_risk": "OPTIONAL",
    "v2_dev_bp_org": "OPTIONAL",
    "v2_dev_bp_effect": "OPTIONAL",
    "v2_dev_bp_dup_sched": "OPTIONAL",
    "v2_dev_bp_en": "OPTIONAL",
    "v2_dev_bp_no_impact": "NOT_APPLICABLE",
    "v2_hol_bp_sched": "OPTIONAL",
    "v2_hol_bp_budget": "OPTIONAL",
    "v2_hol_bp_risk": "OPTIONAL",
    "v2_hol_bp_no_impact": "NOT_APPLICABLE",
}


def _old_eligibility_by_case() -> dict[str, dict[str, Any]]:
    return {
        case_id: {"case_id": case_id, "node_evaluation_mode": mode}
        for case_id, mode in _PRE_CYCLE6_BP_BASELINE_MODE.items()
    }


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------


def main() -> dict[str, Any]:
    old_eligibility = _old_eligibility_by_case()

    fixtures_written = build_new_fixtures()
    new_case_rels = write_new_cases()
    update_manifest(new_case_rels)
    new_case_ids_by_split = {
        "development": [s["case_id"] for s in NEW_DEV_CASES],
        "holdout": [s["case_id"] for s in NEW_HOLDOUT_CASES],
    }
    update_document_and_writer_labels(new_case_ids_by_split)

    all_cases = collect_all_bp_cases()

    pass1_rows = labeling_pass.run_pass1(all_cases, fixtures_root=FX)
    pass2_rows = labeling_pass.run_pass2(all_cases, fixtures_root=FX)
    disagreements = labeling_pass.detect_disagreements(pass1_rows, pass2_rows)
    agreement = labeling_pass.agreement_metrics(pass1_rows, pass2_rows)

    # Inventories for validation cross-checks (wrong-document / unknown-node reject).
    inventories: dict[str, dict[str, Any]] = {}
    change_requests: dict[str, str] = {}
    cases_by_id = {c["case_id"]: c for c in all_cases}
    for case in all_cases:
        docx_path = labeling_pass.resolve_fixture(case, FX)
        doc_id = (case.get("enabled_documents") or [None])[0] or docx_path.stem.upper()
        inventories[case["case_id"]] = inventory_document(docx_path, document_id=doc_id)
        change_requests[case["case_id"]] = case.get("change_request") or ""

    draft_dicts = [r.to_dict() for r in pass1_rows]
    review_dicts = [r.to_dict() for r in pass2_rows]
    protocol.write_draft(draft_dicts)
    protocol.write_review(review_dicts)

    validation_report = validation_mod.validate_gold_rows(
        pass1_rows, inventories=inventories, change_requests=change_requests
    )

    # Finalize: only rows that agree (pass1 == pass2 on substantive fields) AND
    # pass validation are sealed. In this deterministic-policy design pass1/pass2
    # always agree unless the policy or inventory scan regresses.
    disagreement_ids = {d["case_id"] for d in disagreements}
    sealed_rows: list[BusinessProposalGoldRow] = []
    for row in pass1_rows:
        if row.case_id in disagreement_ids:
            continue
        if row.case_id in validation_report["issues_by_case"]:
            continue
        sealed_rows.append(row)

    sealed_dicts = [r.to_dict() for r in sealed_rows]
    seal_manifest = protocol.seal_gold(sealed_dicts)

    case_audit = audit_mod.build_case_label_audit(
        old_eligibility_by_case=old_eligibility, gold_rows=sealed_dicts
    )

    projection = project_gold_into_benchmark(sealed_dicts)

    mode_counts: dict[str, int] = {}
    for r in sealed_dicts:
        mode_counts[r["node_evaluation_mode"]] = mode_counts.get(r["node_evaluation_mode"], 0) + 1

    summary = {
        "generated_at": NOW,
        "n_bp_cases_total": len(all_cases),
        "n_sealed_rows": len(sealed_rows),
        "n_excluded_disagreement": len(disagreement_ids),
        "n_excluded_validation_failed": len(
            [c for c in validation_report["issues_by_case"] if c not in disagreement_ids]
        ),
        "mode_counts": mode_counts,
        "agreement_metrics": agreement,
        "validation_status": validation_report["status"],
        "seal_manifest": seal_manifest,
        "new_fixtures": fixtures_written,
        "new_dev_cases": [s["case_id"] for s in NEW_DEV_CASES],
        "new_holdout_cases": [s["case_id"] for s in NEW_HOLDOUT_CASES],
        "projection": projection,
        "verdict": "READY_FOR_BUSINESS_PROPOSAL_BLIND_NODE_EVALUATION"
        if validation_report["status"] == "VALID" and not disagreement_ids
        else "REVIEW_REQUIRED_BEFORE_BLIND_EVALUATION",
    }

    gold_root = protocol.DEFAULT_GOLD_ROOT
    (gold_root / protocol.DISAGREEMENTS_FILENAME).write_text(
        json.dumps({"n_disagreements": len(disagreements), "disagreements": disagreements}, ensure_ascii=False, indent=2)
        + "\n",
        encoding="utf-8",
    )
    (gold_root / protocol.SUMMARY_FILENAME).write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (gold_root / protocol.AUDIT_FILENAME).write_text(
        json.dumps(case_audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (gold_root / "validation_report.json").write_text(
        json.dumps(validation_report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    return summary


if __name__ == "__main__":
    result = main()
    print(json.dumps(result, ensure_ascii=False, indent=2))
