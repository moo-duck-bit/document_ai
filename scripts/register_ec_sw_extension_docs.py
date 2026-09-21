"""Register EC-SW extension docs (MDTM/MDVP/MDDP/MDCP/MDMP) into data/examples/ec_sw/.

Copies from Desktop (or --src-dir) without modifying originals.
Writes/refreshes document_set_registry.json paths and light analysis.
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
from pathlib import Path

from docx import Document

PROJECT = Path(__file__).resolve().parents[1]
BASE = PROJECT / "data" / "examples" / "ec_sw"
DEFAULT_SRC = Path(os.environ.get("EC_SW_SRC_DIR") or Path.home() / "Desktop")

# short_id prefix, source filename on Desktop, template_id, role, priority
EXTENSION_FILES: list[tuple[str, str, str, str, int]] = [
    (
        "matrix_mdtm",
        "EC-SW-MDTM(XA) 추적성 매트릭스.docx",
        "matrix_traceability",
        "traceability_matrix",
        1,
    ),
    (
        "plan_mdvp",
        "EC-SW-MDVP(XA) 소프트웨어 검증 및 유효성 확인 계획서.docx",
        "plan_verification_validation",
        "vv_plan",
        2,
    ),
    (
        "plan_mddp",
        "EC-SW-MDDP(XA) 소프트웨어 개발 계획서.docx",
        "plan_development",
        "development_plan",
        3,
    ),
    (
        "process_mdcp",
        "EC-SW-MDCP(XA) 소프트웨어 형상관리 프로세스.docx",
        "process_configuration_management",
        "configuration_management",
        4,
    ),
    (
        "process_mdmp",
        "EC-SW-MDMP(XA) 소프트웨어 유지보수 프로세스.docx",
        "process_maintenance",
        "maintenance_process",
        5,
    ),
]


def _light_stats(path: Path) -> dict:
    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return {
        "paragraph_count": len(paras),
        "table_count": len(doc.tables),
        "size_kb": round(path.stat().st_size / 1024, 1),
        "cover_lines": paras[:12],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--src-dir",
        type=Path,
        default=DEFAULT_SRC,
        help="Directory containing original EC-SW-*.docx files",
    )
    args = parser.parse_args()
    src_dir: Path = args.src_dir
    BASE.mkdir(parents=True, exist_ok=True)

    extension_meta: list[dict] = []
    for short_id, filename, template_id, role, priority in EXTENSION_FILES:
        src = src_dir / filename
        if not src.exists():
            raise FileNotFoundError(f"Missing source: {src}")
        dest = BASE / f"{short_id}_{filename}"
        shutil.copy2(src, dest)
        stats = _light_stats(dest)
        rel = str(dest.relative_to(PROJECT)).replace("\\", "/")
        print(f"OK {rel} ({stats['size_kb']} KB, tables={stats['table_count']})")
        extension_meta.append(
            {
                "doc_code": filename.split()[0],
                "short_id": short_id,
                "template_id": template_id,
                "role": role,
                "status": "registered_extension",
                "agent_support": "planned",
                "priority": priority,
                "path": rel,
                "blank_template": None,
                **stats,
            }
        )

    registry_path = BASE / "document_set_registry.json"
    if registry_path.exists():
        registry = json.loads(registry_path.read_text(encoding="utf-8"))
    else:
        registry = {"document_set": "ec_sw", "documents": []}

    # Refresh extension entries by short_id; keep non-extension entries.
    keep = [
        d
        for d in registry.get("documents", [])
        if d.get("short_id")
        not in {m["short_id"] for m in extension_meta}
    ]
    # Merge stats into planned registry shape used by README
    for m in extension_meta:
        slim = {
            "doc_code": m["doc_code"],
            "short_id": m["short_id"],
            "template_id": m["template_id"],
            "role": m["role"],
            "status": m["status"],
            "agent_support": "planned_p0"
            if m["priority"] == 1
            else ("planned_p1" if m["priority"] == 2 else "planned_p2" if m["priority"] == 3 else "planned_p3"),
            "path": m["path"],
            "blank_template": None,
            "priority": m["priority"],
            "paragraph_count": m["paragraph_count"],
            "table_count": m["table_count"],
            "size_kb": m["size_kb"],
        }
        keep.append(slim)

    registry["documents"] = keep
    registry["expansion_order"] = [
        "matrix_traceability",
        "plan_verification_validation",
        "plan_development",
        "process_configuration_management",
        "process_maintenance",
    ]
    registry_path.write_text(
        json.dumps(registry, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {registry_path.relative_to(PROJECT).as_posix()}")


if __name__ == "__main__":
    main()
