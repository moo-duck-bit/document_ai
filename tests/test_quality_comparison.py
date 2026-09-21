"""Tests for gold vs generated comparison."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from document_ai.quality.comparison import compare_documents


def _simple_mdsr(path: Path, product: str, req_body: str) -> None:
    doc = Document()
    product_table = doc.add_table(rows=2, cols=2)
    product_table.rows[0].cells[0].text = "제품명"
    product_table.rows[0].cells[1].text = product
    product_table.rows[1].cells[0].text = "모델명"
    product_table.rows[1].cells[1].text = product
    doc.add_paragraph(f"Overview for {product}")
    req = doc.add_table(rows=2, cols=2)
    req.rows[0].cells[0].text = "Req. 1"
    req.rows[1].cells[0].text = "설명"
    req.rows[1].cells[1].text = req_body
    doc.save(str(path))


def test_compare_documents_metrics(tmp_path: Path):
    gold = tmp_path / "gold.docx"
    generated = tmp_path / "gen.docx"
    _simple_mdsr(gold, "IMS", "Gold requirement body for inventory.")
    _simple_mdsr(generated, "IMS", "Generated requirement body for inventory.")

    result = compare_documents(generated, gold, label="mdsr")
    assert result["field_similarity"] == 1.0
    assert result["requirement_coverage"] == 1.0
    assert result["paragraph_similarity"] > 0.3
