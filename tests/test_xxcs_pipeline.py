"""XXCS learning, synthesis, render, validation, and quality tests."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document

from document_ai.form_fill.security_tests import (
    ensure_security_tests_payload,
    synthesize_security_tests,
)
from document_ai.learn.extract_security_tests import extract_security_tests_docx
from document_ai.quality.analyzer import analyze_document
from document_ai.render.xxcs import fill_xxcs_report, verify_xxcs_completeness
from document_ai.validation.validation_report import score_document
from document_ai.validation.xxcs_compare import compare_xxcs


def _write_mini_xxcs(path: Path, *, req_id: str = "IA-01", body: str = "Auth test passed.") -> Path:
    doc = Document()
    doc.add_paragraph("1. Security Verification")
    doc.add_paragraph(f"{req_id}. Identification and Authentication")
    table = doc.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "시험결과"
    table.rows[0].cells[1].text = "적용"
    table.rows[0].cells[2].text = "만족"
    table.rows[0].cells[3].text = "비고"
    table.rows[1].cells[0].text = body
    table.rows[1].cells[1].text = "해당"
    table.rows[1].cells[2].text = "만족"
    table.rows[1].cells[3].text = ""
    doc.add_paragraph("UC-01. User Control")
    table2 = doc.add_table(rows=2, cols=4)
    table2.rows[0].cells[0].text = "시험결과"
    table2.rows[0].cells[1].text = "적용"
    table2.rows[0].cells[2].text = "만족"
    table2.rows[0].cells[3].text = "비고"
    table2.rows[1].cells[0].text = "Access control verified."
    table2.rows[1].cells[1].text = "해당"
    table2.rows[1].cells[2].text = "만족"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def test_extract_security_tests_ia_uc_si_dc_ra(tmp_path: Path):
    path = _write_mini_xxcs(tmp_path / "filled.docx")
    # add DC/RA headings
    doc = Document(str(path))
    doc.add_paragraph("DC-01. Data confidentiality")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "시험결과"
    t.rows[0].cells[1].text = "적용"
    t.rows[1].cells[0].text = "Encrypted at rest."
    t.rows[1].cells[1].text = "해당"
    doc.add_paragraph("RA-02. Risk assessment")
    t2 = doc.add_table(rows=2, cols=2)
    t2.rows[0].cells[0].text = "시험결과"
    t2.rows[0].cells[1].text = "적용"
    t2.rows[1].cells[0].text = "Risk accepted."
    t2.rows[1].cells[1].text = "해당"
    doc.save(str(path))

    data = extract_security_tests_docx(path)
    ids = {t["req_id"] for t in data["tests"]}
    assert "IA-01" in ids
    assert "UC-01" in ids
    assert "DC-01" in ids
    assert "RA-02" in ids


def test_synthesize_security_tests_from_traceability():
    intake = {"product_name": "HRS", "domain": "hospital_reservation", "facts": {"product_name": "HRS"}}
    requirements = {
        "traceability": [
            {"requirement": "IA-01", "linked_reqs": "Req. 2, Req. 3"},
            {"requirement": "UC-01", "linked_reqs": "Req. 4"},
        ],
        "requirements": [{"req_id": "Req. 2"}, {"req_id": "Req. 3"}, {"req_id": "Req. 4"}],
    }
    design = {"items": [{"req_id": "Req. 2"}, {"req_id": "Req. 4"}]}
    payload = synthesize_security_tests(
        intake=intake,
        requirements_payload=requirements,
        design_payload=design,
        seed=None,
    )
    ids = [t["req_id"] for t in payload["tests"]]
    assert "IA-01" in ids and "UC-01" in ids
    assert payload["tests"][0]["linked_reqs"]


def test_security_tests_link_designs_from_multiple_requirements_without_duplicates():
    requirements = {
        "traceability": [
            {
                "requirement": "IA 04",
                "linked_reqs": "Req. 3, Req. 4\nReq. 102; Req. 4",
                "title": "Credential management",
            }
        ]
    }
    design = {
        "items": [
            {"req_id": "Req. 3"},
            {"req_id": "Req. 4"},
            {"req_id": "Req. 102"},
        ]
    }
    payload = synthesize_security_tests(
        intake={"product_name": "JM COLLECTION"},
        requirements_payload=requirements,
        design_payload=design,
        seed=None,
    )
    ia04 = next(test for test in payload["tests"] if test["req_id"] == "IA-04")
    assert ia04["linked_req_ids"] == ["Req. 3", "Req. 4", "Req. 102"]
    assert ia04["linked_design_ids"] == ["Req. 3", "Req. 4", "Req. 102"]
    assert ia04["provenance"]["linked_design_rule"]


def test_fill_xxcs_and_validate(tmp_path: Path):
    gold = _write_mini_xxcs(tmp_path / "gold_xxcs.docx", body="Gold auth verification text.")
    template = _write_mini_xxcs(tmp_path / "template.docx", body="")  # empty result cell
    # clear result cells
    doc = Document(str(template))
    for table in doc.tables:
        if table.rows[0].cells[0].text.strip() == "시험결과":
            table.rows[1].cells[0].text = ""
            table.rows[1].cells[1].text = ""
            table.rows[1].cells[2].text = ""
    doc.save(str(template))

    payload = {
        "tests": [
            {
                "req_id": "IA-01",
                "test_method": "document_review",
                "test_procedure": "Verify authentication controls.",
                "expected_result": "Auth controls present.",
                "actual_result": "Gold auth verification text.",
                "test_result": "Gold auth verification text.",
                "applied": "해당",
                "satisfaction": "만족",
                "evidence": "review notes",
                "linked_reqs": "Req. 1",
                "linked_design_ids": "Req. 1",
            },
            {
                "req_id": "UC-01",
                "test_method": "document_review",
                "test_procedure": "Verify access control.",
                "expected_result": "Access control verified.",
                "actual_result": "Access control verified.",
                "test_result": "Access control verified.",
                "applied": "해당",
                "satisfaction": "만족",
                "evidence": "review notes",
                "linked_reqs": "Req. 2",
                "linked_design_ids": "Req. 2",
            },
        ]
    }
    out = tmp_path / "output_xxcs.docx"
    stats = fill_xxcs_report(template, {"product_name": "HRS"}, payload, out)
    assert out.exists()
    assert stats["security_tests"]["tables_filled"] >= 1

    cmp = compare_xxcs(
        out,
        gold,
        requirements_payload={
            "requirements": [{"req_id": "Req. 1"}, {"req_id": "Req. 2"}],
        },
        design_payload={"items": [{"req_id": "Req. 1"}, {"req_id": "Req. 2"}]},
    )
    assert cmp["security_coverage"] == 1.0
    assert cmp["test_item_completeness"] >= 0.4
    assert cmp["linked_requirement_coverage"] >= 0.9
    assert cmp["linked_design_coverage"] >= 0.5


def test_xxcs_render_extract_round_trip_status_satisfaction_and_design_links(tmp_path: Path):
    template = _write_mini_xxcs(tmp_path / "template.docx", body="")
    payload = {
        "tests": [
            {
                "req_id": "IA-01",
                "security_test_id": "IA-01-T01",
                "security_category": "IA",
                "title": "Authentication",
                "linked_req_ids": ["Req. 1", "Req. 2"],
                "linked_design_ids": ["Req. 1", "Req. 2"],
                "test_method": "document_review_and_static_analysis",
                "test_procedure": "Review authentication design and traceability.",
                "expected_result": "Authentication design is traceable.",
                "actual_result": "NOT_EXECUTED - no real execution evidence yet.",
                "satisfaction": "NOT_EXECUTED",
                "execution_status": "not_executed",
                "evidence": "NOT_COLLECTED",
            }
        ]
    }
    out = tmp_path / "roundtrip.docx"
    fill_xxcs_report(template, {"product_name": "HRS"}, payload, out)

    extracted = extract_security_tests_docx(out)["tests"]
    ia01 = next(test for test in extracted if test["req_id"] == "IA-01")
    assert ia01["linked_req_ids"] == ["Req. 1", "Req. 2"]
    assert ia01["linked_design_ids"] == ["Req. 1", "Req. 2"]
    assert "NOT_EXECUTED" in ia01["actual_result"]
    assert ia01["satisfaction"] == "NOT_EXECUTED"


def test_not_executed_counts_as_explicit_xxcs_completeness(tmp_path: Path):
    gold = _write_mini_xxcs(tmp_path / "gold.docx", body="Gold text")
    generated = _write_mini_xxcs(tmp_path / "generated.docx", body="")
    doc = Document(str(generated))
    row = doc.tables[0].rows[1]
    row.cells[0].text = (
        "시험방법: document_review\n"
        "절차: Review generated plan.\n"
        "예상결과: Traceability exists.\n"
        "실제결과: NOT_EXECUTED - no execution evidence.\n"
        "linked_reqs: Req. 1\n"
        "linked_design_ids: Req. 1\n"
        "evidence: NOT_COLLECTED"
    )
    row.cells[2].text = "NOT_EXECUTED"
    row2 = doc.tables[1].rows[1]
    row2.cells[0].text = (
        "시험방법: document_review\n"
        "절차: Review generated plan.\n"
        "예상결과: Traceability exists.\n"
        "실제결과: NOT_EXECUTED - no execution evidence.\n"
        "linked_reqs: Req. 2\n"
        "linked_design_ids: Req. 2\n"
        "evidence: NOT_COLLECTED"
    )
    row2.cells[2].text = "NOT_EXECUTED"
    doc.save(str(generated))

    cmp = compare_xxcs(
        generated,
        gold,
        requirements_payload={"requirements": [{"req_id": "Req. 1"}, {"req_id": "Req. 2"}]},
        design_payload={"items": [{"req_id": "Req. 1"}, {"req_id": "Req. 2"}]},
    )
    assert cmp["actual_result_completeness"] == 1.0
    assert cmp["satisfaction_completeness"] == 1.0
    assert cmp["evidence_completeness"] == 1.0


def test_xxcs_quality_analyzer(tmp_path: Path):
    path = _write_mini_xxcs(tmp_path / "out.docx")
    analysis = analyze_document(
        path,
        template="report_security_verification",
        domain="hospital_reservation",
        product_name="HRS",
        security_tests=[{"req_id": "IA-01"}, {"req_id": "UC-01"}],
    )
    assert analysis.structure_review.get("security_ids_found", 0) >= 1
    assert analysis.structure_review.get("tables_with_content", 0) >= 1


def test_ensure_security_tests_payload(tmp_path: Path):
    case = tmp_path / "case"
    case.mkdir()
    (case / "input.json").write_text(
        json.dumps(
            {
                "case_id": "mini",
                "product_name": "HRS",
                "domain": "hospital_reservation",
                "templates_in_set": ["report_security_verification"],
                "facts": {"product_name": "HRS"},
            }
        ),
        encoding="utf-8",
    )
    (case / "requirements.json").write_text(
        json.dumps(
            {
                "traceability": [{"requirement": "IA-01", "linked_reqs": "Req. 1"}],
                "requirements": [{"req_id": "Req. 1"}],
            }
        ),
        encoding="utf-8",
    )
    payload = ensure_security_tests_payload(case)
    assert (case / "security_tests.json").exists()
    assert payload["tests"]
    review = verify_xxcs_completeness(Document(), payload["tests"])
    assert review["security_ids_expected"] >= 1


def test_lab_ec_sw_xxcs_generation_validation_targets_are_met(tmp_path: Path):
    import hashlib

    from document_ai.form_fill.security_execution_review import build_report_security_payload

    case = Path("data/cases/lab_ec_sw")
    requirements = json.loads((case / "requirements.json").read_text(encoding="utf-8"))
    design = json.loads((case / "design_items.json").read_text(encoding="utf-8"))
    # Plan-mode payload: generation target test must not depend on selected overlays.
    security = build_report_security_payload(case, mode="plan")
    template = Path("data/templates/ec_sw/template_xxcs_skeleton.docx")
    out = tmp_path / "lab_xxcs.docx"
    fill_xxcs_report(template, {"product_name": "JM COLLECTION"}, security, out)

    cmp = compare_xxcs(
        out,
        out,
        requirements_payload=requirements,
        design_payload=design,
    )
    score = score_document(cmp)
    diag = {
        "score": score,
        "security_coverage": cmp.get("security_coverage"),
        "linked_requirement_coverage": cmp.get("linked_requirement_coverage"),
        "linked_design_coverage": cmp.get("linked_design_coverage"),
        "plan_test_completeness": cmp.get("plan_test_completeness"),
        "test_item_completeness": cmp.get("test_item_completeness"),
        "expected_result_completeness": cmp.get("expected_result_completeness"),
        "generated_test_count": cmp.get("generated_test_count"),
        "output": str(out),
        "output_sha256": hashlib.sha256(out.read_bytes()).hexdigest(),
        "cwd": str(Path.cwd()),
        "report_mode": security.get("report_mode"),
        "selected_count": (security.get("execution_overlay") or {}).get("selected_count"),
        "plan_test_count": len(security.get("tests") or []),
    }
    assert score >= 85.0, diag
    assert cmp["security_coverage"] == 1.0, diag
    assert cmp["linked_requirement_coverage"] == 1.0, diag
    assert cmp["linked_design_coverage"] >= 0.70, diag
    assert cmp["plan_test_completeness"] >= 0.85, diag
    assert cmp["test_item_completeness"] >= 0.85, diag
