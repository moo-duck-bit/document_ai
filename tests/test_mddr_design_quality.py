"""Tests for MDDR design-block extraction, matching, and force-generate quality."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from document_ai.learn.extract_design_items import (
    _design_text_from_item,
    extract_design_items_docx,
)
from document_ai.render.design_items import patch_mddr_design_items
from document_ai.validation.runner import _compare_design_blocks, validate_mddr


def _save(doc: Document, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def test_extract_paragraph_req_with_embedded_body(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("5. Design")
    doc.add_paragraph("Req. 1\nArchitecture uses web, api, and admin tiers.")
    doc.add_paragraph("Req. 2 Login flow stores JWT refresh tokens.")
    path = _save(doc, tmp_path / "para_mddr.docx")

    items = extract_design_items_docx(path)["items"]
    by_id = {item["req_id"]: item for item in items}
    assert set(by_id) == {"Req. 1", "Req. 2"}
    assert "Architecture uses web" in by_id["Req. 1"]["design_description"]
    assert _design_text_from_item(by_id["Req. 2"])
    assert by_id["Req. 1"]["block_kind"] == "paragraph"


def test_extract_section_prefixed_req_ids(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("4.2.1 Req. 1\nSection-prefixed design for overview.")
    doc.add_paragraph("5.2.3 Req. 102\nSecure session design for Req 102.")
    path = _save(doc, tmp_path / "section_mddr.docx")

    items = extract_design_items_docx(path)["items"]
    by_id = {item["req_id"]: item for item in items}
    assert "Req. 1" in by_id
    assert "Req. 102" in by_id
    assert by_id["Req. 1"].get("section_prefix") == "4.2.1"
    assert by_id["Req. 102"].get("section_prefix") == "5.2.3"


def test_extract_table_design_block(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Req. 7"
    table.rows[1].cells[0].text = "목적"
    table.rows[1].cells[1].text = "Protect inventory integrity"
    table.rows[2].cells[0].text = "설명"
    table.rows[2].cells[1].text = "WarehouseService locks stock rows."
    path = _save(doc, tmp_path / "table_mddr.docx")

    items = extract_design_items_docx(path)["items"]
    assert len(items) == 1
    assert items[0]["req_id"] == "Req. 7"
    assert items[0]["block_kind"] == "table"
    assert "WarehouseService" in _design_text_from_item(items[0])


def test_extract_run_split_req_id(tmp_path: Path):
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Req. ")
    para.add_run("9")
    para.add_run("\nDesign body from split runs.")
    path = _save(doc, tmp_path / "runs_mddr.docx")

    items = extract_design_items_docx(path)["items"]
    assert len(items) == 1
    assert items[0]["req_id"] == "Req. 9"
    assert "split runs" in items[0]["design_description"]


def test_match_ignores_section_number_difference(tmp_path: Path):
    gold = Document()
    gold.add_paragraph("4.1 Req. 3\nGold design body for requirement three.")
    gold_path = _save(gold, tmp_path / "gold.docx")

    gen = Document()
    gen.add_paragraph("9.9.9 Req. 3\nGold design body for requirement three.")
    gen_path = _save(gen, tmp_path / "gen.docx")

    result = _compare_design_blocks(gen_path, gold_path)
    assert result["design_block_coverage"] == 1.0
    assert result["matched_design_block_count"] == 1
    assert result["missing_in_generated"] == []
    assert result["design_block_text_similarity"] >= 0.99


def test_patch_does_not_duplicate_existing_req(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Req. 1\nExisting design body already present.")
    before = extract_design_items_docx(_save(doc, tmp_path / "before.docx"))["items"]
    assert len(before) == 1

    patched = patch_mddr_design_items(
        doc,
        [{"req_id": "Req. 1", "design_description": "Replacement design body text."}],
    )
    assert "Req. 1" in patched
    out = _save(doc, tmp_path / "after.docx")
    items = extract_design_items_docx(out)["items"]
    assert len(items) == 1
    assert items[0]["req_id"] == "Req. 1"
    assert "Replacement" in items[0]["design_description"]


@pytest.mark.skipif(
    not Path("data/cases/lab_ec_sw/output_mddr.docx").exists()
    or not Path("data/cases/jm_collection/output_mddr.docx").exists(),
    reason="lab_ec_sw / jm_collection MDDR outputs required",
)
def test_lab_ec_sw_mddr_validation_targets():
    gen = Path("data/cases/lab_ec_sw/output_mddr.docx")
    gold = Path("data/cases/jm_collection/output_mddr.docx")
    result = validate_mddr(gen, gold, domain="ecommerce_b2c")
    metrics = result["metrics"]
    assert metrics["design_block_coverage"] >= 0.50
    assert result["score"] >= 85.0


@pytest.mark.skipif(
    not Path("data/cases/lab_ec_sw/input.json").exists(),
    reason="lab_ec_sw case required",
)
def test_force_generate_quality_pass(tmp_path: Path):
    """force_generate must run on an isolated case copy — never mutate lab_ec_sw."""
    import hashlib
    import shutil

    from document_ai.eval.real_project_runner import run_project_e2e_validation

    src = Path("data/cases/lab_ec_sw")
    watched = [
        "output_mdsr.docx",
        "output_mddr.docx",
        "output_xxcs.docx",
        "security_tests.json",
        "design_items.json",
        "requirements.json",
    ]
    before = {
        name: hashlib.sha256((src / name).read_bytes()).hexdigest()
        for name in watched
        if (src / name).exists()
    }

    case_dir = tmp_path / "lab_ec_sw"
    shutil.copytree(
        src,
        case_dir,
        ignore=shutil.ignore_patterns(
            "executions",
            "execution_reviews",
            "__pycache__",
            "workdir",
        ),
    )
    result = run_project_e2e_validation(
        case_dir,
        force_generate=True,
    )
    assert result["harness"]["ok"] is True
    assert result["quality"]["scores"]["overall"] >= 85.0
    assert result["checks"]["quality_pass"] is True
    validation = result["validation"]["scores"]
    if validation.get("mddr") is not None:
        assert validation["mddr"] >= 85.0
    blocks = json.loads((case_dir / "validation_report.json").read_text(encoding="utf-8"))
    mddr_metrics = (blocks.get("mddr") or {}).get("metrics") or blocks.get("metrics") or {}
    coverage = mddr_metrics.get("design_block_coverage")
    if coverage is not None:
        assert coverage >= 0.50

    after = {
        name: hashlib.sha256((src / name).read_bytes()).hexdigest()
        for name in before
    }
    assert after == before, "lab_ec_sw source files must not change during force_generate test"
