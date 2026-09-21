# -*- coding: utf-8 -*-
"""PR-27 Generic Document Set E2E Workflow tests."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from document_ai.pilot_ui.app import create_app
from document_ai.workflow.catalog import get_catalog, list_document_sets
from document_ai.workflow.orchestrator import (
    approve_workflow,
    catalog_payload,
    create_workflow,
    get_workflow,
    list_workflows,
    run_analysis,
    run_result,
    run_writer,
)
from document_ai.workflow.schema import WORKFLOW_STATES

REPO = Path(__file__).resolve().parents[1]
MDTM = next((REPO / "data" / "examples" / "ec_sw").glob("matrix_mdtm*.docx"))
DESKTOP_MDTM = Path.home() / "Desktop" / "EC-SW-MDTM(XA) 추적성 매트릭스.docx"


def _mini_docx(tmp_path: Path, name: str = "report.docx", text: str = "방법론 데이터 결과") -> Path:
    path = tmp_path / name
    doc = Document()
    doc.add_heading("일반 보고서", level=1)
    doc.add_paragraph(text)
    doc.add_paragraph("결론과 제언")
    doc.save(str(path))
    return path


def test_catalog_three_sets():
    sets = {c["document_set_id"] for c in list_document_sets()}
    assert sets == {"ec_sw", "general_report", "business_proposal"}


def test_catalog_payload():
    p = catalog_payload()
    assert len(p["document_sets"]) == 3


def test_get_catalog_unknown():
    with pytest.raises(KeyError):
        get_catalog("nope")


def test_workflow_states_include_required():
    for s in (
        "UPLOADED",
        "ANALYZING",
        "REVIEW_READY",
        "WAITING_APPROVAL",
        "WRITING",
        "VALIDATING",
        "COMPLETED",
        "FAILED",
    ):
        assert s in WORKFLOW_STATES


def test_create_requires_change_request(tmp_path):
    with pytest.raises(ValueError):
        create_workflow(document_set="ec_sw", change_request="  ", root=tmp_path)


def test_create_ec_sw_upload(tmp_path):
    data = MDTM.read_bytes()
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101 변경",
        files=[("MDTM.docx", data, "traceability")],
        name="t",
        root=tmp_path,
    )
    assert rec["state"] == "UPLOADED"
    assert (tmp_path / rec["workflow_id"] / "input" / "MDTM.docx").exists()
    assert (tmp_path / rec["workflow_id"] / "workflow" / "workflow.json").exists()


def test_analyze_ec_sw_mdtm(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101 권한관리",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert out["state"] in {"WAITING_APPROVAL", "REVIEW_READY", "COMPLETED"}
    assert out["validation"]["ok"] is True
    assert any(c.get("status") == "PATCH_CANDIDATE" for c in out.get("patch_candidates") or []) or out.get(
        "patch_candidates"
    ) is not None


def test_analyze_general_report(tmp_path):
    docx = _mini_docx(tmp_path)
    rec = create_workflow(
        document_set="general_report",
        change_request="방법론 데이터 출처를 보강한다",
        files=[(docx.name, docx.read_bytes(), "general_report")],
        root=tmp_path,
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert out["state"] == "WAITING_APPROVAL"
    assert out["metadata"]["analysis"]["engine"] == "generic_template_sections"
    assert len(out["review_required"]) >= 1


def test_analyze_business_proposal(tmp_path):
    docx = _mini_docx(tmp_path, "proposal.docx", "실행 일정 예산 계획")
    rec = create_workflow(
        document_set="business_proposal",
        change_request="실행 일정과 예산을 수정",
        files=[(docx.name, docx.read_bytes(), None)],
        root=tmp_path,
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert out["validation"]["ok"] is True


def test_approve_all_pending(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 204",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    out = approve_workflow(rec["workflow_id"], approve_all_pending=True, root=tmp_path)
    assert all(a["decision"] == "APPROVED" for a in out["approvals"]) or not out["approvals"]


def test_item_level_reject(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
    )
    analyzed = run_analysis(rec["workflow_id"], root=tmp_path)
    if not analyzed["approvals"]:
        pytest.skip("no approvals seeded")
    item_id = analyzed["approvals"][0]["item_id"]
    out = approve_workflow(
        rec["workflow_id"],
        decisions=[{"item_id": item_id, "decision": "REJECTED"}],
        root=tmp_path,
    )
    hit = next(a for a in out["approvals"] if a["item_id"] == item_id)
    assert hit["decision"] == "REJECTED"


def test_document_level_block(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
    )
    analyzed = run_analysis(rec["workflow_id"], root=tmp_path)
    if not analyzed["approvals"]:
        pytest.skip("no approvals")
    doc_id = analyzed["approvals"][0]["document_id"]
    out = approve_workflow(
        rec["workflow_id"],
        document_decisions={doc_id: "BLOCKED"},
        root=tmp_path,
    )
    assert all(
        a["decision"] == "BLOCKED" for a in out["approvals"] if a["document_id"] == doc_id
    )


def test_writer_default_no_mutation(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    approve_workflow(rec["workflow_id"], approve_all_pending=True, root=tmp_path)
    before = hashlib.sha256(MDTM.read_bytes()).hexdigest()
    out = run_writer(rec["workflow_id"], enable_write=False, root=tmp_path)
    after = hashlib.sha256(MDTM.read_bytes()).hexdigest()
    assert before == after
    assert out["workflow"]["writer_result"]["controlled_writer_invoked"] is False
    assert out["workflow"]["state"] == "COMPLETED"


def test_result_rows_schema(tmp_path):
    rec = create_workflow(
        document_set="general_report",
        change_request="결과 핵심 발견",
        files=[("r.docx", _mini_docx(tmp_path).read_bytes(), None)],
        root=tmp_path,
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    approve_workflow(rec["workflow_id"], approve_all_pending=True, root=tmp_path)
    result = run_writer(rec["workflow_id"], root=tmp_path)
    rows = result["result"]["rows"]
    assert rows
    for key in (
        "document_id",
        "status",
        "review",
        "applied",
        "rejected",
        "blocked",
        "validation",
        "diff_available",
        "download",
    ):
        assert key in rows[0]


def test_artifacts_written(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 1",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    wid = rec["workflow_id"]
    run_analysis(wid, root=tmp_path)
    base = tmp_path / wid / "workflow"
    for name in (
        "workflow.json",
        "workflow_summary.json",
        "workflow_validation.json",
        "workflow_timeline.json",
    ):
        assert (base / name).is_file()


def test_timeline_events(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    got = get_workflow(rec["workflow_id"], root=tmp_path)
    events = [e["event"] for e in got["workflow"]["timeline"]]
    assert "created" in events
    assert "analysis_completed" in events


def test_list_workflows(tmp_path):
    create_workflow(
        document_set="ec_sw",
        change_request="Req. 1",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    rows = list_workflows(root=tmp_path)
    assert rows


def test_get_workflow(tmp_path):
    rec = create_workflow(
        document_set="business_proposal",
        change_request="일정",
        files=[("p.docx", _mini_docx(tmp_path, "p.docx").read_bytes(), None)],
        root=tmp_path,
    )
    got = get_workflow(rec["workflow_id"], root=tmp_path)
    assert got["workflow"]["document_set"] == "business_proposal"


def test_desktop_untouched_during_e2e(tmp_path):
    if not DESKTOP_MDTM.exists():
        return
    before = hashlib.sha256(DESKTOP_MDTM.read_bytes()).hexdigest()
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    run_writer(rec["workflow_id"], root=tmp_path)
    after = hashlib.sha256(DESKTOP_MDTM.read_bytes()).hexdigest()
    assert before == after


def test_examples_mdtm_untouched(tmp_path):
    before = hashlib.sha256(MDTM.read_bytes()).hexdigest()
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 108",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    approve_workflow(rec["workflow_id"], approve_all_pending=True, root=tmp_path)
    run_writer(rec["workflow_id"], enable_write=True, root=tmp_path)
    after = hashlib.sha256(MDTM.read_bytes()).hexdigest()
    assert before == after


def test_api_catalog_and_workflow_page():
    client = TestClient(create_app())
    assert client.get("/api/workflow/catalog").status_code == 200
    r = client.get("/workflow")
    assert r.status_code == 200
    assert "Generic Document Set Workflow" in r.text


def test_api_create_analyze(tmp_path, monkeypatch):
    from document_ai import workflow as wfmod

    monkeypatch.setattr(wfmod.orchestrator, "DEFAULT_ROOT", tmp_path)
    client = TestClient(create_app())
    files = [("files", ("MDTM.docx", MDTM.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))]
    r = client.post(
        "/api/workflow/create",
        data={
            "document_set": "ec_sw",
            "change_request": "Req. 101",
            "name": "api",
        },
        files=files,
    )
    assert r.status_code == 200, r.text
    wid = r.json()["workflow_id"]
    a = client.post(f"/api/workflow/{wid}/analyze")
    assert a.status_code == 200
    assert a.json()["state"] in WORKFLOW_STATES


def test_legacy_pilot_still_works():
    client = TestClient(create_app())
    assert client.get("/").status_code == 200
    assert client.get("/api/pilot/health").json()["ok"] is True


def test_run_result_idempotent(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 2",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    approve_workflow(rec["workflow_id"], approve_all_pending=True, root=tmp_path)
    run_writer(rec["workflow_id"], root=tmp_path)
    a = run_result(rec["workflow_id"], root=tmp_path)
    b = run_result(rec["workflow_id"], root=tmp_path)
    assert a["workflow"]["state"] == b["workflow"]["state"] == "COMPLETED"
    assert b.get("idempotent") is True
    completed = [e for e in a["workflow"]["timeline"] if e.get("event") == "completed"]
    completed_b = [e for e in b["workflow"]["timeline"] if e.get("event") == "completed"]
    assert len(completed) == 1
    assert len(completed_b) == 1


def test_failed_unknown_set(tmp_path):
    with pytest.raises(KeyError):
        create_workflow(document_set="xyz", change_request="x", root=tmp_path)


def test_copy_created_for_upload(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 1",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    copies = list((tmp_path / rec["workflow_id"] / "copies").glob("*.docx"))
    assert copies


def test_multi_file_upload(tmp_path):
    a = _mini_docx(tmp_path, "a.docx")
    b = _mini_docx(tmp_path, "b.docx")
    rec = create_workflow(
        document_set="general_report",
        change_request="방법론",
        files=[
            ("a.docx", a.read_bytes(), None),
            ("b.docx", b.read_bytes(), None),
        ],
        root=tmp_path,
    )
    assert len(rec["documents"]) == 2


def test_validation_ok_flag(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert "ok" in out["validation"]


def test_writer_gated_even_if_enable_true(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 101",
        files=[("MDTM.docx", MDTM.read_bytes(), None)],
        root=tmp_path,
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    approve_workflow(rec["workflow_id"], approve_all_pending=True, root=tmp_path)
    out = run_writer(rec["workflow_id"], enable_write=True, root=tmp_path)
    assert out["workflow"]["writer_result"]["applied"] == 0
    assert out["workflow"]["writer_result"]["controlled_writer_invoked"] is False


def test_index_has_workflow_link():
    client = TestClient(create_app())
    html = client.get("/").text
    assert "/workflow" in html
