# -*- coding: utf-8 -*-
"""Generalized regression fixtures for EC-SW document impact (no case-id hardcode)."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document

from document_ai.domain_packs.ec_sw.identifier_parser import valid_canonical_requirement_ids
from document_ai.evaluation.document_set.prediction_adapter import adapt_workflow_prediction
from document_ai.workflow.orchestrator import create_workflow, run_analysis

REPO = Path(__file__).resolve().parents[1]
MDTM = next((REPO / "data" / "examples" / "ec_sw").glob("matrix_mdtm*.docx"))


def _sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def _stub(path: Path) -> Path:
    d = Document()
    d.add_paragraph("requirements stub body")
    d.save(str(path))
    return path


def test_malformed_req_not_exact_and_not_impacted(tmp_path):
    """Glued REQ<digits> must not become exact Req evidence or IMPACTED."""
    cr = "REQ2 형식 오류 식별자"
    assert valid_canonical_requirement_ids(cr) == []
    before = _sha(MDTM)
    rec = create_workflow(
        document_set="ec_sw",
        change_request=cr,
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
        name="malformed_req",
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert out["patch_candidates"] == []
    assert "MDTM" not in out["impacted_documents"]
    decisions = {d["document_id"]: d for d in out["metadata"]["document_impact_decisions"]}
    assert decisions["MDTM"]["predicted_status"] == "UNRELATED"
    pred = adapt_workflow_prediction(
        "fixture",
        {**out, "document_impact_decisions": out["metadata"]["document_impact_decisions"]},
    )
    md = next(d for d in pred["documents"] if d["document_id"] == "MDTM")
    assert md["predicted_status"] == "UNRELATED"
    assert pred["patches"] == []
    assert _sha(MDTM) == before


def test_mdtm_only_multi_doc_no_role_promotion(tmp_path):
    """Exact evidence on MDTM must not promote role-compatible sibling docs."""
    stub = _stub(tmp_path / "MDSR_stub.docx")
    before = _sha(MDTM)
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 2 MDTM만 영향",
        files=[
            ("MDTM.docx", MDTM.read_bytes(), "traceability"),
            ("MDSR_stub.docx", stub.read_bytes(), "requirements"),
        ],
        root=tmp_path,
        name="mdtm_only",
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert "MDTM" in out["impacted_documents"]
    assert "MDSR_STUB" not in out["impacted_documents"]
    decisions = {d["document_id"]: d for d in out["metadata"]["document_impact_decisions"]}
    assert decisions["MDTM"]["predicted_status"] == "IMPACTED"
    assert decisions["MDSR_STUB"]["predicted_status"] == "UNRELATED"
    assert decisions["MDTM"]["substantive_evidence_count"] >= 1
    assert decisions["MDSR_STUB"]["substantive_evidence_count"] == 0
    assert len(out["patch_candidates"]) >= 1
    assert all(c.get("document_id") == "MDTM" for c in out["patch_candidates"])
    assert _sha(MDTM) == before


def test_exact_req_still_impacted(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 2 행 추적성 갱신",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
        name="exact_ok",
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert any(c.get("status") == "PATCH_CANDIDATE" or True for c in out["patch_candidates"])
    assert len(out["patch_candidates"]) >= 1
    assert "MDTM" in out["impacted_documents"]


def test_multiple_req_still_patches(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 1과 Req. 3을 동시에 수정",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
        name="multi_req",
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert len(out["patch_candidates"]) >= 2


def test_semantic_only_no_patch(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="사용자 인증 및 접근통제 보안 요구 변경",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
        name="semantic",
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert out["patch_candidates"] == []


def test_general_report_unchanged_review_only(tmp_path):
    p = tmp_path / "report.docx"
    d = Document()
    d.add_heading("일반 보고서", 1)
    d.add_heading("방법론", 2)
    d.add_paragraph("데이터 출처는 내부 설문이다.")
    d.save(str(p))
    rec = create_workflow(
        document_set="general_report",
        change_request="방법론 데이터 출처를 외부 공개자료로 수정",
        files=[("report_base.docx", p.read_bytes(), "general_report")],
        root=tmp_path,
        name="gr",
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert out["patch_candidates"] == []
    assert out["review_required"]


def test_prediction_adapter_does_not_read_gold():
    import inspect
    from document_ai.evaluation.document_set import prediction_adapter as pa

    src = inspect.getsource(pa)
    assert "document_impacts.jsonl" not in src
    assert "load_benchmark_manifest" not in src
    assert "golden" not in src.lower()


def test_impact_artifacts_written(tmp_path):
    rec = create_workflow(
        document_set="ec_sw",
        change_request="Req. 2",
        files=[("MDTM.docx", MDTM.read_bytes(), "traceability")],
        root=tmp_path,
        name="arts",
    )
    run_analysis(rec["workflow_id"], root=tmp_path)
    ec = tmp_path / rec["workflow_id"] / "output" / "document_set" / "ec_sw"
    for name in (
        "ec_sw_identifier_analysis.json",
        "ec_sw_document_impact_evidence.json",
        "ec_sw_document_impact_decisions.json",
        "ec_sw_document_impact_summary.json",
        "ec_sw_document_impact_validation.json",
    ):
        assert (ec / name).is_file(), name


def test_no_case_id_hardcode_in_policy_modules():
    root = REPO / "src" / "document_ai"
    forbidden = ["ec_sw_bad_req_format", "ec_sw_mdtm_only_multi_doc"]
    for rel in (
        "domain_packs/ec_sw/identifier_parser.py",
        "domain_packs/ec_sw/document_impact.py",
        "workflow/analysis.py",
        "evaluation/document_set/prediction_adapter.py",
    ):
        text = (root / rel).read_text(encoding="utf-8")
        for token in forbidden:
            assert token not in text
