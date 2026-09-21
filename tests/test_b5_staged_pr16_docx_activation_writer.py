# -*- coding: utf-8 -*-
"""PR-16: Feature-Flag DOCX Activation Writer tests."""

from __future__ import annotations

import copy
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from document_ai.impact.docx_activation_writer import (
    DocxActivationPlanItem,
    build_docx_activation_plan,
    is_docx_activation_enabled,
    locate_targets,
    replace_text_preserving_runs,
    resolve_execution_policy,
    run_docx_activation_writer,
    sha256_file,
    validate_activated_docx,
    validate_docx_activation_plan,
)


def _gate(
    pid: str,
    *,
    status: str = "PASS",
    eligible: bool = True,
    decision: str = "AUTO_APPLY",
    applied: bool = True,
    patch_status: str = "VALID",
    lang_status: str = "VALID",
    acu: str = "ACU-001",
    rid: str = "Req. 1",
    doc: str = "MDSR",
    field: str = "description",
) -> dict:
    return {
        "gate_result_id": f"PG-{pid}",
        "patch_id": pid,
        "draft_id": f"SD-{pid}",
        "atomic_change_id": acu,
        "requirement_id": rid,
        "document": doc,
        "field": field,
        "patch_validation_status": patch_status,
        "language_validation_status": lang_status,
        "activation_decision": decision,
        "preview_applied": applied,
        "preview_validation_status": "VALID",
        "final_status": status,
        "eligible_for_docx_activation": eligible,
        "actual_docx_changed": False,
        "actual_generation_changed": False,
        "reason_codes": [],
    }


def _preview(
    pid: str,
    before: str,
    after: str,
    *,
    rid: str = "Req. 1",
    op: str = "UPDATE",
) -> dict:
    return {
        "patch_id": pid,
        "requirement_id": rid,
        "document": "MDSR",
        "field": "description",
        "operation": op,
        "original_requirement": before,
        "proposed_requirement": after,
        "preview_requirement": after,
        "applied_in_preview": True,
        "activation_decision": "AUTO_APPLY",
    }


def _patch(pid: str, *, op: str = "UPDATE", before: str = "", after: str = "") -> dict:
    return {
        "patch_id": pid,
        "operation": op,
        "original_requirement": before,
        "patched_requirement": after,
        "requirement_id": "Req. 1",
        "document": "MDSR",
        "field": "description",
    }


def _make_fixture_docx(path: Path, *, before: str, extra: str = "UNRELATED KEEP") -> Path:
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "HEADER_STABLE"
    section.footer.paragraphs[0].text = "FOOTER_STABLE"
    doc.add_paragraph(extra)
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Req. 1"
    table.rows[0].cells[1].text = "Title"
    table.rows[1].cells[0].text = "설명"
    cell = table.rows[1].cells[1]
    cell.paragraphs[0].clear()
    run = cell.paragraphs[0].add_run(before)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
    table.rows[2].cells[0].text = "목적"
    table.rows[2].cells[1].text = "PURPOSE_STABLE"
    doc.add_paragraph("TRAILING_STABLE")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


@pytest.fixture
def fixture_docx(tmp_path: Path) -> Path:
    before = "시스템은 계정을 잠가야 한다."
    return _make_fixture_docx(tmp_path / "source_mdsr.docx", before=before)


def _run_writer(
    tmp_path: Path,
    fixture: Path,
    *,
    gates: list[dict],
    previews: list[dict],
    patches: list[dict],
    flag: bool,
    policy: str,
    global_status: str,
):
    return run_docx_activation_writer(
        gate_payload={"results": gates, "global_status": global_status},
        preview_entries=previews,
        patches=patches,
        source_docx_by_document={"MDSR": fixture},
        output_dir=tmp_path / "activated",
        feature_flag_enabled=flag,
        execution_policy=policy,
    )


def test_feature_flag_default_false():
    assert is_docx_activation_enabled(env={}) is False
    assert is_docx_activation_enabled(env={"DOCX_ACTIVATION_ENABLED": "false"}) is False
    assert is_docx_activation_enabled(env={"DOCX_ACTIVATION_ENABLED": "true"}) is True


def test_policy_default_strict():
    assert resolve_execution_policy(env={}) == "strict_global"
    assert resolve_execution_policy(env={"DOCX_ACTIVATION_POLICY": "eligible_only"}) == (
        "eligible_only"
    )


def test_1_2_3_flag_off_no_write(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    h0 = sha256_file(fixture_docx)
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=False,
        policy="strict_global",
        global_status="PASS",
    )
    assert payload["summary"]["attempted_count"] == 0
    assert payload["summary"]["succeeded_count"] == 0
    assert payload["results"]["output_files"] == []
    assert sha256_file(fixture_docx) == h0
    assert payload["summary"]["actual_docx_changed"] is False


def test_4_strict_global_block_stops_write(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[
            _gate("RP-1"),
            _gate(
                "RP-B",
                status="BLOCK",
                eligible=False,
                decision="BLOCK",
                applied=False,
            ),
        ],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="BLOCK",
    )
    assert payload["plan"]["write_allowed"] is False
    assert payload["summary"]["succeeded_count"] == 0
    assert payload["results"]["output_files"] == []


def test_5_strict_global_pass_plans_pass_items(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    planned = [i for i in payload["plan"]["items"] if i["planned"]]
    assert len(planned) == 1
    assert payload["summary"]["succeeded_count"] == 1
    assert payload["results"]["output_files"]


def test_6_eligible_only_with_global_block(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[
            _gate("RP-1"),
            _gate(
                "RP-B",
                status="BLOCK",
                eligible=False,
                decision="BLOCK",
                applied=False,
                acu="ACU-B",
            ),
        ],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="eligible_only",
        global_status="BLOCK",
    )
    planned = [i for i in payload["plan"]["items"] if i["planned"]]
    assert len(planned) == 1
    assert planned[0]["patch_id"] == "RP-1"
    blocked = [i for i in payload["plan"]["items"] if i["patch_id"] == "RP-B"][0]
    assert blocked["planned"] is False
    assert payload["summary"]["succeeded_count"] == 1


def test_7_8_9_10_11_eligibility_skips(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "AFTER"
    cases = [
        (_gate("RP-E", eligible=False), "ITEM_NOT_ELIGIBLE"),
        (
            _gate(
                "RP-R",
                status="REVIEW",
                eligible=False,
                decision="REVIEW",
                applied=False,
            ),
            "ITEM_NOT_PASS",
        ),
        (
            _gate(
                "RP-B",
                status="BLOCK",
                eligible=False,
                decision="BLOCK",
                applied=False,
            ),
            "ITEM_NOT_PASS",
        ),
        (
            _gate(
                "RP-U",
                status="UNKNOWN",
                eligible=False,
                decision="UNKNOWN",
                applied=False,
            ),
            "UNKNOWN_STATE",
        ),
    ]
    for g, code in cases:
        plan = build_docx_activation_plan(
            gate_results=[g],
            preview_entries=[_preview(g["patch_id"], before, after)],
            patches=[_patch(g["patch_id"], before=before, after=after)],
            source_docx_by_document={"MDSR": fixture_docx},
            output_dir=tmp_path / "activated",
            global_status="PASS",
            feature_flag_enabled=True,
            execution_policy="strict_global",
        )
        item = plan["items"][0]
        assert item.planned is False
        assert code in item.skip_reason_codes


def test_12_13_auto_apply_and_preview_required(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "AFTER"
    cases = [
        (
            _gate("RP-1", decision="REVIEW", applied=True, status="PASS", eligible=True),
            "ACTIVATION_NOT_AUTO_APPLY",
        ),
        (_gate("RP-2", applied=False), "PREVIEW_NOT_APPLIED"),
    ]
    for g, code in cases:
        plan = build_docx_activation_plan(
            gate_results=[g],
            preview_entries=[_preview(g["patch_id"], before, after)],
            patches=[_patch(g["patch_id"], before=before, after=after)],
            source_docx_by_document={"MDSR": fixture_docx},
            output_dir=tmp_path / f"act_{g['patch_id']}",
            global_status="PASS",
            feature_flag_enabled=True,
            execution_policy="strict_global",
        )
        assert plan["items"][0].planned is False
        assert code in plan["items"][0].skip_reason_codes


def test_14_15_supported_and_unsupported_ops(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    ok = _run_writer(
        tmp_path / "ok",
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after, op="UPDATE")],
        patches=[_patch("RP-1", op="UPDATE", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    assert ok["summary"]["succeeded_count"] == 1

    skip = build_docx_activation_plan(
        gate_results=[_gate("RP-2")],
        preview_entries=[_preview("RP-2", before, after, op="ADD")],
        patches=[_patch("RP-2", op="ADD", before=before, after=after)],
        source_docx_by_document={"MDSR": fixture_docx},
        output_dir=tmp_path / "skip",
        global_status="PASS",
        feature_flag_enabled=True,
        execution_policy="strict_global",
    )
    assert skip["items"][0].planned is False
    assert "UNSUPPORTED_DOCX_OPERATION" in skip["items"][0].skip_reason_codes


def test_16_17_18_locator_rules(fixture_docx: Path, tmp_path: Path):
    doc = Document(str(fixture_docx))
    m0 = locate_targets(
        doc, requirement_id="Req. 1", field="description", before_text="NO_SUCH_TEXT"
    )
    assert len(m0) == 0
    m1 = locate_targets(
        doc,
        requirement_id="Req. 1",
        field="description",
        before_text="시스템은 계정을 잠가야 한다.",
    )
    assert len(m1) == 1
    p = tmp_path / "dup.docx"
    d2 = Document()
    t = d2.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Req. 1"
    t.rows[1].cells[0].text = "설명"
    t.rows[1].cells[1].paragraphs[0].add_run("SAME")
    t.rows[1].cells[1].add_paragraph("SAME")
    d2.save(str(p))
    d2r = Document(str(p))
    m2 = locate_targets(d2r, requirement_id="Req. 1", field="description", before_text="SAME")
    assert len(m2) >= 2


def test_19_20_21_source_immutable_output_separate(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    h0 = sha256_file(fixture_docx)
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    assert sha256_file(fixture_docx) == h0
    outs = payload["results"]["output_files"]
    assert outs
    assert Path(outs[0]).resolve() != fixture_docx.resolve()
    assert "__activated_pr16" in Path(outs[0]).name


def test_22_23_24_25_output_content(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    outp = Path(payload["results"]["output_files"][0])
    doc = Document(str(outp))
    table_text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert after in table_text
    texts = [p.text for p in doc.paragraphs] + [
        c.text for t in doc.tables for r in t.rows for c in r.cells
    ]
    assert any("UNRELATED KEEP" in t for t in texts)
    assert "PURPOSE_STABLE" in table_text
    assert "HEADER_STABLE" in doc.sections[0].header.paragraphs[0].text
    assert "FOOTER_STABLE" in doc.sections[0].footer.paragraphs[0].text
    assert "BLOCK_SHOULD_NOT_APPEAR" not in table_text


def test_26_formatting_preservation(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    doc = Document(payload["results"]["output_files"][0])
    cell = doc.tables[0].rows[1].cells[1]
    run = cell.paragraphs[0].runs[0]
    assert run.bold is True
    assert after in run.text


def test_27_28_table_header_footer(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "NEW TEXT FOR LOCKOUT"
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    doc = Document(payload["results"]["output_files"][0])
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 3
    assert doc.sections[0].header.paragraphs[0].text == "HEADER_STABLE"
    assert doc.sections[0].footer.paragraphs[0].text == "FOOTER_STABLE"


def test_30_31_write_failure_keeps_source(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "AFTER"
    h0 = sha256_file(fixture_docx)
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", "WRONG_BEFORE", after)],
        patches=[_patch("RP-1", before="WRONG_BEFORE", after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    assert sha256_file(fixture_docx) == h0
    assert payload["summary"]["succeeded_count"] == 0


def test_32_no_overwrite_existing_output(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "AFTER_ONE"
    out_dir = tmp_path / "activated"
    p1 = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    first = Path(p1["results"]["output_files"][0])
    assert first.exists()
    plan2 = build_docx_activation_plan(
        gate_results=[_gate("RP-1")],
        preview_entries=[_preview("RP-1", before, "AFTER_TWO")],
        patches=[_patch("RP-1", before=before, after="AFTER_TWO")],
        source_docx_by_document={"MDSR": fixture_docx},
        output_dir=out_dir,
        global_status="PASS",
        feature_flag_enabled=True,
        execution_policy="strict_global",
    )
    assert Path(plan2["items"][0].output_docx_path) != first
    assert "__activated_pr16_" in Path(plan2["items"][0].output_docx_path).name


def test_33_34_deterministic_plan(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "AFTER"
    gates = [_gate("RP-2", acu="ACU-002"), _gate("RP-1", acu="ACU-001")]
    previews = [_preview("RP-1", before, after), _preview("RP-2", before, after)]
    patches = [
        _patch("RP-1", before=before, after=after),
        _patch("RP-2", before=before, after=after),
    ]
    a = build_docx_activation_plan(
        gate_results=gates,
        preview_entries=previews,
        patches=patches,
        source_docx_by_document={"MDSR": fixture_docx},
        output_dir=tmp_path / "a",
        global_status="PASS",
        feature_flag_enabled=False,
        execution_policy="strict_global",
    )
    b = build_docx_activation_plan(
        gate_results=gates,
        preview_entries=previews,
        patches=patches,
        source_docx_by_document={"MDSR": fixture_docx},
        output_dir=tmp_path / "b",
        global_status="PASS",
        feature_flag_enabled=False,
        execution_policy="strict_global",
    )
    assert [i.patch_id for i in a["items"]] == [i.patch_id for i in b["items"]]
    assert [i.patch_id for i in a["items"]] == ["RP-1", "RP-2"]


def test_35_input_non_mutation(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "AFTER"
    gates = [_gate("RP-1")]
    previews = [_preview("RP-1", before, after)]
    patches = [_patch("RP-1", before=before, after=after)]
    snap = (copy.deepcopy(gates), copy.deepcopy(previews), copy.deepcopy(patches))
    _run_writer(
        tmp_path,
        fixture_docx,
        gates=gates,
        previews=previews,
        patches=patches,
        flag=False,
        policy="strict_global",
        global_status="PASS",
    )
    assert gates == snap[0]
    assert previews == snap[1]
    assert patches == snap[2]


def test_36_preview_vs_output(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "시스템은 계정을 잠그고 알림을 제공해야 한다."
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    vs = payload["vs_preview"]
    assert vs["pairs"][0]["matched"] is True
    assert vs["pairs"][0]["applied"] is True


def test_37_validation_source_hash_change_detected():
    results = {
        "feature_flag_enabled": False,
        "items": [],
        "output_files": [],
        "source_hashes_before": {"MDSR": "aaa"},
        "source_hashes_after": {"MDSR": "bbb"},
        "actual_docx_changed": False,
        "actual_generation_changed": False,
    }
    v = validate_activated_docx(plan={}, results=results)
    assert any("source_hash_changed" in i for i in v["issues"])


def test_38_validation_illegal_block_write(fixture_docx: Path, tmp_path: Path):
    bad = DocxActivationPlanItem(
        activation_item_id="DA-X",
        gate_result_id="PG-X",
        patch_id="RP-X",
        draft_id="",
        atomic_change_id="",
        requirement_id="Req. 1",
        document="MDSR",
        field="description",
        source_docx_path=str(fixture_docx),
        output_docx_path=str(tmp_path / "out.docx"),
        gate_final_status="BLOCK",
        eligible_for_docx_activation=False,
        activation_decision="BLOCK",
        preview_applied=False,
        operation="UPDATE",
        before_text="a",
        after_text="b",
        write_succeeded=True,
    )
    v = validate_activated_docx(
        plan={},
        results={
            "feature_flag_enabled": True,
            "items": [bad],
            "output_files": [],
            "source_hashes_before": {},
            "source_hashes_after": {},
            "actual_docx_changed": False,
            "actual_generation_changed": False,
        },
    )
    assert any("review_block_applied" in i or "illegal_non_pass" in i for i in v["issues"])


def test_39_unrelated_preserved(fixture_docx: Path, tmp_path: Path):
    before = "시스템은 계정을 잠가야 한다."
    after = "NEW"
    payload = _run_writer(
        tmp_path,
        fixture_docx,
        gates=[_gate("RP-1")],
        previews=[_preview("RP-1", before, after)],
        patches=[_patch("RP-1", before=before, after=after)],
        flag=True,
        policy="strict_global",
        global_status="PASS",
    )
    doc = Document(payload["results"]["output_files"][0])
    texts = [p.text for p in doc.paragraphs] + [
        c.text for t in doc.tables for r in t.rows for c in r.cells
    ]
    assert any("UNRELATED KEEP" in t for t in texts)
    assert any("TRAILING_STABLE" in t for t in texts)
    assert any("PURPOSE_STABLE" in t for t in texts)


def test_plan_validation_rejects_non_pass_planned(fixture_docx: Path, tmp_path: Path):
    item = DocxActivationPlanItem(
        activation_item_id="DA-1",
        gate_result_id="PG-1",
        patch_id="RP-1",
        draft_id="",
        atomic_change_id="",
        requirement_id="Req. 1",
        document="MDSR",
        field="description",
        source_docx_path=str(fixture_docx),
        output_docx_path=str(tmp_path / "x.docx"),
        gate_final_status="REVIEW",
        eligible_for_docx_activation=False,
        activation_decision="REVIEW",
        preview_applied=False,
        operation="UPDATE",
        before_text="a",
        after_text="b",
        planned=True,
    )
    v = validate_docx_activation_plan({"items": [item]})
    assert v["status"] == "INVALID"


def test_run_aware_replace_unsafe_multi_run():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("AA")
    p.add_run("BB")
    ok, reason = replace_text_preserving_runs(p, "AABB", "CC")
    assert ok is False
    assert reason == "FORMAT_PRESERVATION_UNSAFE"
