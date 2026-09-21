"""Tests for real document validation (gold vs generated)."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from document_ai.validation.docx_compare import (
    paragraph_similarity,
    requirement_table_counts,
    traceability_table_counts,
)
from document_ai.validation.runner import run_document_validation
from document_ai.validation.section_compare import compare_sections
from document_ai.validation.traceability_compare import compare_traceability


def _write_mdsr(path: Path, *, product: str, req_body: str, with_trace: bool = True) -> None:
    doc = Document()
    doc.add_paragraph("1. Introduction")
    doc.add_paragraph("2. References")
    doc.add_paragraph("5. Functional Requirements")

    product_table = doc.add_table(rows=2, cols=2)
    product_table.rows[0].cells[0].text = "제품명"
    product_table.rows[0].cells[1].text = product
    product_table.rows[1].cells[0].text = "모델명"
    product_table.rows[1].cells[1].text = product

    doc.add_paragraph(f"Overview paragraph for {product} system.")

    req = doc.add_table(rows=4, cols=2)
    req.rows[0].cells[0].text = "Req. 1"
    req.rows[1].cells[0].text = "설명"
    req.rows[1].cells[1].text = req_body
    req.rows[2].cells[0].text = "목적"
    req.rows[2].cells[1].text = f"Purpose for {product}"
    req.rows[3].cells[0].text = "기준"
    req.rows[3].cells[1].text = "Acceptance criteria for requirement one."

    if with_trace:
        trace = doc.add_table(rows=2, cols=4)
        trace.rows[0].cells[0].text = "IA-01"
        trace.rows[0].cells[1].text = "Authentication"
        trace.rows[0].cells[2].text = "해당"
        trace.rows[0].cells[3].text = "Req. 1"
        trace.rows[1].cells[0].text = "UC-01"
        trace.rows[1].cells[1].text = "User control"
        trace.rows[1].cells[2].text = "해당"
        trace.rows[1].cells[3].text = "Req. 1"

    doc.save(str(path))


def _write_mddr(path: Path, *, product: str, design_body: str) -> None:
    doc = Document()
    doc.add_paragraph("1. Introduction")
    doc.add_paragraph("5. Design Description")
    doc.add_paragraph("Req. 1 Product overview design")
    doc.add_paragraph(design_body)
    doc.add_paragraph(f"Figure 1. {product} architecture")
    doc.save(str(path))


@pytest.fixture()
def validation_case(tmp_path: Path) -> Path:
    case_dir = tmp_path / "validate_case"
    case_dir.mkdir()
    (case_dir / "input.json").write_text(
        json.dumps(
            {
                "case_id": "validate_case",
                "domain": "inventory_b2b",
                "facts": {"product_name": "IMS"},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    gold_mdsr = case_dir / "gold_mdsr.docx"
    gold_mddr = case_dir / "gold_mddr.docx"
    gen_mdsr = case_dir / "output_mdsr.docx"
    gen_mddr = case_dir / "output_mddr.docx"

    body = "Inventory management requirement body with warehouse and stock control."
    design = "InventoryService handles stock reservation and warehouse lookup."
    _write_mdsr(gold_mdsr, product="IMS", req_body=body)
    _write_mdsr(gen_mdsr, product="IMS", req_body=body)
    _write_mddr(gold_mddr, product="IMS", design_body=design)
    _write_mddr(gen_mddr, product="IMS", design_body=design)
    return case_dir


def test_identical_documents_score_near_100(validation_case: Path):
    result = run_document_validation(validation_case)
    assert result["scores"]["overall"] >= 95
    assert result["scores"]["mdsr"] >= 95
    assert result["scores"]["mddr"] >= 95
    assert result["status"] == "PASS"


def test_section_coverage_calculation(validation_case: Path):
    sections = compare_sections(
        validation_case / "output_mdsr.docx",
        validation_case / "gold_mdsr.docx",
    )
    assert sections["section_coverage"] == 1.0
    assert sections["matched_section_count"] >= 2


def test_requirement_coverage_calculation(validation_case: Path):
    filled, total = requirement_table_counts(validation_case / "output_mdsr.docx")
    assert total == 1
    assert filled == 1
    assert paragraph_similarity(
        ["Overview paragraph for IMS system."],
        ["Overview paragraph for IMS system."],
    ) == 1.0


def test_traceability_coverage_calculation(validation_case: Path):
    filled, total = traceability_table_counts(validation_case / "output_mdsr.docx")
    assert total == 2
    assert filled == 2
    trace = compare_traceability(
        validation_case / "output_mdsr.docx",
        validation_case / "gold_mdsr.docx",
    )
    assert trace["traceability_coverage"] == 1.0
    assert trace["traceability_row_match"] == 1.0


def test_validation_report_files_created(validation_case: Path):
    result = run_document_validation(validation_case)
    md_path = Path(result["report_paths"]["markdown"])
    json_path = Path(result["report_paths"]["json"])
    assert md_path.exists()
    assert json_path.exists()
    md_text = md_path.read_text(encoding="utf-8")
    assert "# Document Validation Report" in md_text
    assert "Human Review Checklist" in md_text
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert payload["scores"]["overall"] >= 95


def test_cli_document_validate_subprocess(validation_case: Path):
    import subprocess
    import sys

    proc = subprocess.run(
        [
            sys.executable,
            "-m",
            "document_ai.cli",
            "document-validate",
            "--case",
            str(validation_case),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert proc.returncode == 0, proc.stderr
    assert "PASS" in proc.stdout
