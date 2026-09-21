# -*- coding: utf-8 -*-
"""PR-20: Document Structure Mapping Engine tests."""

from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path

from document_ai.document_parser.docx_parser import parse_docx_file
from document_ai.document_parser.locator_candidates import build_locator_candidates
from document_ai.document_parser.markdown_parser import parse_markdown_text
from document_ai.document_parser.normalization import (
    build_section_tree,
    document_from_object,
    normalize_document,
)
from document_ai.document_parser.parser import (
    SAMPLE_MARKDOWN_PROPOSAL,
    SAMPLE_MARKDOWN_REPORT,
    map_document_structure,
    run_document_structure_mapping_engine,
)
from document_ai.document_parser.structure import DocumentModel, SectionModel
from document_ai.document_parser.validation import validate_document_structure
from document_ai.impact.docx_activation_writer import is_docx_activation_enabled
from document_ai.template.generic_mapping import run_generic_document_template_pack
from document_ai.template.inventory import run_template_abstraction_layer

REPO = Path(__file__).resolve().parents[1]
FROZEN_PATHS = [
    REPO / "data/user_scenarios/scenario-001",
    REPO / "data/trials/trial-001-mindrium-xa",
    REPO / "data/trials/trial-002-lockout-multireq",
]


def _make_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("DOCX 샘플 보고서", level=1)
    doc.add_heading("요약", level=2)
    doc.add_paragraph("요약 본문입니다.")
    doc.add_heading("방법론", level=2)
    doc.add_heading("데이터 출처", level=3)
    doc.add_paragraph("내부 fixture")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "출처"
    table.rows[0].cells[1].text = "설명"
    table.rows[1].cells[0].text = "fixture"
    table.rows[1].cells[1].text = "샘플"
    doc.add_heading("결과", level=2)
    doc.add_paragraph("항목 하나", style="List Bullet")
    doc.add_paragraph("항목 둘", style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def test_01_markdown_parsing():
    doc = parse_markdown_text(
        SAMPLE_MARKDOWN_REPORT,
        document_id="md1",
        document_type="general_report",
    )
    assert doc.source_format == "markdown"
    assert doc.title == "연구 진행 보고서"
    assert any(s.heading == "방법론" for s in doc.sections)


def test_02_docx_parsing(tmp_path: Path):
    p = _make_docx(tmp_path / "sample.docx")
    h0 = hashlib.sha256(p.read_bytes()).hexdigest()
    doc = parse_docx_file(p, document_id="docx1", document_type="general_report")
    assert doc.source_format == "docx"
    assert any(s.heading == "요약" for s in doc.sections)
    assert hashlib.sha256(p.read_bytes()).hexdigest() == h0


def test_03_nested_heading():
    md = "# Root\n\n## Parent\n\n### Child\n\ntext\n"
    doc = normalize_document(parse_markdown_text(md, document_id="nest"))
    parent = next(s for s in doc.sections if s.heading == "Parent")
    child = next(s for s in doc.sections if s.heading == "Child")
    assert child.parent == parent.section_id
    assert child.section_id in parent.children
    assert child.heading_level == 3
    assert parent.heading_level == 2


def test_04_table_parsing_markdown():
    doc = parse_markdown_text(SAMPLE_MARKDOWN_REPORT, document_id="tbl")
    tables = [t for s in doc.sections for t in s.tables]
    assert tables
    assert "출처" in tables[0].headers


def test_05_list_parsing_markdown():
    doc = parse_markdown_text(SAMPLE_MARKDOWN_REPORT, document_id="lst")
    lists = [lst for s in doc.sections for lst in s.lists]
    assert lists
    assert any("requirement_id" in item for lst in lists for item in lst.items)


def test_06_ordered_list_parsing():
    doc = parse_markdown_text(SAMPLE_MARKDOWN_PROPOSAL, document_id="ol")
    lists = [lst for s in doc.sections for lst in s.lists]
    assert any(lst.ordered for lst in lists)


def test_07_docx_table_and_list(tmp_path: Path):
    p = _make_docx(tmp_path / "t2.docx")
    doc = parse_docx_file(p)
    assert any(s.tables for s in doc.sections)
    assert any(s.lists for s in doc.sections)


def test_08_in_memory_object():
    doc = document_from_object(
        {
            "sample_document_id": "obj1",
            "document_type": "general_report",
            "title": "오브젝트 문서",
            "sections": {
                "methodology": {
                    "title": "방법론",
                    "data_sources": "fixture",
                    "key_findings": ["a", "b"],
                }
            },
        }
    )
    assert doc.source_format == "object"
    assert any(s.heading == "방법론" for s in doc.sections)
    assert any(s.lists for s in doc.sections)


def test_09_normalize_rebuilds_children():
    doc = parse_markdown_text("# A\n\n## B\n\nx\n", document_id="n")
    doc.sections[1].children = []
    fixed = normalize_document(doc)
    parent = next(s for s in fixed.sections if s.heading == "A")
    child = next(s for s in fixed.sections if s.heading == "B")
    assert child.section_id in parent.children


def test_10_section_tree():
    doc = normalize_document(
        parse_markdown_text(SAMPLE_MARKDOWN_REPORT, document_id="tree")
    )
    tree = build_section_tree(doc)
    assert tree["document_id"] == "tree"
    assert tree["roots"]


def test_11_12_13_locator_candidates():
    doc = normalize_document(
        parse_markdown_text(SAMPLE_MARKDOWN_REPORT, document_id="cand")
    )
    cands = build_locator_candidates(doc)
    assert cands
    types = {c.locator_type for c in cands}
    assert "heading_path" in types
    assert "section_name" in types
    assert "heading_text" in types
    assert all(0 <= c.score <= 1 for c in cands)
    meth = next(s for s in doc.sections if s.heading == "데이터 출처")
    path_cands = [c for c in cands if c.section_id == meth.section_id]
    assert any(c.heading_path for c in path_cands)


def test_14_map_document_structure():
    mapped = map_document_structure(
        parse_markdown_text(SAMPLE_MARKDOWN_REPORT, document_id="map1")
    )
    assert mapped["document"]["document_id"] == "map1"
    assert mapped["locator_candidates"]
    assert mapped["validation"]["status"] in ("VALID", "VALID_WITH_WARNINGS", "INVALID")


def test_15_engine_default_samples():
    pkg = run_document_structure_mapping_engine()
    assert pkg["stage"] == "document_structure_mapping_engine"
    assert len(pkg["documents"]) >= 2
    assert pkg["summary"]["candidate_count"] > 0
    assert pkg["actual_docx_changed"] is False


def test_16_deterministic_output():
    a = run_document_structure_mapping_engine(
        markdown_texts=[
            {
                "text": SAMPLE_MARKDOWN_REPORT,
                "document_id": "det",
                "document_type": "general_report",
            }
        ],
        objects=[],
    )
    b = run_document_structure_mapping_engine(
        markdown_texts=[
            {
                "text": SAMPLE_MARKDOWN_REPORT,
                "document_id": "det",
                "document_type": "general_report",
            }
        ],
        objects=[],
    )
    assert json.dumps(a["documents"], sort_keys=True) == json.dumps(
        b["documents"], sort_keys=True
    )
    assert json.dumps(a["locator_candidates"], sort_keys=True) == json.dumps(
        b["locator_candidates"], sort_keys=True
    )


def test_17_summary_fields():
    pkg = run_document_structure_mapping_engine(
        markdown_texts=[
            {
                "text": SAMPLE_MARKDOWN_REPORT,
                "document_id": "sum",
                "document_type": "general_report",
            }
        ],
        objects=[],
    )
    s = pkg["summary"]
    for k in (
        "heading_count",
        "section_count",
        "paragraph_count",
        "table_count",
        "list_count",
        "candidate_count",
        "validation_status",
    ):
        assert k in s


def test_18_statistics_match_document():
    pkg = run_document_structure_mapping_engine(
        markdown_texts=[
            {
                "text": SAMPLE_MARKDOWN_PROPOSAL,
                "document_id": "st",
                "document_type": "business_proposal",
            }
        ],
        objects=[],
    )
    doc = pkg["documents"][0]
    st = pkg["statistics"]["documents"][0]
    assert st["section_count"] == len(doc["sections"])
    assert st["paragraph_count"] == sum(len(s["paragraphs"]) for s in doc["sections"])


def test_19_negative_duplicate_heading():
    doc = DocumentModel(
        document_id="dup",
        document_type="x",
        title="t",
        source_format="object",
        sections=[
            SectionModel("a", "동일", 1, None, order=1),
            SectionModel("b", "동일", 1, None, order=2),
        ],
    )
    v = validate_document_structure(doc)
    assert v["status"] == "INVALID"
    assert any("duplicate_heading" in i for i in v["issues"])


def test_20_negative_invalid_hierarchy():
    doc = DocumentModel(
        document_id="hier",
        document_type="x",
        title="t",
        source_format="object",
        sections=[
            SectionModel("p", "Parent", 3, None, children=["c"], order=1),
            SectionModel("c", "Child", 2, "p", order=2),
        ],
    )
    v = validate_document_structure(doc)
    assert any("invalid_hierarchy" in i for i in v["issues"])


def test_21_negative_parent_cycle():
    doc = DocumentModel(
        document_id="cyc",
        document_type="x",
        title="t",
        source_format="object",
        sections=[
            SectionModel("a", "A", 1, "b", order=1),
            SectionModel("b", "B", 2, "a", order=2),
        ],
    )
    v = validate_document_structure(doc)
    assert any("parent_cycle" in i for i in v["issues"])


def test_22_negative_broken_tree():
    doc = DocumentModel(
        document_id="brk",
        document_type="x",
        title="t",
        source_format="object",
        sections=[
            SectionModel("a", "A", 1, "missing", order=1),
        ],
    )
    v = validate_document_structure(doc)
    assert any("broken_tree" in i for i in v["issues"])


def test_23_negative_empty_heading():
    doc = DocumentModel(
        document_id="eh",
        document_type="x",
        title="t",
        source_format="object",
        sections=[SectionModel("a", "  ", 1, None, order=1)],
    )
    v = validate_document_structure(doc)
    assert any("empty_heading" in i for i in v["issues"])


def test_24_negative_invalid_heading_level():
    doc = DocumentModel(
        document_id="lvl",
        document_type="x",
        title="t",
        source_format="object",
        sections=[SectionModel("a", "X", 9, None, order=1)],
    )
    v = validate_document_structure(doc)
    assert any("invalid_heading_level" in i for i in v["issues"])


def test_25_empty_section_warning():
    doc = DocumentModel(
        document_id="es",
        document_type="x",
        title="t",
        source_format="object",
        sections=[SectionModel("a", "Empty", 1, None, order=1)],
    )
    v = validate_document_structure(doc)
    assert any("empty_section" in w for w in v["warnings"])


def test_26_malformed_markdown_still_parses():
    doc = parse_markdown_text("not a heading\n\n### orphan child\n", document_id="mal")
    assert doc.sections


def test_27_no_semantic_matching_invariants():
    mapped = map_document_structure(
        parse_markdown_text(SAMPLE_MARKDOWN_REPORT, document_id="nosem")
    )
    inv = mapped["validation"]["invariants"]
    assert inv["no_semantic_matching"] is True
    assert inv["no_llm"] is True
    assert inv["no_docx_mutation"] is True


def test_28_pr18_non_mutation():
    items = [
        {
            "review_item_id": "CRI-RP-1",
            "patch_id": "RP-1",
            "document": "MDSR",
            "requirement_id": "Req. 105",
            "field": "criteria",
            "change_type": "UPDATE",
            "review_status": "READY",
            "acu_id": "ACU-001",
        }
    ]
    snap = copy.deepcopy(items)
    a = run_template_abstraction_layer(review_items=items)
    _ = run_document_structure_mapping_engine()
    b = run_template_abstraction_layer(review_items=items)
    assert items == snap
    assert a["summary"]["mapped_count"] == b["summary"]["mapped_count"]
    tids = {t["template_id"] for t in a["registry"]["templates"]}
    assert tids == {"mdsr_v1", "mddr_v1"}


def test_29_pr19_non_mutation():
    a = run_generic_document_template_pack()
    _ = run_document_structure_mapping_engine()
    b = run_generic_document_template_pack()
    assert a["summary"]["mapped_count"] == b["summary"]["mapped_count"]
    assert {t["template_id"] for t in a["registry"]["templates"]} == {
        "general_report_v1",
        "business_proposal_v1",
    }


def test_30_docx_writer_flag_off():
    assert is_docx_activation_enabled(env={}) is False
    pkg = run_document_structure_mapping_engine()
    assert pkg["actual_docx_changed"] is False
    assert pkg["actual_generation_changed"] is False


def test_31_source_docx_immutable(tmp_path: Path):
    p = _make_docx(tmp_path / "imm.docx")
    h0 = hashlib.sha256(p.read_bytes()).hexdigest()
    _ = run_document_structure_mapping_engine(
        docx_paths=[p], markdown_texts=[], objects=[]
    )
    assert hashlib.sha256(p.read_bytes()).hexdigest() == h0


def test_32_freeze_paths_exist():
    for p in FROZEN_PATHS:
        assert p.exists()


def test_33_heading_path_for_nested():
    doc = normalize_document(
        parse_markdown_text(SAMPLE_MARKDOWN_REPORT, document_id="hp")
    )
    cands = build_locator_candidates(doc)
    ds = next(s for s in doc.sections if s.heading == "데이터 출처")
    hp = next(
        c
        for c in cands
        if c.section_id == ds.section_id and c.locator_type == "heading_path"
    )
    assert "방법론" in hp.heading_path
    assert "데이터 출처" in hp.heading_path


def test_34_candidate_count_three_per_section():
    doc = normalize_document(
        parse_markdown_text("# T\n\n## S\n\nbody\n", document_id="c3")
    )
    cands = build_locator_candidates(doc)
    assert len(cands) == len(doc.sections) * 3


def test_35_engine_with_docx_and_markdown(tmp_path: Path):
    p = _make_docx(tmp_path / "mix.docx")
    pkg = run_document_structure_mapping_engine(
        markdown_texts=[
            {
                "text": "# Mix\n\n## A\n\nx\n",
                "document_id": "mix_md",
                "document_type": "general_report",
            }
        ],
        docx_paths=[p],
        objects=[],
    )
    assert len(pkg["documents"]) == 2
    assert pkg["summary"]["document_count"] == 2
