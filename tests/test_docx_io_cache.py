from docx import Document

from document_ai.learn.docx_io import load_document


def test_load_document_cache_refreshes_on_mtime(tmp_path):
    """Different paths with the same filename must not share a stale cache entry."""
    a = tmp_path / "a" / "same.docx"
    b = tmp_path / "b" / "same.docx"
    a.parent.mkdir()
    b.parent.mkdir()

    doc_a = Document()
    doc_a.add_paragraph("version-a")
    doc_a.save(a)

    loaded_a = load_document(a)
    assert loaded_a.paragraphs[0].text == "version-a"

    doc_b = Document()
    doc_b.add_paragraph("version-b")
    doc_b.save(b)

    loaded_b = load_document(b)
    assert loaded_b.paragraphs[0].text == "version-b"

    doc_a2 = Document()
    doc_a2.add_paragraph("version-a2")
    doc_a2.save(a)

    reloaded_a = load_document(a)
    assert reloaded_a.paragraphs[0].text == "version-a2"
