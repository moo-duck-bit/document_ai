# -*- coding: utf-8 -*-
"""Document identity signal extraction tests."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document

from document_ai.document_identity.filename_signals import extract_filename_signals
from document_ai.document_identity.signal_extractor import (
    extract_all_signals,
    extract_content_structure_signals,
)


REPO = Path(__file__).resolve().parents[1]
MDTM = REPO / "data" / "eval" / "document_set_benchmark_v2" / "fixtures" / "ec_sw" / "mdtm_base.docx"
REPORT = (
    REPO / "data" / "eval" / "document_set_benchmark_v2" / "fixtures" / "general_report" / "report_std.docx"
)
PROPOSAL = (
    REPO
    / "data"
    / "eval"
    / "document_set_benchmark_v2"
    / "fixtures"
    / "business_proposal"
    / "proposal_base.docx"
)


def test_filename_short_id_mdtm():
    sigs = extract_filename_signals(source_document_id="S1", filename="trace_matrix_final_v3.docx")
    assert any(s.signal_type == "FILENAME_TOKEN" and s.normalized_value == "MDTM" for s in sigs)


def test_filename_short_id_mdsr():
    sigs = extract_filename_signals(source_document_id="S1", filename="요구사항_스펙.docx")
    assert any(s.normalized_value == "MDSR" for s in sigs)


def test_filename_short_id_mddr():
    sigs = extract_filename_signals(source_document_id="S1", filename="설계문서_MDDR.docx")
    assert any(s.normalized_value == "MDDR" for s in sigs)


def test_filename_report_hint():
    sigs = extract_filename_signals(source_document_id="S1", filename="2026_중간성과자료.docx")
    assert any(s.normalized_value == "REPORT" for s in sigs)


def test_filename_proposal_hint():
    sigs = extract_filename_signals(source_document_id="S1", filename="사업_제안서.docx")
    assert any(s.normalized_value == "PROPOSAL" for s in sigs)


def test_title_and_heading_signals_from_mdtm():
    if not MDTM.is_file():
        return
    sigs = extract_content_structure_signals(source_document_id="S1", docx_path=MDTM)
    assert any(s.signal_type in {"DOCUMENT_TITLE", "HEADING_TEXT"} for s in sigs)


def test_table_header_signal():
    if not MDTM.is_file():
        return
    sigs = extract_content_structure_signals(source_document_id="S1", docx_path=MDTM)
    assert any(s.signal_type == "TABLE_HEADER" for s in sigs)


def test_identifier_pattern_signal_mdtm_row():
    if not MDTM.is_file():
        return
    sigs = extract_content_structure_signals(source_document_id="S1", docx_path=MDTM)
    assert any(s.signal_type == "IDENTIFIER_PATTERN" for s in sigs)
    assert any("mdtm_strong_evidence" in s.reason_codes for s in sigs)


def test_template_coverage_report():
    if not REPORT.is_file():
        return
    sigs = extract_content_structure_signals(source_document_id="S1", docx_path=REPORT)
    assert any(s.signal_type == "CONTENT_CONCEPT" and s.normalized_value == "general_report" for s in sigs)


def test_template_coverage_proposal():
    if not PROPOSAL.is_file():
        return
    sigs = extract_content_structure_signals(source_document_id="S1", docx_path=PROPOSAL)
    assert any(
        s.signal_type == "CONTENT_CONCEPT" and s.normalized_value == "business_proposal" for s in sigs
    )


def test_file_format_signal():
    sigs = extract_filename_signals(source_document_id="S1", filename="x.docx")
    assert any(s.signal_type == "FILE_FORMAT" for s in sigs)


def test_user_hint_signal():
    sigs = extract_all_signals(
        source_document_id="S1",
        filename="unknown.docx",
        user_hints={"short_id": "MDTM"},
    )
    assert any(s.signal_type == "EXPLICIT_USER_HINT" for s in sigs)


def test_bytes_docx_extraction():
    doc = Document()
    doc.add_heading("Traceability Matrix", 1)
    t = doc.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text = "Requirement"
    t.rows[0].cells[1].text = "Design"
    t.rows[0].cells[2].text = "Test"
    t.rows[1].cells[0].text = "Req. 1"
    t.rows[1].cells[1].text = "1.0"
    t.rows[1].cells[2].text = "TC-1"
    buf = io.BytesIO()
    doc.save(buf)
    sigs = extract_content_structure_signals(
        source_document_id="S1", docx_path=None, file_bytes=buf.getvalue()
    )
    assert any(s.signal_type == "IDENTIFIER_PATTERN" for s in sigs)
