from docx import Document

from document_ai.impact.change import merge_requirement_changes
from document_ai.render.patch import patch_mdsr_requirements


def _req_table(req_id: str) -> Document:
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = req_id
    table.rows[1].cells[0].text = "설명"
    table.rows[2].cells[0].text = "목적"
    table.rows[3].cells[0].text = "기준"
    return doc


def test_merge_requirement_changes_updates_purpose_and_criteria():
    payload = {
        "requirements": [
            {
                "req_id": "Req. 6",
                "description": "old desc",
                "purpose": "old purpose",
                "criteria": "old criteria",
            }
        ]
    }
    change = {
        "requirement_changes": [
            {
                "req_id": "Req. 6",
                "description": "new desc",
                "purpose": "new purpose",
                "criteria": "new criteria",
            }
        ]
    }
    updated = merge_requirement_changes(payload, change)
    row = payload["requirements"][0]
    assert updated == ["Req. 6"]
    assert row["description"] == "new desc"
    assert row["purpose"] == "new purpose"
    assert row["criteria"] == "new criteria"


def test_patch_mdsr_requirements_fills_all_labeled_rows():
    doc = _req_table("Req. 6")
    table = doc.tables[0]
    patched = patch_mdsr_requirements(
        doc,
        [
            {
                "req_id": "Req. 6",
                "description": "desc text",
                "purpose": "purpose text",
                "criteria": "criteria text",
            }
        ],
    )
    assert patched == ["Req. 6"]
    assert table.rows[1].cells[1].text == "desc text"
    assert table.rows[2].cells[1].text == "purpose text"
    assert table.rows[3].cells[1].text == "criteria text"
