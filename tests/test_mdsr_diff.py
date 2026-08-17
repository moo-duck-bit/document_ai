from document_ai.learn.mdsr_diff import compare_mdsr_documents, extract_req_blocks


def test_extract_req_blocks_by_section():
    from docx import Document

    doc = Document()
    t = doc.add_table(rows=4, cols=2)
    t.rows[0].cells[0].text = "Req. 5"
    t.rows[1].cells[0].text = "설명"
    t.rows[1].cells[1].text = "desc"
    blocks = extract_req_blocks(doc)
    assert "Req. 5" in blocks
    assert blocks["Req. 5"].section == "functional"
    assert blocks["Req. 5"].by_label()["설명"] == "desc"


def test_compare_detects_missing_criteria(tmp_path):
    from docx import Document

    gold = Document()
    gt = gold.add_table(rows=4, cols=2)
    gt.rows[0].cells[0].text = "Req. 1"
    gt.rows[1].cells[1].text = "gold only"

    out = Document()
    ot = out.add_table(rows=4, cols=2)
    ot.rows[0].cells[0].text = "Req. 1"
    ot.rows[1].cells[0].text = "설명"
    ot.rows[1].cells[1].text = "out desc"

    gp = tmp_path / "gold.docx"
    op = tmp_path / "out.docx"
    gold.save(gp)
    out.save(op)

    report = compare_mdsr_documents(gp, op)
    kinds = {g.kind for g in report.quality_issues + report.gaps}
    assert "missing_criteria" in kinds or "filled_vs_empty_gold" in kinds
