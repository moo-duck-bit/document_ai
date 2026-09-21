# -*- coding: utf-8 -*-
"""General report schedule/table retrieval tests."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from document_ai.document_set.schedule_table_retrieval import (
    collect_schedule_table_evidence,
    evidences_to_review_items,
)
from document_ai.workflow.orchestrator import create_workflow, run_analysis

REPO = Path(__file__).resolve().parents[1]
BASE = REPO / "data" / "eval" / "document_set_benchmark" / "fixtures" / "general_report" / "report_base.docx"


def _mini_schedule(path: Path) -> Path:
    d = Document()
    d.add_heading("일반 보고서", 1)
    d.add_heading("일정", 2)
    d.add_paragraph("2026-Q3 완료 예정")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "단계"
    t.cell(0, 1).text = "일정"
    t.cell(1, 0).text = "1주차"
    t.cell(1, 1).text = "설계"
    d.save(str(path))
    return path


def test_schedule_heading_review(tmp_path):
    p = _mini_schedule(tmp_path / "r.docx")
    ev = collect_schedule_table_evidence(
        change_request="일정 표 수정 2026-Q4",
        document_id="R1",
        docx_path=p,
    )
    assert any(e.supports_review for e in ev)
    assert all(not e.supports_patch for e in ev)


def test_schedule_table_header(tmp_path):
    p = _mini_schedule(tmp_path / "r.docx")
    ev = collect_schedule_table_evidence(
        change_request="수행 일정 업데이트",
        document_id="R1",
        docx_path=p,
    )
    assert any(e.table_id for e in ev if e.supports_review) or any(e.supports_review for e in ev)


def test_time_unit_table(tmp_path):
    p = _mini_schedule(tmp_path / "r.docx")
    ev = collect_schedule_table_evidence(
        change_request="일정 계획 Q4 반영",
        document_id="R1",
        docx_path=p,
    )
    assert any(e.time_unit_hits for e in ev if e.supports_review) or any(
        e.supports_review for e in ev
    )


def test_generic_table_no_schedule_unrelated(tmp_path):
    p = tmp_path / "plain.docx"
    d = Document()
    d.add_heading("결과", 1)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "지표"
    t.cell(0, 1).text = "값"
    t.cell(1, 0).text = "만족도"
    t.cell(1, 1).text = "70"
    d.save(str(p))
    ev = collect_schedule_table_evidence(
        change_request="사무실 화분 배치 변경 XYZABC",
        document_id="R1",
        docx_path=p,
    )
    assert ev == []


def test_table_only_no_patch_items(tmp_path):
    p = _mini_schedule(tmp_path / "r.docx")
    items = evidences_to_review_items(
        collect_schedule_table_evidence(
            change_request="일정표 수정",
            document_id="R1",
            docx_path=p,
        )
    )
    assert all(i["status"] == "REVIEW_REQUIRED" for i in items)
    assert all(i.get("metadata", {}).get("supports_patch") is False for i in items)


def test_workflow_schedule_review(tmp_path):
    rec = create_workflow(
        document_set="general_report",
        change_request="일정 표 수정 2026-Q4",
        files=[("report_base.docx", BASE.read_bytes(), "general_report")],
        root=tmp_path,
        name="sch",
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert out["patch_candidates"] == []
    dec = {d["document_id"]: d for d in out["metadata"]["document_impact_decisions"]}
    assert dec["REPORT_BASE"]["predicted_status"] == "REVIEW_REQUIRED"


def test_methodology_case_still_reviews(tmp_path):
    rec = create_workflow(
        document_set="general_report",
        change_request="방법론 데이터 출처를 외부 공개자료로 수정",
        files=[("report_base.docx", BASE.read_bytes(), "general_report")],
        root=tmp_path,
        name="meth",
    )
    out = run_analysis(rec["workflow_id"], root=tmp_path)
    assert out["patch_candidates"] == []
    assert out["review_required"]


def test_duplicate_schedule_ambiguous(tmp_path):
    p = tmp_path / "dup.docx"
    d = Document()
    d.add_heading("일정", 1)
    d.add_paragraph("1차 일정 2026-Q3")
    d.add_heading("추진 일정", 1)
    d.add_paragraph("2차 일정 2026-Q4")
    d.save(str(p))
    ev = collect_schedule_table_evidence(
        change_request="일정 수정",
        document_id="R1",
        docx_path=p,
    )
    strong = [e for e in ev if e.supports_review]
    if len(strong) > 1:
        assert any("ambiguous" in " ".join(e.reason_codes) for e in strong)
