# -*- coding: utf-8 -*-
"""Draft/review/sealed protocol + Pass1/Pass2 labeling tests (Cycle 6)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from document_ai.evaluation.business_proposal_gold import labeling_pass, protocol

REPO = Path(__file__).resolve().parents[1]


def _build_base(path: Path) -> None:
    doc = Document()
    doc.add_heading("사업 제안서", level=0)
    doc.add_heading("실행 일정", level=1)
    doc.add_paragraph("2026-Q3 착수, 2026-Q4 완료")
    doc.add_heading("예산", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "금액"
    table.rows[1].cells[0].text = "인건비"
    table.rows[1].cells[1].text = "100"
    doc.save(str(path))


def test_gold_dirs_layout(tmp_path):
    dirs = protocol.gold_dirs(tmp_path)
    assert dirs["root"] == tmp_path
    assert dirs["draft"] == tmp_path / "draft"
    assert dirs["review"] == tmp_path / "review"
    assert dirs["sealed"] == tmp_path / "sealed"


def test_write_and_read_jsonl_round_trip(tmp_path):
    rows = [{"case_id": "a", "x": 1}, {"case_id": "b", "x": 2}]
    path = tmp_path / "out.jsonl"
    protocol.write_jsonl(path, rows)
    read_back = protocol.read_jsonl(path)
    assert read_back == rows


def test_read_jsonl_missing_file_returns_empty(tmp_path):
    assert protocol.read_jsonl(tmp_path / "missing.jsonl") == []


def test_write_draft_and_review(tmp_path):
    protocol.write_draft([{"case_id": "a"}], gold_root=tmp_path)
    protocol.write_review([{"case_id": "a"}], gold_root=tmp_path)
    assert (tmp_path / "draft" / protocol.DRAFT_FILENAME).is_file()
    assert (tmp_path / "review" / protocol.REVIEW_FILENAME).is_file()


def test_seal_gold_writes_manifest_and_hashes(tmp_path):
    rows = [{"case_id": "a", "node_evaluation_mode": "REQUIRED"}]
    manifest = protocol.seal_gold(rows, gold_root=tmp_path)
    assert manifest["status"] == "SEALED"
    assert manifest["n_rows"] == 1
    assert "aggregate_hash" in manifest
    assert (tmp_path / "sealed" / protocol.SEALED_FILENAME).is_file()
    assert (tmp_path / protocol.MANIFEST_FILENAME).is_file()
    assert (tmp_path / protocol.HASHES_FILENAME).is_file()


def test_load_sealed_gold_round_trip(tmp_path):
    rows = [{"case_id": "a", "node_evaluation_mode": "OPTIONAL"}]
    protocol.seal_gold(rows, gold_root=tmp_path)
    loaded = protocol.load_sealed_gold(gold_root=tmp_path)
    assert loaded == rows


def test_prediction_code_must_not_read_sealed_flags_offending_file(tmp_path):
    bad_file = tmp_path / "prediction_adapter_fake.py"
    bad_file.write_text(
        "x = 'data/eval/.../business_proposal_node_gold/sealed/foo.jsonl'\n", encoding="utf-8"
    )
    bad = protocol.prediction_code_must_not_read_sealed_bp_gold([str(bad_file)])
    assert str(bad_file) in bad


def test_prediction_code_must_not_read_sealed_ignores_gold_package_itself(tmp_path):
    own_file = tmp_path / "business_proposal_gold_helper.py"
    own_file.write_text("SEALED_TOKEN = 'business_proposal_node_gold/sealed'\n", encoding="utf-8")
    bad = protocol.prediction_code_must_not_read_sealed_bp_gold([str(own_file)])
    assert bad == []


def test_prediction_code_must_not_read_sealed_clean_for_clean_file(tmp_path):
    clean_file = tmp_path / "prediction_adapter_clean.py"
    clean_file.write_text("x = 1\n", encoding="utf-8")
    bad = protocol.prediction_code_must_not_read_sealed_bp_gold([str(clean_file)])
    assert bad == []


def test_real_prediction_code_paths_do_not_read_sealed_bp_gold():
    paths = [
        REPO / "src" / "document_ai" / "evaluation" / "document_set" / "prediction_adapter.py",
        REPO / "src" / "document_ai" / "workflow" / "analysis.py",
        REPO / "src" / "document_ai" / "domain_packs" / "business_proposal" / "node_ranking.py",
        REPO / "src" / "document_ai" / "domain_packs" / "business_proposal" / "structural_match.py",
        REPO / "src" / "document_ai" / "domain_packs" / "business_proposal" / "query_intent.py",
        REPO / "src" / "document_ai" / "document_set" / "proposal_table_retrieval.py",
    ]
    existing = [str(p) for p in paths if p.is_file()]
    assert existing, "expected at least the prediction_adapter.py file to exist"
    bad = protocol.prediction_code_must_not_read_sealed_bp_gold(existing)
    assert bad == []


def test_resolve_fixture_direct_path(tmp_path):
    fixture = tmp_path / "x.docx"
    _build_base(fixture)
    case = {"case_id": "c1", "input_documents": [{"path": str(fixture)}]}
    resolved = labeling_pass.resolve_fixture(case, tmp_path)
    assert resolved == fixture


def test_resolve_fixture_missing_raises(tmp_path):
    case = {"case_id": "c1", "input_documents": [{"path": "fixtures/business_proposal/nope.docx"}]}
    try:
        labeling_pass.resolve_fixture(case, tmp_path)
        assert False, "expected FileNotFoundError"
    except FileNotFoundError:
        pass


def test_resolve_fixture_no_documents_raises(tmp_path):
    case = {"case_id": "c1", "input_documents": []}
    try:
        labeling_pass.resolve_fixture(case, tmp_path)
        assert False, "expected ValueError"
    except ValueError:
        pass


def test_label_case_produces_gold_row(tmp_path):
    fixture = tmp_path / "proposal_base.docx"
    _build_base(fixture)
    case = {
        "case_id": "c_sched",
        "change_request": "실행 일정 2026-Q4 조정",
        "tags": ["schedule"],
        "enabled_documents": ["PROPOSAL_BASE"],
    }
    row = labeling_pass.label_case(case=case, docx_path=fixture, labeled_by="pass1_structure_policy")
    assert row.case_id == "c_sched"
    assert row.node_evaluation_mode == "REQUIRED"
    assert row.labeled_by == "pass1_structure_policy"


def test_pass1_pass2_agree_deterministically(tmp_path):
    fixture_dir = tmp_path / "fixtures"
    fixture_dir.mkdir()
    fixture = fixture_dir / "proposal_base.docx"
    _build_base(fixture)
    cases = [
        {
            "case_id": "c_sched",
            "change_request": "실행 일정 2026-Q4 조정",
            "tags": ["schedule"],
            "enabled_documents": ["PROPOSAL_BASE"],
            "input_documents": [{"path": str(fixture)}],
        },
        {
            "case_id": "c_budget",
            "change_request": "예산 표 인건비 수정",
            "tags": ["budget", "table"],
            "enabled_documents": ["PROPOSAL_BASE"],
            "input_documents": [{"path": str(fixture)}],
        },
    ]
    p1 = labeling_pass.run_pass1(cases, fixtures_root=fixture_dir)
    p2 = labeling_pass.run_pass2(cases, fixtures_root=fixture_dir)
    disagreements = labeling_pass.detect_disagreements(p1, p2)
    assert disagreements == []
    agreement = labeling_pass.agreement_metrics(p1, p2)
    assert agreement["raw_agreement_rate"] == 1.0
    assert agreement["mode_agreement_rate"] == 1.0
    assert agreement["reference_agreement_rate"] == 1.0
    assert agreement["n_common_cases"] == 2


def test_detect_disagreements_flags_field_mismatch():
    from dataclasses import replace

    from document_ai.evaluation.business_proposal_gold.schema import BusinessProposalGoldRow

    r1 = BusinessProposalGoldRow(
        case_id="c1",
        document_id="D1",
        node_evaluation_mode="REQUIRED",
        primary_reference={"template_node_id": "t1", "document_node_id": "p1"},
    )
    r2 = replace(r1, node_evaluation_mode="OPTIONAL", primary_reference={"template_node_id": None, "document_node_id": None})
    disagreements = labeling_pass.detect_disagreements([r1], [r2])
    assert len(disagreements) == 1
    assert disagreements[0]["case_id"] == "c1"
    assert disagreements[0]["issue"] == "FIELD_MISMATCH"


def test_detect_disagreements_missing_in_one_pass():
    from document_ai.evaluation.business_proposal_gold.schema import BusinessProposalGoldRow

    r1 = BusinessProposalGoldRow(case_id="only_in_1", document_id="D1", node_evaluation_mode="REQUIRED")
    disagreements = labeling_pass.detect_disagreements([r1], [])
    assert len(disagreements) == 1
    assert disagreements[0]["issue"] == "MISSING_IN_ONE_PASS"


def test_agreement_metrics_empty_inputs_do_not_crash():
    agreement = labeling_pass.agreement_metrics([], [])
    assert agreement["n_common_cases"] == 0
