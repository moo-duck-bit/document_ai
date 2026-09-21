# -*- coding: utf-8 -*-
"""Real User Document Pilot v2 — orchestrator flow tests (create -> upload -> identity ->
analyze -> decisions -> result), using an isolated tmp_path session root and pilot-owned
fixture docx files (never the frozen benchmark corpus).
"""

from __future__ import annotations

import pytest
from docx import Document

from document_ai.pilot_v2 import orchestrator as pv2
from document_ai.pilot_v2 import scenarios
from document_ai.pilot_v2.orchestrator import PilotOrchestratorError


def _small_docx_bytes(tmp_path, name: str = "src.docx") -> bytes:
    path = tmp_path / name
    doc = Document()
    doc.add_heading("Traceability Matrix (Test)", level=1)
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = "Requirement", "Design", "Test"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text, table.rows[1].cells[2].text = "Req. 11", "5.1.2", "TC-11"
    doc.save(str(path))
    return path.read_bytes()


# --- create_session -------------------------------------------------------------


def test_create_session_success(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="Req. 11 update", root=tmp_path)
    assert session["status"] == "CREATED"
    assert session["participant_id"].startswith("P")
    assert session["session_id"]


def test_create_session_rejects_unknown_document_set(tmp_path):
    with pytest.raises(PilotOrchestratorError) as exc:
        pv2.create_session(document_set="unknown_set", change_request="x", root=tmp_path)
    assert exc.value.reason_code == "UNKNOWN_DOCUMENT_SET"


def test_create_session_rejects_empty_change_request(tmp_path):
    with pytest.raises(PilotOrchestratorError) as exc:
        pv2.create_session(document_set="ec_sw", change_request="   ", root=tmp_path)
    assert exc.value.reason_code == "EMPTY_CHANGE_REQUEST"


def test_create_session_uses_explicit_participant_id(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", participant_id="P042", root=tmp_path)
    assert session["participant_id"] == "P042"


# --- upload_documents -------------------------------------------------------------


def test_upload_documents_success(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    sid = session["session_id"]
    content = _small_docx_bytes(tmp_path)
    session = pv2.upload_documents(sid, [("MDTM.docx", content, "traceability")], root=tmp_path)
    assert session["status"] == "UPLOADED"
    assert len(session["documents"]) == 1
    assert session["documents"][0]["document_id"]


def test_upload_documents_requires_at_least_one_file(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    with pytest.raises(PilotOrchestratorError) as exc:
        pv2.upload_documents(session["session_id"], [], root=tmp_path)
    assert exc.value.reason_code == "NO_DOCUMENTS"


def test_upload_documents_rejects_duplicate_filenames(tmp_path):
    from document_ai.pilot_v2 import security

    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    content = _small_docx_bytes(tmp_path)
    with pytest.raises(security.PilotSecurityError) as exc:
        pv2.upload_documents(
            session["session_id"], [("A.docx", content, None), ("A.docx", content, None)], root=tmp_path
        )
    assert exc.value.reason_code == "DUPLICATE_FILENAME"


def test_upload_documents_never_writes_to_original_bytes_source(tmp_path):
    # The bytes passed in are in-memory; assert the session copy lives strictly
    # under the session's own input/ directory (never any path outside root).
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    content = _small_docx_bytes(tmp_path)
    session = pv2.upload_documents(session["session_id"], [("A.docx", content, None)], root=tmp_path)
    doc_path = session["documents"][0]["path"]
    assert "/input/" in doc_path.replace("\\", "/")
    assert str(tmp_path).replace("\\", "/") in doc_path.replace("\\", "/")


# --- resolve_identity --------------------------------------------------------------


def test_resolve_identity_explicit_mode_resolves_immediately(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    sid = session["session_id"]
    content = _small_docx_bytes(tmp_path)
    pv2.upload_documents(sid, [("MDTM.docx", content, "traceability")], root=tmp_path)
    session = pv2.resolve_identity(sid, routing_mode="explicit", root=tmp_path)
    assert session["status"] == "IDENTITY_RESOLVED"


def test_resolve_identity_before_upload_is_invalid_state(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    with pytest.raises(PilotOrchestratorError) as exc:
        pv2.resolve_identity(session["session_id"], root=tmp_path)
    assert exc.value.reason_code == "INVALID_SESSION_STATE"


# --- analyze -----------------------------------------------------------------------


def test_analyze_requires_identity_resolved_state(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    sid = session["session_id"]
    content = _small_docx_bytes(tmp_path)
    pv2.upload_documents(sid, [("MDTM.docx", content, "traceability")], root=tmp_path)
    with pytest.raises(PilotOrchestratorError) as exc:
        pv2.analyze(sid, root=tmp_path)
    assert exc.value.reason_code == "INVALID_SESSION_STATE"


def test_analyze_produces_review_items_and_transitions(tmp_path):
    scenario = scenarios.get_scenario("pilot_ec_req_single")
    fixture_path = scenarios.ensure_fixture(scenario)
    session = pv2.create_session(
        document_set=scenario.document_set, change_request=scenario.change_request, scenario_id=scenario.scenario_id, root=tmp_path
    )
    sid = session["session_id"]
    pv2.upload_documents(sid, [(fixture_path.name, fixture_path.read_bytes(), scenario.fixture_role)], root=tmp_path)
    pv2.resolve_identity(sid, routing_mode="explicit", root=tmp_path)
    session = pv2.analyze(sid, root=tmp_path)
    assert session["status"] in {"REVIEW_IN_PROGRESS", "REVIEW_COMPLETE"}
    assert isinstance(session["review_items"], list)
    assert session["workflow_ref"]["workflow_id"]


def test_analyze_no_impact_scenario_yields_no_or_review_required_items(tmp_path):
    scenario = scenarios.get_scenario("pilot_ec_no_impact")
    fixture_path = scenarios.ensure_fixture(scenario)
    session = pv2.create_session(
        document_set=scenario.document_set, change_request=scenario.change_request, scenario_id=scenario.scenario_id, root=tmp_path
    )
    sid = session["session_id"]
    pv2.upload_documents(sid, [(fixture_path.name, fixture_path.read_bytes(), scenario.fixture_role)], root=tmp_path)
    pv2.resolve_identity(sid, routing_mode="explicit", root=tmp_path)
    session = pv2.analyze(sid, root=tmp_path)
    # Analysis never raises for an unrelated change request; it just yields
    # few/no confident patch candidates.
    assert session["status"] in {"REVIEW_IN_PROGRESS", "REVIEW_COMPLETE"}


# --- apply_decisions -----------------------------------------------------------------


def test_apply_decisions_requires_workflow(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    sid = session["session_id"]
    content = _small_docx_bytes(tmp_path)
    pv2.upload_documents(sid, [("MDTM.docx", content, "traceability")], root=tmp_path)
    pv2.resolve_identity(sid, routing_mode="explicit", root=tmp_path)
    with pytest.raises(PilotOrchestratorError) as exc:
        pv2.apply_decisions(sid, [], root=tmp_path)
    assert exc.value.reason_code in {"INVALID_SESSION_STATE", "NO_WORKFLOW"}


def test_apply_decisions_rejects_invalid_decision_value(tmp_path):
    scenario = scenarios.get_scenario("pilot_gr_schedule_table")
    fixture_path = scenarios.ensure_fixture(scenario)
    session = pv2.create_session(
        document_set=scenario.document_set, change_request=scenario.change_request, scenario_id=scenario.scenario_id, root=tmp_path
    )
    sid = session["session_id"]
    pv2.upload_documents(sid, [(fixture_path.name, fixture_path.read_bytes(), scenario.fixture_role)], root=tmp_path)
    pv2.resolve_identity(sid, routing_mode="explicit", root=tmp_path)
    session = pv2.analyze(sid, root=tmp_path)
    items = session["review_items"]
    if not items:
        pytest.skip("scenario produced no review items to exercise decision validation")
    with pytest.raises(PilotOrchestratorError) as exc:
        pv2.apply_decisions(sid, [{"item_id": items[0]["item_id"], "decision": "NOT_A_REAL_DECISION"}], root=tmp_path)
    assert exc.value.reason_code == "INVALID_DECISION"


# --- get_review_items / get_result / download_artifact / list_sessions ---------------


def test_get_review_items_matches_session_state(tmp_path):
    scenario = scenarios.get_scenario("pilot_ec_req_single")
    fixture_path = scenarios.ensure_fixture(scenario)
    session = pv2.create_session(
        document_set=scenario.document_set, change_request=scenario.change_request, scenario_id=scenario.scenario_id, root=tmp_path
    )
    sid = session["session_id"]
    pv2.upload_documents(sid, [(fixture_path.name, fixture_path.read_bytes(), scenario.fixture_role)], root=tmp_path)
    pv2.resolve_identity(sid, routing_mode="explicit", root=tmp_path)
    session = pv2.analyze(sid, root=tmp_path)
    items = pv2.get_review_items(sid, root=tmp_path)
    assert items == session["review_items"]


def test_get_result_includes_session_and_root(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    sid = session["session_id"]
    result = pv2.get_result(sid, root=tmp_path)
    assert result["session"]["session_id"] == sid
    assert "session_root" in result


def test_download_artifact_resolves_valid_path(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    sid = session["session_id"]
    from document_ai.pilot_v2 import session_store

    session_root = session_store.session_dir(sid, root=tmp_path)
    session_store.write_trace(session_root, "output", "note", {"a": 1})
    resolved = pv2.download_artifact(sid, "output/note.json", root=tmp_path)
    assert resolved.is_file()


def test_download_artifact_blocks_traversal(tmp_path):
    session = pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    sid = session["session_id"]
    from document_ai.pilot_v2 import security

    with pytest.raises(security.PilotSecurityError):
        pv2.download_artifact(sid, "../../etc/passwd", root=tmp_path)


def test_list_sessions_returns_registered_sessions(tmp_path):
    pv2.create_session(document_set="ec_sw", change_request="x", root=tmp_path)
    pv2.create_session(document_set="general_report", change_request="y", root=tmp_path)
    rows = pv2.list_sessions(root=tmp_path)
    assert len(rows) == 2
