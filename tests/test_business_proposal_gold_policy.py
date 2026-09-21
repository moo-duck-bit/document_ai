# -*- coding: utf-8 -*-
"""Deterministic policy tests (Cycle 6): REQUIRED/AMBIGUOUS/OPTIONAL/NOT_APPLICABLE."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from document_ai.evaluation.business_proposal_gold.document_inventory import inventory_document
from document_ai.evaluation.business_proposal_gold.policy import (
    classify_case,
    resolve_target_concept,
    select_section_representative,
)


def _build_base(path: Path) -> None:
    doc = Document()
    doc.add_heading("사업 제안서", level=0)
    doc.add_heading("실행 일정", level=1)
    doc.add_paragraph("본 사업의 실행 일정에 대한 설명입니다.")
    doc.add_paragraph("2026-Q3 착수, 2026-Q4 완료")
    doc.add_heading("예산", level=1)
    doc.add_paragraph("본 사업의 예산에 대한 설명입니다.")
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "금액"
    table.rows[1].cells[0].text = "인건비"
    table.rows[1].cells[1].text = "100"
    table.rows[2].cells[0].text = "장비"
    table.rows[2].cells[1].text = "50"
    doc.add_heading("위험 관리", level=1)
    doc.add_paragraph("위험 요소와 대응 방안을 기술합니다.")
    doc.save(str(path))


def _build_dup_schedule(path: Path) -> None:
    doc = Document()
    doc.add_heading("사업 제안서", level=0)
    doc.add_heading("실행 일정", level=1)
    doc.add_paragraph("2026-Q3 착수")
    doc.add_heading("예산", level=1)
    doc.add_paragraph("100만원")
    doc.add_heading("추진 일정", level=1)
    doc.add_paragraph("2027-Q1 마무리")
    doc.save(str(path))


def _inv(path: Path, document_id: str = "PROPOSAL_BASE"):
    return inventory_document(path, document_id=document_id)


def test_resolve_target_concept_from_tags():
    assert resolve_target_concept("아무 문구", ["schedule"]) == "SCHEDULE"
    assert resolve_target_concept("아무 문구", ["budget"]) == "BUDGET"
    assert resolve_target_concept("아무 문구", ["risk"]) == "RISK_MANAGEMENT"


def test_resolve_target_concept_from_text_when_no_tags():
    assert resolve_target_concept("예산 표 수정", []) == "BUDGET"


def test_resolve_target_concept_none_for_unrelated_text():
    assert resolve_target_concept("표지 로고 위치 변경", []) is None


def test_required_schedule_paragraph_no_table(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    row = classify_case(
        case_id="c_sched",
        document_id="PROPOSAL_BASE",
        change_request="실행 일정 2026-Q4 조정",
        tags=["schedule"],
        inventory=inv,
    )
    assert row.node_evaluation_mode == "REQUIRED"
    assert row.expected_physical_node_type == "PARAGRAPH"
    assert row.primary_reference["document_node_id"] is not None
    assert row.primary_reference["template_node_id"] == "business_proposal_v1.schedule"


def test_required_budget_resolves_table(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    row = classify_case(
        case_id="c_budget",
        document_id="PROPOSAL_BASE",
        change_request="예산 표 인건비 수정",
        tags=["budget", "table"],
        inventory=inv,
    )
    assert row.node_evaluation_mode == "REQUIRED"
    assert row.expected_physical_node_type == "TABLE"
    assert row.primary_reference["document_node_id"].startswith("table_")


def test_not_applicable_for_no_impact_tag(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    row = classify_case(
        case_id="c_no_impact",
        document_id="PROPOSAL_BASE",
        change_request="표지 로고 위치 변경",
        tags=["no_impact"],
        inventory=inv,
    )
    assert row.node_evaluation_mode == "NOT_APPLICABLE"
    assert row.primary_reference is None


def test_optional_for_unmapped_concept(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    row = classify_case(
        case_id="c_style",
        document_id="PROPOSAL_BASE",
        change_request="문서 전반 표현 개선",
        tags=["style_overall"],
        inventory=inv,
    )
    assert row.node_evaluation_mode == "OPTIONAL"
    assert row.expected_physical_node_type == "NONE"


def test_optional_for_add_missing_section(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    row = classify_case(
        case_id="c_market",
        document_id="PROPOSAL_BASE",
        change_request="시장분석 절을 추가",
        tags=["market_add"],
        inventory=inv,
    )
    assert row.node_evaluation_mode == "OPTIONAL"
    assert row.primary_reference["document_node_id"] is None


def test_ambiguous_add_onto_existing_section(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    row = classify_case(
        case_id="c_risk_add",
        document_id="PROPOSAL_BASE",
        change_request="위험 관리 항목 추가",
        tags=["risk"],
        inventory=inv,
    )
    assert row.node_evaluation_mode == "AMBIGUOUS"
    assert len(row.acceptable_groups) == 1
    assert len(row.acceptable_groups[0]) >= 2


def test_ambiguous_duplicate_sections(tmp_path):
    path = tmp_path / "dup.docx"
    _build_dup_schedule(path)
    inv = _inv(path, document_id="PROPOSAL_DUP")
    row = classify_case(
        case_id="c_dup",
        document_id="PROPOSAL_DUP",
        change_request="일정 섹션 중복 정리",
        tags=["duplicate_schedule"],
        inventory=inv,
    )
    assert row.node_evaluation_mode == "AMBIGUOUS"
    assert len(row.acceptable_groups[0]) >= 2


def test_force_ambiguous_tag_overrides_unique_section(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    row = classify_case(
        case_id="c_sched_section",
        document_id="PROPOSAL_BASE",
        change_request="일정 섹션 전반 수정",
        tags=["schedule_section"],
        inventory=inv,
    )
    assert row.node_evaluation_mode == "AMBIGUOUS"


def test_select_section_representative_prefers_table(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    node_index = {n["node_id"]: n for n in inv["nodes"]}
    budget_section = next(s for s in inv["sections"] if "BUDGET" in s["concepts"])
    node_id, score = select_section_representative(
        budget_section, node_index, target_concept="BUDGET", prefer_table=True
    )
    assert node_id is not None
    assert node_id.startswith("table_")
    assert score > 0


def test_classify_case_rationale_non_empty(tmp_path):
    path = tmp_path / "base.docx"
    _build_base(path)
    inv = _inv(path)
    row = classify_case(
        case_id="c_x",
        document_id="PROPOSAL_BASE",
        change_request="예산 표 인건비 수정",
        tags=["budget"],
        inventory=inv,
    )
    assert row.label_rationale.strip() != ""
    assert row.label_source == "document_structure_and_change_request"
