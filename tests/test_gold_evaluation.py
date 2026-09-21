"""Tests for independent gold dataset and human evaluation framework."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from document_ai.eval.bootstrap_gold import bootstrap_gold_from_case
from document_ai.eval.gold_dataset import (
    load_gold_manifest,
    resolve_independent_gold,
)
from document_ai.eval.gold_fields import (
    build_gold_fields,
    compare_against_gold_fields,
    load_gold_fields,
    save_gold_fields,
)
from document_ai.eval.human_review import (
    build_human_review_items,
    render_human_review_checklist_md,
    write_human_review_checklist,
)
from document_ai.eval.harness_benchmark import run_harness_benchmark
from document_ai.validation.runner import run_document_validation


def _minimal_docs(tmp_path: Path) -> Path:
    case = tmp_path / "gold_case"
    case.mkdir()
    (case / "input.json").write_text(
        json.dumps(
            {
                "case_id": "gold_case",
                "domain": "inventory_b2b",
                "product_name": "IMS",
                "facts": {"product_name": "IMS"},
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    mdsr = Document()
    mdsr.add_paragraph("1. Introduction")
    table = mdsr.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "제품명"
    table.rows[0].cells[1].text = "IMS"
    req = mdsr.add_table(rows=2, cols=2)
    req.rows[0].cells[0].text = "Req. 1"
    req.rows[1].cells[0].text = "설명"
    req.rows[1].cells[1].text = "Track warehouse stock levels."
    trace = mdsr.add_table(rows=1, cols=4)
    trace.rows[0].cells[0].text = "IA-01"
    trace.rows[0].cells[1].text = "Auth"
    trace.rows[0].cells[2].text = "해당"
    trace.rows[0].cells[3].text = "Req. 1"
    mdsr.save(str(case / "output_mdsr.docx"))

    mddr = Document()
    mddr.add_paragraph("5. Design Description")
    mddr.add_paragraph("Req. 1\nInventoryService locks stock rows for reservations.")
    mddr.save(str(case / "output_mddr.docx"))
    return case


def test_build_and_compare_gold_fields(tmp_path: Path):
    case = _minimal_docs(tmp_path)
    fields = build_gold_fields(
        case_id="gold_case",
        mdsr_path=case / "output_mdsr.docx",
        mddr_path=case / "output_mddr.docx",
        product_name="IMS",
        domain="inventory_b2b",
    )
    assert "Req. 1" in fields["requirement_ids"]
    assert "Req. 1" in fields["design_ids"]
    assert fields["product_name"] == "IMS"
    assert fields["traceability_rows"]

    path = tmp_path / "fields.json"
    save_gold_fields(fields, path)
    loaded = load_gold_fields(path)
    assert loaded["case_id"] == "gold_case"

    metrics = compare_against_gold_fields(
        generated_mdsr=case / "output_mdsr.docx",
        generated_mddr=case / "output_mddr.docx",
        gold_fields=fields,
    )
    assert metrics["semantic_requirement_id_coverage"] == 1.0
    assert metrics["semantic_design_id_coverage"] == 1.0
    assert metrics["semantic_overall"] >= 90


def test_bootstrap_gold_and_manifest(tmp_path: Path):
    case = _minimal_docs(tmp_path)
    gold_root = tmp_path / "gold"
    result = bootstrap_gold_from_case(
        case,
        split="holdout",
        approve=True,
        root=gold_root,
    )
    assert result["status"] == "human_approved"
    assert Path(result["copied"]["mdsr"]).exists()
    assert Path(result["copied"]["mddr"]).exists()
    assert Path(result["copied"]["fields"]).exists()

    manifest = load_gold_manifest(gold_root)
    assert any(c["case_id"] == "gold_case" and c["split"] == "holdout" for c in manifest["cases"])
    resolved = resolve_independent_gold("gold_case", root=gold_root)
    assert resolved["mdsr"] and resolved["fields"]


def test_human_review_checklist_generator(tmp_path: Path):
    result = {
        "case": "demo",
        "product_name": "IMS",
        "domain": "inventory_b2b",
        "scores": {"overall": 90, "mdsr": 92, "mddr": 88},
        "metrics": {"residual_placeholder_count": 1},
        "semantic": {
            "semantic_overall": 95,
            "semantic_requirement_id_coverage": 1.0,
            "missing_requirement_ids": [],
        },
        "high_risk_differences": [],
    }
    items = build_human_review_items(result)
    categories = {i["category"] for i in items}
    assert "figure" in categories
    assert "table" in categories
    assert "terminology" in categories
    assert "regulation" in categories
    assert "traceability" in categories
    assert "placeholders" in categories

    md = render_human_review_checklist_md(result)
    assert "Human Review Checklist" in md
    path = write_human_review_checklist(result, tmp_path / "human_review_checklist.md")
    assert path.exists()
    assert "figure" in path.read_text(encoding="utf-8").lower()


def test_validation_writes_human_review_and_semantic(tmp_path: Path):
    case = _minimal_docs(tmp_path)
    # seed gold as same outputs
    gold_mdsr = case / "gold_mdsr.docx"
    gold_mddr = case / "gold_mddr.docx"
    gold_mdsr.write_bytes((case / "output_mdsr.docx").read_bytes())
    gold_mddr.write_bytes((case / "output_mddr.docx").read_bytes())
    fields = build_gold_fields(
        case_id="gold_case",
        mdsr_path=gold_mdsr,
        mddr_path=gold_mddr,
        product_name="IMS",
    )
    fields_path = case / "gold_fields.json"
    save_gold_fields(fields, fields_path)

    result = run_document_validation(
        case,
        gold_mdsr=gold_mdsr,
        gold_mddr=gold_mddr,
        gold_fields_path=fields_path,
    )
    assert result["semantic"]["semantic_overall"] >= 90
    assert result["scores"].get("semantic") is not None
    checklist = Path(result["report_paths"]["human_review_checklist"])
    assert checklist.exists()
    assert "Review Items" in checklist.read_text(encoding="utf-8")


def test_holdout_benchmark_filter(tmp_path: Path):
    config = {
        "version": "1.0",
        "cases": [
            {
                "case_id": "inventory_mgmt",
                "case_dir": "data/cases/inventory_mgmt",
                "split": "train",
                "skip_generate": True,
            },
            {
                "case_id": "hospital_reservation",
                "case_dir": "data/cases/hospital_reservation",
                "split": "holdout",
                "skip_generate": True,
            },
        ],
    }
    config_path = tmp_path / "bench.json"
    config_path.write_text(json.dumps(config), encoding="utf-8")
    if not Path("data/cases/hospital_reservation/output_mdsr.docx").exists():
        pytest.skip("hospital_reservation outputs required")

    report = run_harness_benchmark(
        config_path=config_path,
        out_dir=tmp_path / "out",
        holdout_only=True,
    )
    assert report["cases_run"] == 1
    assert report["case_results"][0]["case_id"] == "hospital_reservation"
    assert report["holdout"]["cases_run"] == 1
    assert report["train"]["cases_run"] == 0
