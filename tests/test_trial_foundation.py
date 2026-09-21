"""Tests for Evaluation Foundation / Trial framework."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from document_ai.trial.baseline import freeze_baseline
from document_ai.trial.check_input import check_trial_input
from document_ai.trial.dataset import (
    default_dataset_manifest,
    validate_dataset_manifest,
)
from document_ai.trial.diff import compare_docx_pair, compare_token_bags
from document_ai.trial.generate import generate_trial
from document_ai.trial.init import init_trial
from document_ai.trial.leakage import (
    build_input_inventory,
    check_dataset_role_conflicts,
    path_looks_like_leakage,
)
from document_ai.trial.metrics import time_saving_rate
from document_ai.trial.models import TRIAL_STATUSES, human_revision_template
from document_ai.trial.paths import read_json, write_json
from document_ai.trial.review import prepare_trial_review
from document_ai.trial.analyze import analyze_trial
from document_ai.trial.summary import summarize_trial


def _write_docx(path: Path, paragraphs: list[str], table_rows: list[list[str]] | None = None) -> None:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, cell in enumerate(row):
                table.rows[i].cells[j].text = cell
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def test_trial_manifest_status_enum():
    assert "prepared" in TRIAL_STATUSES
    assert "completed" in TRIAL_STATUSES


def test_invalid_trial_type_rejected(tmp_path: Path):
    case = tmp_path / "case"
    case.mkdir()
    (case / "input.json").write_text("{}", encoding="utf-8")
    with pytest.raises(ValueError):
        init_trial(
            trial_id="bad-type",
            case=case,
            trial_type="not_a_type",
            trials_root=tmp_path / "trials",
        )


def test_frozen_tag_commit_recorded(tmp_path: Path):
    case = tmp_path / "lab_ec_sw"
    case.mkdir()
    (case / "input.json").write_text("{}", encoding="utf-8")
    result = init_trial(
        trial_id="trial-freeze-check",
        case=case,
        system_version="v0.5-document-harness",
        system_commit="d73fc18",
        trials_root=tmp_path / "trials",
        synthetic=True,
    )
    manifest = result["manifest"]
    assert manifest["system_version"] == "v0.5-document-harness"
    assert manifest["system_commit"] == "d73fc18"
    assert "git" in manifest["frozen"]
    assert "commit" in manifest["frozen"]["git"]


def test_input_hash_and_gold_leakage(tmp_path: Path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    good = input_dir / "input.json"
    good.write_text('{"case_id":"x"}', encoding="utf-8")
    inv = build_input_inventory(input_dir)
    assert inv["file_count"] == 1
    assert inv["files"][0]["sha256"]
    assert inv["data_leakage_clean"] is True

    leak = input_dir / "gold_mdsr.docx"
    leak.write_bytes(b"not-a-real-docx")
    inv2 = build_input_inventory(input_dir)
    assert inv2["data_leakage_clean"] is False
    assert any("gold" in w.lower() for w in inv2["leakage_warnings"])


def test_path_leakage_markers():
    assert path_looks_like_leakage("data/gold/mdsr/x.docx")
    assert path_looks_like_leakage(r"C:\proj\human_revised\final.docx")
    assert not path_looks_like_leakage("data/cases/lab_ec_sw/input.json")


def test_human_revised_leakage_before_review(tmp_path: Path):
    case = tmp_path / "lab_ec_sw"
    case.mkdir()
    (case / "input.json").write_text("{}", encoding="utf-8")
    result = init_trial(
        trial_id="trial-leak-hr",
        case=case,
        trials_root=tmp_path / "trials",
        synthetic=True,
    )
    trial = Path(tmp_path / "trials" / "trial-leak-hr")
    # Put a premature revised doc
    _write_docx(trial / "human_revised" / "revised_mdsr.docx", ["secret final answer"])
    check = check_trial_input(trial)
    assert check["ok"] is False
    assert any("human_revised" in f for f in check["findings"])


def test_generated_output_isolation(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    case = tmp_path / "lab_ec_sw"
    case.mkdir()
    (case / "input.json").write_text('{"facts":{}}', encoding="utf-8")
    _write_docx(case / "output_mdsr.docx", ["Req. 1 original"])
    _write_docx(case / "output_mddr.docx", ["Design original"])
    original_hash = (case / "output_mdsr.docx").read_bytes()

    init_trial(
        trial_id="trial-iso",
        case=case,
        trials_root=tmp_path / "trials",
        synthetic=True,
    )
    trial = tmp_path / "trials" / "trial-iso"
    # Remove premature human_revised findings: only .gitkeep
    result = generate_trial(trial, skip_generate=True)
    assert result["source_case_preserved"] is True
    assert (trial / "generated" / "output_mdsr.docx").exists()
    assert (case / "output_mdsr.docx").read_bytes() == original_hash
    assert result["generation_manifest"]["overwrite_source_case"] is False


def test_time_saving_na_without_baseline():
    result = time_saving_rate(
        manual_baseline_minutes=None,
        generation_minutes=5,
        revision_minutes=10,
    )
    assert result["status"] == "N/A"
    assert result["time_saving_rate"] is None


def test_time_saving_calculation():
    result = time_saving_rate(
        manual_baseline_minutes=100,
        generation_minutes=10,
        revision_minutes=20,
    )
    assert result["status"] == "ok"
    assert result["time_saving_rate"] == 0.7


def test_docx_paragraph_and_table_diff(tmp_path: Path):
    gen = tmp_path / "gen.docx"
    rev = tmp_path / "rev.docx"
    _write_docx(
        gen,
        ["Alpha", "Beta", "Gamma"],
        [["ID", "Value"], ["Req. 1", "old"]],
    )
    _write_docx(
        rev,
        ["Alpha", "Beta changed", "Delta"],
        [["ID", "Value"], ["Req. 1", "new"]],
    )
    bags = compare_token_bags(["a", "b", "c"], ["a", "b", "d"])
    assert bags["unchanged_count"] == 2
    assert bags["modified_count"] + bags["added_count"] + bags["deleted_count"] >= 1
    pair = compare_docx_pair(gen, rev)
    assert "unchanged_paragraph_ratio" in pair
    assert "modified_table_cell_ratio" in pair
    assert pair["document_level_edit_burden_score"] >= 0


def test_error_taxonomy_and_unverified_auto(tmp_path: Path):
    case = tmp_path / "lab_ec_sw"
    case.mkdir()
    (case / "input.json").write_text("{}", encoding="utf-8")
    _write_docx(case / "output_mdsr.docx", ["Hello world"])
    _write_docx(case / "output_mddr.docx", ["Design"])
    init_trial(
        trial_id="trial-err",
        case=case,
        trials_root=tmp_path / "trials",
        synthetic=True,
    )
    trial = tmp_path / "trials" / "trial-err"
    generate_trial(trial, skip_generate=True)
    prepare_trial_review(trial)
    # Provide revised with changes so auto candidates appear
    _write_docx(trial / "human_revised" / "revised_mdsr.docx", ["Hello world", "Extra section"])
    _write_docx(trial / "human_revised" / "revised_mddr.docx", ["Design"])
    analyzed = analyze_trial(trial)
    ann_path = trial / "error_annotations.json"
    data = read_json(ann_path)
    assert data["annotations"]
    assert any(not a.get("verified") for a in data["annotations"])
    assert analyzed["ok"]


def test_presentation_and_csv_reports(tmp_path: Path):
    case = tmp_path / "lab_ec_sw"
    case.mkdir()
    (case / "input.json").write_text("{}", encoding="utf-8")
    _write_docx(case / "output_mdsr.docx", ["A"])
    _write_docx(case / "output_mddr.docx", ["B"])
    init_trial(
        trial_id="trial-pres",
        case=case,
        trials_root=tmp_path / "trials",
        synthetic=True,
    )
    trial = tmp_path / "trials" / "trial-pres"
    generate_trial(trial, skip_generate=True)
    prepare_trial_review(trial)
    analyze_trial(trial)
    summary = summarize_trial(trial)
    assert summary["completion_label"] == "synthetic_framework_validation"
    assert (trial / "presentation_summary.md").exists()
    assert (trial / "trial_report.md").exists()
    assert (trial / "reports" / "tables" / "system_baseline.csv").exists()
    assert (trial / "reports" / "figures" / "data" / "score_comparison.json").exists()
    # Must not claim real completion
    metrics = read_json(trial / "metrics.json")
    assert metrics["real_world_trial_complete"] is False


def test_dataset_role_conflict_detection():
    conflicts = check_dataset_role_conflicts(
        [
            {"case_id": "hospital_reservation", "role": "holdout"},
            {"case_id": "hospital_reservation", "role": "development"},
        ]
    )
    assert conflicts
    ok = validate_dataset_manifest(default_dataset_manifest())
    assert ok["ok"] is True


def test_holdout_role_in_default_manifest():
    manifest = default_dataset_manifest()
    holdout = [c for c in manifest["cases"] if c["case_id"] == "hospital_reservation"]
    assert holdout and holdout[0]["role"] == "holdout"


def test_baseline_freeze_no_case_overwrite(tmp_path: Path):
    out = tmp_path / "baselines" / "v0.5-document-harness"
    # Create minimal sources if missing in sandbox — freeze copies if exist
    result = freeze_baseline(out_root=out, run_pytest=False)
    assert result["ok"] is True
    assert (out / "baseline_report.md").exists()
    assert (out / "git_info.json").exists()
    assert (out / "environment.json").exists()


def test_human_revision_template_fields():
    template = human_revision_template("trial-x", ["MDSR", "MDDR"])
    assert template["manual_baseline_minutes"] is None
    assert len(template["document_reviews"]) == 2
    assert "content_accuracy" in template["document_reviews"][0]["rating"]


def test_review_template_generation(tmp_path: Path):
    case = tmp_path / "lab_ec_sw"
    case.mkdir()
    (case / "input.json").write_text("{}", encoding="utf-8")
    _write_docx(case / "output_mdsr.docx", ["A"])
    _write_docx(case / "output_mddr.docx", ["B"])
    init_trial(
        trial_id="trial-rev",
        case=case,
        trials_root=tmp_path / "trials",
        synthetic=True,
    )
    trial = tmp_path / "trials" / "trial-rev"
    generate_trial(trial, skip_generate=True)
    prep = prepare_trial_review(trial)
    assert prep["ok"]
    assert (trial / "review" / "reviewer_instructions.md").exists()
    assert (trial / "review" / "review_checklist.md").exists()
    assert (trial / "review" / "human_revision_record.template.json").exists()
