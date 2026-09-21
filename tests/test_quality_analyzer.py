"""Tests for document quality analyzer."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from document_ai.quality.analyzer import analyze_document, analyze_case


def _write_mdsr(path: Path, *, body: str, product: str = "Test Product") -> None:
    doc = Document()
    doc.add_paragraph(f"EC-SW-MDSR(TP) {product}")
    doc.add_paragraph(body)
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "Req. 1"
    table.rows[1].cells[0].text = "설명"
    table.rows[1].cells[1].text = body
    table.rows[2].cells[0].text = "목적"
    table.rows[2].cells[1].text = "purpose"
    table.rows[3].cells[0].text = "기준"
    table.rows[3].cells[1].text = "criteria"
    doc.save(str(path))


def test_analyzer_detects_placeholder(tmp_path: Path):
    doc_path = tmp_path / "out.docx"
    _write_mdsr(doc_path, body="Document number XX-XX-XXXX pending review.")
    analysis = analyze_document(
        doc_path,
        template="spec_requirements",
        domain="inventory_b2b",
        product_name="Inventory Management System",
        requirements=[{"req_id": "Req. 1"}],
    )
    categories = {f.category for f in analysis.findings}
    assert "placeholder" in categories


def test_analyzer_detects_mindrium_residual(tmp_path: Path):
    doc_path = tmp_path / "out.docx"
    _write_mdsr(doc_path, body="Android 12 and Python 3.7 training data module.")
    analysis = analyze_document(
        doc_path,
        template="spec_requirements",
        domain="inventory_b2b",
        product_name="Inventory Management System",
        requirements=[{"req_id": "Req. 1"}],
    )
    assert any(f.category == "mindrium" for f in analysis.findings)


def test_analyzer_detects_domain_mismatch(tmp_path: Path):
    doc_path = tmp_path / "out.docx"
    _write_mdsr(
        doc_path,
        body="B2C 온라인 쇼핑몰에서 장바구니와 PG 결제를 제공한다.",
        product="Inventory Management System",
    )
    analysis = analyze_document(
        doc_path,
        template="spec_requirements",
        domain="inventory_b2b",
        product_name="Inventory Management System",
        requirements=[{"req_id": "Req. 1"}],
    )
    assert any(f.category == "domain_mismatch" for f in analysis.findings)


def test_analyzer_detects_duplicate_paragraph(tmp_path: Path):
    doc_path = tmp_path / "out.docx"
    doc = Document()
    repeated = "이 문단은 재고관리 시스템의 운영 정책을 설명하는 중복 본문입니다. " * 2
    doc.add_paragraph(repeated)
    doc.add_paragraph(repeated)
    doc.save(str(doc_path))
    analysis = analyze_document(doc_path, template="spec_requirements")
    assert any(f.category == "duplicate_paragraph" for f in analysis.findings)


def test_analyze_case_reads_input(tmp_path: Path):
    case_dir = tmp_path / "case"
    case_dir.mkdir()
    (case_dir / "input.json").write_text(
        json.dumps(
            {
                "case_id": "case",
                "domain": "inventory_b2b",
                "product_name": "IMS",
                "templates_in_set": ["spec_requirements"],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (case_dir / "requirements.json").write_text(
        json.dumps({"requirements": [{"req_id": "Req. 1"}]}, ensure_ascii=False),
        encoding="utf-8",
    )
    _write_mdsr(case_dir / "output_mdsr.docx", body="IMS 재고관리 본문", product="IMS")
    result = analyze_case(case_dir)
    assert result["domain"] == "inventory_b2b"
    assert "mdsr" in result["documents"]
