from docx import Document
from docx.table import Table

from document_ai.render.mdsr import (
    apply_document_layout,
    apply_mdsr_content,
    apply_operating_principle,
    apply_traceability_details,
)
from document_ai.render.requirements import _fill_req_table_full, fill_traceability_matrix


def _req_table(*rows: tuple[str, str]) -> Table:
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=2)
    for ri, (a, b) in enumerate(rows):
        table.rows[ri].cells[0].text = a
        table.rows[ri].cells[1].text = b
    return table


def test_fill_req_table_labeled_rows():
    table = _req_table(
        ("Req. 105", ""),
        ("설명", ""),
        ("목적", ""),
        ("기준", ""),
    )
    req = {
        "description": "desc text",
        "purpose": "purpose text",
        "criteria": "criteria text",
    }
    assert _fill_req_table_full(table, req, overwrite=True)
    assert table.rows[1].cells[1].text == "desc text"
    assert table.rows[2].cells[1].text == "purpose text"
    assert table.rows[3].cells[1].text == "criteria text"


def test_fill_req_table_unlabeled_functional_gets_labels():
    table = _req_table(
        ("Req. 2", ""),
        ("", ""),
        ("", ""),
        ("", ""),
    )
    req = {
        "description": "main desc",
        "purpose": "main purpose",
        "criteria": "main criteria",
    }
    assert _fill_req_table_full(table, req, overwrite=True)
    assert table.rows[1].cells[0].text == "설명"
    assert table.rows[1].cells[1].text == "main desc"
    assert table.rows[2].cells[0].text == "목적"
    assert table.rows[2].cells[1].text == "main purpose"
    assert table.rows[3].cells[0].text == "기준"
    assert table.rows[3].cells[1].text == "main criteria"


def test_traceability_fills_title_and_applicability():
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "IA-01"
    table.rows[1].cells[0].text = "UC-02"
    doc._body._body.insert(-1, table._tbl)

    trace = [
        {"requirement": "IA-01", "title": "사용자 식별 및 인증", "linked_reqs": "Req. 2", "applicability": "해당"},
        {"requirement": "UC-02", "title": "모바일 코드 사용 통제", "linked_reqs": "N/A", "applicability": "비해당"},
    ]
    count = apply_traceability_details(doc, trace)
    assert count >= 2
    assert "사용자" in table.rows[0].cells[0].text
    assert table.rows[0].cells[2].text == "해당"
    assert table.rows[0].cells[3].text == "Req. 2"
    assert table.rows[1].cells[2].text == "비해당"


def test_fill_req_table_applies_labels_for_high_numbered_req():
    table = _req_table(
        ("Req. 100", ""),
        ("", ""),
        ("", ""),
        ("", ""),
    )
    req = {
        "description": "desc for 100",
        "purpose": "purpose for 100",
        "criteria": "criteria for 100",
    }
    assert _fill_req_table_full(table, req, overwrite=True)
    assert table.rows[1].cells[0].text == "설명"
    assert table.rows[1].cells[1].text == "desc for 100"
    assert table.rows[2].cells[0].text == "목적"
    assert table.rows[2].cells[1].text == "purpose for 100"
    assert table.rows[3].cells[0].text == "기준"
    assert table.rows[3].cells[1].text == "criteria for 100"


def test_apply_document_layout_replaces_standards_and_cover_line():
    doc = Document()
    doc.add_paragraph("Software Requirement Specification")
    doc.add_paragraph("")
    doc.add_paragraph("IEC 62304:2015 CSV Medical device software")
    doc.add_paragraph("ISO 14971:2019 Medical devices")
    doc.add_paragraph("XX-XX-XXXX)")
    doc.add_paragraph("")
    doc.add_paragraph("")

    content = {
        "cover_product_line": "JM COLLECTION Co., Ltd.",
        "standards": ["PCI DSS v4.0", "개인정보보호법 / OWASP ASVS"],
        "document_title": "EC-SW-MDSR(JM) JM COLLECTION",
        "document_version_line": "Ver 1.0 / JM-SRS-2026-001",
        "narrative": ["Overview paragraph one."],
    }
    stats = apply_document_layout(doc, content)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    assert stats["cover_product_line"] == 1
    assert stats["standards_filled"] == 2
    assert "JM COLLECTION Co., Ltd." in paragraphs
    assert "PCI DSS v4.0" in paragraphs
    assert "Overview paragraph one." in paragraphs


def test_apply_operating_principle_replaces_residual_and_clears_placeholder():
    from docx.shared import RGBColor
    from docx.text.paragraph import Paragraph

    from document_ai.learn.docx_io import iter_blocks
    from document_ai.render.mdsr import _find_operating_principle_slots

    doc = Document()
    doc.add_paragraph("ok")
    doc.add_paragraph("범불안장애 환자의 증상 개선을 위한 환자용 어플리케이션")
    doc.add_paragraph("범불안장애 인지행동치료법을 모바일 앱으로 구현")
    placeholder = doc.add_paragraph(".")
    placeholder.runs[0].font.color.rgb = RGBColor(0x00, 0x70, 0xC0)

    slots = _find_operating_principle_slots(doc)
    assert len(slots) == 3
    filled = apply_operating_principle(
        doc,
        ["JM COLLECTION 동작원리 본문", "JM COLLECTION 거래 흐름", ""],
        slot_indices=slots,
    )
    assert filled == 2
    blocks = list(iter_blocks(doc))
    assert isinstance(blocks[slots[0]], Paragraph)
    assert "JM COLLECTION 동작원리 본문" in blocks[slots[0]].text
    assert blocks[slots[2]].text == ""
