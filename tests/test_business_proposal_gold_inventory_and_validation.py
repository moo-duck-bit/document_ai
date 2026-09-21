# -*- coding: utf-8 -*-
"""DOCX inventory scanning + gold row validation tests (Cycle 6)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from document_ai.evaluation.business_proposal_gold.document_inventory import (
    find_matching_sections,
    inventory_document,
    nodes_by_id,
)
from document_ai.evaluation.business_proposal_gold.schema import BusinessProposalGoldRow
from document_ai.evaluation.business_proposal_gold.validation import (
    validate_change_request_table_alignment,
    validate_gold_row,
    validate_gold_rows,
)


def _build_boilerplate_doc(path: Path) -> None:
    doc = Document()
    doc.add_heading("일정", level=1)
    doc.add_paragraph("일정에 대한 설명입니다.")
    doc.add_heading("예산", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "금액"
    table.rows[1].cells[0].text = "인건비"
    table.rows[1].cells[1].text = "100"
    doc.save(str(path))


def test_inventory_node_id_conventions(tmp_path):
    path = tmp_path / "doc.docx"
    _build_boilerplate_doc(path)
    inv = inventory_document(path, document_id="D1")
    headings = [n for n in inv["nodes"] if n["node_type"] == "HEADING"]
    paragraphs = [n for n in inv["nodes"] if n["node_type"] == "PARAGRAPH"]
    tables = [n for n in inv["nodes"] if n["node_type"] == "TABLE"]
    assert all(n["node_id"].startswith("heading_") for n in headings)
    assert all(n["node_id"].startswith("paragraph_") for n in paragraphs)
    assert all(n["node_id"].startswith("table_") for n in tables)
    assert tables[0]["node_id"] == "table_00"


def test_inventory_boilerplate_detection(tmp_path):
    path = tmp_path / "doc.docx"
    _build_boilerplate_doc(path)
    inv = inventory_document(path, document_id="D1")
    para = next(n for n in inv["nodes"] if n["node_type"] == "PARAGRAPH")
    assert para["is_boilerplate"] is True


def test_inventory_table_row_col_counts(tmp_path):
    path = tmp_path / "doc.docx"
    _build_boilerplate_doc(path)
    inv = inventory_document(path, document_id="D1")
    table = next(n for n in inv["nodes"] if n["node_type"] == "TABLE")
    assert table["row_count"] == 2
    assert table["column_count"] == 2


def test_inventory_table_concepts_include_budget(tmp_path):
    path = tmp_path / "doc.docx"
    _build_boilerplate_doc(path)
    inv = inventory_document(path, document_id="D1")
    table = next(n for n in inv["nodes"] if n["node_type"] == "TABLE")
    assert "BUDGET" in table["concepts"]


def test_inventory_sections_group_members(tmp_path):
    path = tmp_path / "doc.docx"
    _build_boilerplate_doc(path)
    inv = inventory_document(path, document_id="D1")
    assert len(inv["sections"]) == 2
    sched_section = next(s for s in inv["sections"] if "SCHEDULE" in s["concepts"])
    assert "paragraph_0001" in sched_section["member_node_ids"] or any(
        "paragraph" in m for m in sched_section["member_node_ids"]
    )


def test_find_matching_sections(tmp_path):
    path = tmp_path / "doc.docx"
    _build_boilerplate_doc(path)
    inv = inventory_document(path, document_id="D1")
    matches = find_matching_sections(inv, "BUDGET")
    assert len(matches) == 1
    assert find_matching_sections(inv, "") == []


def test_nodes_by_id(tmp_path):
    path = tmp_path / "doc.docx"
    _build_boilerplate_doc(path)
    inv = inventory_document(path, document_id="D1")
    idx = nodes_by_id(inv)
    assert "table_00" in idx
    assert idx["table_00"]["node_type"] == "TABLE"


def test_validate_required_missing_primary_rejected():
    row = BusinessProposalGoldRow(
        case_id="c1",
        document_id="D1",
        node_evaluation_mode="REQUIRED",
        primary_reference=None,
        label_rationale="no target",
    )
    issues = validate_gold_row(row)
    assert "REQUIRED_MISSING_PRIMARY" in issues


def test_validate_required_virtual_only_rejected():
    row = BusinessProposalGoldRow(
        case_id="c1",
        document_id="D1",
        node_evaluation_mode="REQUIRED",
        primary_reference={"reference_type": "VIRTUAL"},
        label_rationale="virtual only",
    )
    issues = validate_gold_row(row)
    assert "REQUIRED_VIRTUAL_ONLY_REJECTED" in issues


def test_validate_ambiguous_missing_group_rejected():
    row = BusinessProposalGoldRow(
        case_id="c1",
        document_id="D1",
        node_evaluation_mode="AMBIGUOUS",
        primary_reference={"reference_type": "TEMPLATE", "template_node_id": "x"},
        acceptable_groups=[],
        acceptable_references=[],
        label_rationale="ambiguous but no group",
    )
    issues = validate_gold_row(row)
    assert "AMBIGUOUS_MISSING_GROUP" in issues


def test_validate_rationale_required():
    row = BusinessProposalGoldRow(
        case_id="c1", document_id="D1", node_evaluation_mode="OPTIONAL", label_rationale=""
    )
    issues = validate_gold_row(row)
    assert "RATIONALE_REQUIRED" in issues


def test_validate_rationale_rejects_prediction_artifact_reference():
    row = BusinessProposalGoldRow(
        case_id="c1",
        document_id="D1",
        node_evaluation_mode="OPTIONAL",
        label_rationale="derived from node_ranking output top-1",
    )
    issues = validate_gold_row(row)
    assert any("RATIONALE_REFERENCES_PREDICTION_ARTIFACT" in i for i in issues)


def test_validate_valid_required_row_no_issues():
    row = BusinessProposalGoldRow(
        case_id="c1",
        document_id="D1",
        node_evaluation_mode="REQUIRED",
        primary_reference={
            "reference_type": "TEMPLATE",
            "template_node_id": "business_proposal_v1.budget",
            "document_node_id": "table_00",
        },
        expected_physical_node_type="TABLE",
        expected_location_type="TABLE",
        label_rationale="unique matching table node",
    )
    issues = validate_gold_row(row)
    assert issues == []


def test_validate_change_request_table_alignment_literal_table_word_rejected():
    row = BusinessProposalGoldRow(
        case_id="c1",
        document_id="D1",
        node_evaluation_mode="REQUIRED",
        expected_physical_node_type="PARAGRAPH",
        label_rationale="ok",
    )
    issues = validate_change_request_table_alignment(row, change_request="일정 표를 수정")
    assert "TABLE_INTENT_PARAGRAPH_ONLY_REJECTED" in issues


def test_validate_change_request_table_alignment_no_literal_word_no_issue():
    row = BusinessProposalGoldRow(
        case_id="c1",
        document_id="D1",
        node_evaluation_mode="REQUIRED",
        expected_physical_node_type="PARAGRAPH",
        label_rationale="ok",
    )
    # "SCHEDULE" concept implies table_intent broadly, but no literal 표/table wording here —
    # must NOT be flagged (this is the exact regression this cross-check must avoid).
    issues = validate_change_request_table_alignment(row, change_request="실행 일정 2026-Q4 조정")
    assert issues == []


def test_validate_gold_rows_aggregate_reports_invalid_cases():
    good = BusinessProposalGoldRow(
        case_id="good",
        document_id="D1",
        node_evaluation_mode="NOT_APPLICABLE",
        label_rationale="off topic",
    )
    bad = BusinessProposalGoldRow(
        case_id="bad",
        document_id="D1",
        node_evaluation_mode="REQUIRED",
        primary_reference=None,
        label_rationale="missing primary",
    )
    report = validate_gold_rows([good, bad])
    assert report["n_rows"] == 2
    assert report["n_invalid"] == 1
    assert "bad" in report["issues_by_case"]
    assert report["status"] == "INVALID"


def test_validate_gold_rows_all_valid_status():
    good = BusinessProposalGoldRow(
        case_id="good",
        document_id="D1",
        node_evaluation_mode="NOT_APPLICABLE",
        label_rationale="off topic",
    )
    report = validate_gold_rows([good])
    assert report["status"] == "VALID"
    assert report["n_invalid"] == 0
