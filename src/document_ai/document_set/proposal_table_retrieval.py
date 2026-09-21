# -*- coding: utf-8 -*-
"""Business Proposal table REVIEW retrieval (schedule/budget/org/kpi; PATCH forbidden)."""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table

from document_ai.domain_packs.business_proposal.concepts import (
    BUDGET,
    KPI,
    ORGANIZATION,
    SCHEDULE,
    TABLE,
    normalize_proposal_concepts,
)
from document_ai.domain_packs.business_proposal.structural_roles import classify_table_header_role
from document_ai.template.concept_normalization import (
    EXECUTION_PLAN,
    MILESTONE,
    TIMELINE,
    extract_time_unit_hits,
    normalize_concepts,
    tokenize,
)

SCHEDULE_FAMILY = frozenset({SCHEDULE, TIMELINE, MILESTONE, EXECUTION_PLAN})


@dataclass
class ProposalTableEvidence:
    document_id: str
    node_id: str
    table_id: str
    table_role: str
    header_terms: list[str]
    canonical_concepts: list[str]
    time_unit_hits: list[str]
    schedule_term_hits: list[str]
    row_count: int
    column_count: int
    section_context: str
    score: float
    reason_codes: list[str] = field(default_factory=list)
    supports_review: bool = False
    supports_patch: bool = False
    source_stage: str = "proposal_table_retrieval"
    score_components: dict[str, float] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _table_headers(table: Table) -> list[str]:
    if not table.rows:
        return []
    return [(c.text or "").strip() for c in table.rows[0].cells]


def _table_blob(table: Table) -> str:
    parts = []
    for row in table.rows:
        for cell in row.cells:
            t = (cell.text or "").strip()
            if t:
                parts.append(t)
    return "\n".join(parts)


def _cr_wants_table(cr_concepts: set[str], primary: str | None = None) -> bool:
    return bool(
        cr_concepts
        & (SCHEDULE_FAMILY | {BUDGET, ORGANIZATION, KPI, TABLE})
        or primary in {SCHEDULE, BUDGET, ORGANIZATION, KPI}
    )


def collect_proposal_table_evidence(
    *,
    change_request: str,
    document_id: str,
    docx_path: Path,
    template_section_hits: list[dict[str, Any]] | None = None,
) -> list[ProposalTableEvidence]:
    cr = change_request or ""
    cr_concepts = normalize_proposal_concepts(cr) | normalize_concepts(cr)
    proposal_primary = next(
        (c for c in (SCHEDULE, BUDGET, ORGANIZATION, KPI) if c in cr_concepts),
        None,
    )
    cr_time = extract_time_unit_hits(cr)
    if not _cr_wants_table(cr_concepts, proposal_primary) and not cr_time:
        return []

    path = Path(docx_path)
    if not path.is_file():
        return []

    try:
        document = Document(str(path))
    except Exception:
        return []

    evidences: list[ProposalTableEvidence] = []

    for i, p in enumerate(document.paragraphs):
        text = (p.text or "").strip()
        if not text:
            continue
        concepts = normalize_proposal_concepts(text) | normalize_concepts(text)
        time_hits = extract_time_unit_hits(text)
        score = 0.0
        reasons = ["PATCH_FORBIDDEN_PROPOSAL_TABLE"]
        comps: dict[str, float] = {}
        table_role = classify_table_header_role(text) or "PARAGRAPH"
        if concepts & SCHEDULE_FAMILY:
            score += 0.5
            comps["heading_schedule_concept"] = 0.5
            reasons.append("schedule_concept_heading")
            table_role = "SCHEDULE_TABLE"
        if BUDGET in concepts:
            score += 0.45
            comps["heading_budget_concept"] = 0.45
            table_role = "BUDGET_TABLE"
        if ORGANIZATION in concepts:
            score += 0.4
            comps["heading_org_concept"] = 0.4
            table_role = "ORGANIZATION_SECTION"
        if KPI in concepts:
            score += 0.4
            comps["heading_kpi_concept"] = 0.4
            table_role = "KPI_TABLE"
        if cr_concepts & concepts:
            score += 0.15
            comps["cr_concept_overlap"] = 0.15
        if time_hits and (concepts & SCHEDULE_FAMILY or proposal_primary == SCHEDULE):
            score += 0.2
            comps["time_unit"] = 0.2
        if score < 0.4:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        is_heading = "heading" in style.lower() or len(text) <= 40
        node_id = f"heading_{i:04d}" if is_heading else f"paragraph_{i:04d}"
        evidences.append(
            ProposalTableEvidence(
                document_id=document_id,
                node_id=node_id,
                table_id="",
                table_role=table_role,
                header_terms=[],
                canonical_concepts=sorted(concepts),
                time_unit_hits=time_hits,
                schedule_term_hits=sorted(concepts & SCHEDULE_FAMILY),
                row_count=0,
                column_count=0,
                section_context=text[:120],
                score=round(min(score, 1.0), 4),
                reason_codes=sorted(set(reasons)),
                supports_review=True,
                supports_patch=False,
                score_components=comps,
            )
        )

    for ti, table in enumerate(document.tables):
        headers = _table_headers(table)
        blob = _table_blob(table)
        header_join = " ".join(headers)
        concepts = normalize_proposal_concepts(header_join + "\n" + blob[:800]) | normalize_concepts(
            header_join + "\n" + blob[:800]
        )
        table_role = classify_table_header_role(header_join, blob=blob[:400]) or TABLE
        time_hits = extract_time_unit_hits(blob) + extract_time_unit_hits(header_join)
        time_hits = list(dict.fromkeys(time_hits))
        score = 0.0
        reasons = ["PATCH_FORBIDDEN_PROPOSAL_TABLE"]
        comps: dict[str, float] = {}

        if table_role == "SCHEDULE_TABLE" or concepts & SCHEDULE_FAMILY:
            score += 0.45
            comps["schedule_table_header"] = 0.45
            table_role = "SCHEDULE_TABLE"
            reasons.append("schedule_table_structure")
        if table_role == "BUDGET_TABLE" or BUDGET in concepts:
            score += 0.45
            comps["budget_table_header"] = 0.45
            table_role = "BUDGET_TABLE"
            reasons.append("budget_table_structure")
        if table_role == "ORGANIZATION_SECTION" or ORGANIZATION in concepts:
            score += 0.4
            comps["org_table_header"] = 0.4
            table_role = "ORGANIZATION_SECTION"
            reasons.append("organization_table_structure")
        if table_role == "KPI_TABLE" or KPI in concepts:
            score += 0.4
            comps["kpi_table_header"] = 0.4
            table_role = "KPI_TABLE"
            reasons.append("kpi_table_structure")
        if cr_concepts & concepts:
            score += 0.15
            comps["cr_overlap"] = 0.15
        if time_hits and (proposal_primary == SCHEDULE or concepts & SCHEDULE_FAMILY):
            score += 0.25
            comps["time_unit_in_table"] = 0.25
        if proposal_primary and proposal_primary in concepts:
            score += 0.2
            comps["primary_concept_table"] = 0.2
        if score < 0.42:
            continue
        nrows = len(table.rows)
        ncols = len(table.columns) if table.rows else 0
        evidences.append(
            ProposalTableEvidence(
                document_id=document_id,
                node_id=f"table_{ti:02d}",
                table_id=f"table_{ti:02d}",
                table_role=table_role,
                header_terms=headers,
                canonical_concepts=sorted(concepts),
                time_unit_hits=time_hits,
                schedule_term_hits=sorted(concepts & SCHEDULE_FAMILY),
                row_count=nrows,
                column_count=ncols,
                section_context=header_join[:120] or blob[:120],
                score=round(min(score, 1.0), 4),
                reason_codes=sorted(set(reasons)),
                supports_review=True,
                supports_patch=False,
                score_components=comps,
            )
        )

    for hit in template_section_hits or []:
        if hit.get("status") != "REVIEW_REQUIRED":
            continue
        sec = str(hit.get("node_id") or "")
        sec_low = sec.lower()
        role = None
        if "schedule" in sec_low:
            role = "SCHEDULE_TABLE"
        elif "budget" in sec_low:
            role = "BUDGET_TABLE"
        elif "organization" in sec_low:
            role = "ORGANIZATION_SECTION"
        elif "expected_outcomes" in sec_low or "kpi" in sec_low:
            role = "KPI_TABLE"
        if not role:
            continue
        evidences.append(
            ProposalTableEvidence(
                document_id=document_id,
                node_id=sec,
                table_id="",
                table_role=role,
                header_terms=[],
                canonical_concepts=sorted(normalize_proposal_concepts(sec)),
                time_unit_hits=cr_time,
                schedule_term_hits=[SCHEDULE] if role == "SCHEDULE_TABLE" else [],
                row_count=0,
                column_count=0,
                section_context=str(hit.get("display_name") or sec),
                score=0.32,
                reason_codes=[
                    "template_section_hint",
                    "template_only_not_document_grounded",
                    "PATCH_FORBIDDEN_PROPOSAL_TABLE",
                ],
                supports_review=False,
                supports_patch=False,
                score_components={"template_section": 0.32},
            )
        )

    evidences.sort(key=lambda e: (-e.score, e.document_id, e.node_id))
    seen: set[str] = set()
    uniq: list[ProposalTableEvidence] = []
    for e in evidences:
        if e.node_id in seen:
            continue
        seen.add(e.node_id)
        uniq.append(e)
    if len([e for e in uniq if e.supports_review]) > 1:
        for e in uniq:
            if e.supports_review and "proposal_table_ambiguous" not in e.reason_codes:
                e.reason_codes.append("proposal_table_ambiguous")
    return uniq


def evidences_to_review_items(
    evidences: list[ProposalTableEvidence],
    *,
    limit: int = 12,
) -> list[dict[str, Any]]:
    items = []
    for e in evidences:
        if not e.supports_review:
            continue
        items.append(
            {
                "item_id": f"BP-TBL-{e.node_id}",
                "candidate_id": f"BP-TBL-{e.node_id}",
                "document_id": e.document_id,
                "node_id": e.node_id,
                "status": "REVIEW_REQUIRED",
                "overlap": e.score,
                "display_name": e.section_context,
                "reason_codes": e.reason_codes,
                "human_review_required": True,
                "source_stage": "proposal_table_retrieval",
                "metadata": {
                    "overlap": e.score,
                    "evidence_type": "TABLE_STRUCTURE" if e.table_id else "TABLE_CONTEXT",
                    "supports_patch": False,
                    "supports_review": True,
                    "table_role": e.table_role,
                    "canonical_concepts": e.canonical_concepts,
                },
            }
        )
        if len(items) >= limit:
            break
    return items


def write_proposal_table_artifacts(
    out_dir: Path,
    *,
    change_request: str,
    evidences: list[ProposalTableEvidence],
) -> dict[str, str]:
    out_dir.mkdir(parents=True, exist_ok=True)
    concepts = sorted(normalize_proposal_concepts(change_request) | normalize_concepts(change_request))
    summary = {
        "query_concepts": concepts,
        "time_units": extract_time_unit_hits(change_request),
        "review_candidates": sum(1 for e in evidences if e.supports_review),
        "patch_from_proposal_table": 0,
        "table_roles": sorted({e.table_role for e in evidences if e.supports_review}),
    }
    validation = {
        "ok": all(not e.supports_patch for e in evidences),
        "issues": [],
        "invariants": [
            "no_table_only_patch",
            "proposal_table_candidate_has_structure_evidence",
            "review_requires_substantive_evidence",
        ],
    }
    files = {
        "business_proposal_table_analysis.json": {
            "change_request": change_request,
            "concepts": concepts,
            "tokens": sorted(tokenize(change_request)),
            "evidences": [e.to_dict() for e in evidences],
            "review_items": evidences_to_review_items(evidences),
            "summary": summary,
            "validation": validation,
        },
    }
    written = {}
    for name, payload in files.items():
        path = out_dir / name
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        written[name] = str(path).replace("\\", "/")
    return written
