# -*- coding: utf-8 -*-
"""General Report schedule / table REVIEW retrieval (PATCH forbidden)."""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table

from document_ai.template.concept_normalization import (
    EXECUTION_PLAN,
    MILESTONE,
    SCHEDULE,
    TABLE,
    TIMELINE,
    extract_time_unit_hits,
    normalize_concepts,
    tokenize,
)

SCHEDULE_FAMILY = frozenset({SCHEDULE, TIMELINE, MILESTONE, EXECUTION_PLAN})


@dataclass
class TableStructureEvidence:
    document_id: str
    node_id: str
    table_id: str
    header_terms: list[str]
    canonical_concepts: list[str]
    time_unit_hits: list[str]
    schedule_term_hits: list[str]
    row_count: int
    column_count: int
    section_context: str
    score: float
    reason_codes: list[str] = field(default_factory=list)
    supports_review: bool = False
    supports_patch: bool = False
    source_stage: str = "schedule_table_retrieval"
    score_components: dict[str, float] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _table_headers(table: Table) -> list[str]:
    if not table.rows:
        return []
    return [(c.text or "").strip() for c in table.rows[0].cells]


def _table_blob(table: Table) -> str:
    parts = []
    for row in table.rows:
        for cell in row.cells:
            t = (cell.text or "").strip()
            if t:
                parts.append(t)
    return "\n".join(parts)


def collect_schedule_table_evidence(
    *,
    change_request: str,
    document_id: str,
    docx_path: Path,
    template_section_hits: list[dict[str, Any]] | None = None,
) -> list[TableStructureEvidence]:
    cr = change_request or ""
    cr_concepts = normalize_concepts(cr)
    cr_wants_schedule = bool(cr_concepts & SCHEDULE_FAMILY) or bool(cr_concepts & {TABLE})
    cr_time = extract_time_unit_hits(cr)
    if not cr_wants_schedule and not cr_time:
        return []

    path = Path(docx_path)
    if not path.is_file():
        return []

    try:
        document = Document(str(path))
    except Exception:
        return []

    evidences: list[TableStructureEvidence] = []
    # Headings / paragraphs
    for i, p in enumerate(document.paragraphs):
        text = (p.text or "").strip()
        if not text:
            continue
        concepts = normalize_concepts(text)
        time_hits = extract_time_unit_hits(text)
        schedule_hits = sorted(concepts & SCHEDULE_FAMILY)
        score = 0.0
        reasons = ["PATCH_FORBIDDEN_TABLE_OR_SCHEDULE_ONLY"]
        comps: dict[str, float] = {}
        if concepts & SCHEDULE_FAMILY:
            score += 0.55
            comps["heading_schedule_concept"] = 0.55
            reasons.append("schedule_concept_heading")
        if TABLE in concepts:
            score += 0.25
            comps["table_concept"] = 0.25
            reasons.append("table_concept")
        if time_hits and (concepts & SCHEDULE_FAMILY or cr_wants_schedule):
            score += 0.25
            comps["time_unit"] = 0.25
            reasons.append("time_unit_evidence")
        # CR concept alignment
        if cr_concepts & concepts:
            score += 0.15
            comps["cr_concept_overlap"] = 0.15
        if score < 0.4:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        is_heading = "heading" in style.lower() or len(text) <= 40
        node_id = f"heading_{i:04d}" if is_heading else f"paragraph_{i:04d}"
        evidences.append(
            TableStructureEvidence(
                document_id=document_id,
                node_id=node_id,
                table_id="",
                header_terms=[],
                canonical_concepts=sorted(concepts),
                time_unit_hits=time_hits,
                schedule_term_hits=schedule_hits,
                row_count=0,
                column_count=0,
                section_context=text[:120],
                score=round(min(score, 1.0), 4),
                reason_codes=sorted(set(reasons)),
                supports_review=True,
                supports_patch=False,
                score_components=comps,
            )
        )

    # Tables
    for ti, table in enumerate(document.tables):
        headers = _table_headers(table)
        blob = _table_blob(table)
        header_join = " ".join(headers)
        concepts = normalize_concepts(header_join + "\n" + blob[:500])
        time_hits = extract_time_unit_hits(blob) + extract_time_unit_hits(header_join)
        time_hits = list(dict.fromkeys(time_hits))
        schedule_hits = sorted(concepts & SCHEDULE_FAMILY)
        score = 0.0
        reasons = ["PATCH_FORBIDDEN_TABLE_OR_SCHEDULE_ONLY"]
        comps = {}
        if concepts & SCHEDULE_FAMILY:
            score += 0.45
            comps["table_header_schedule"] = 0.45
            reasons.append("schedule_concept_table_header")
        if TABLE in cr_concepts or TABLE in concepts:
            score += 0.20
            comps["table_presence"] = 0.20
            reasons.append("table_structure")
        if time_hits and (cr_wants_schedule or concepts & SCHEDULE_FAMILY):
            score += 0.30
            comps["time_unit_in_table"] = 0.30
            reasons.append("time_unit_table")
        if cr_concepts & concepts:
            score += 0.15
            comps["cr_overlap"] = 0.15
        # schedule CR + any table with time units
        if cr_wants_schedule and time_hits:
            score += 0.20
            comps["schedule_cr_plus_time_table"] = 0.20
        if score < 0.45:
            continue
        nrows = len(table.rows)
        ncols = len(table.columns) if table.rows else 0
        evidences.append(
            TableStructureEvidence(
                document_id=document_id,
                node_id=f"table_{ti:02d}",
                table_id=f"table_{ti:02d}",
                header_terms=headers,
                canonical_concepts=sorted(concepts),
                time_unit_hits=time_hits,
                schedule_term_hits=schedule_hits,
                row_count=nrows,
                column_count=ncols,
                section_context=header_join[:120],
                score=round(min(score, 1.0), 4),
                reason_codes=sorted(set(reasons)),
                supports_review=True,
                supports_patch=False,
                score_components=comps,
            )
        )

    # Template section hint — keep as weak hint only; NOT document-grounded evidence
    for hit in template_section_hits or []:
        if hit.get("status") != "REVIEW_REQUIRED":
            continue
        sec = str(hit.get("node_id") or "")
        if "schedule" not in sec and "일정" not in str(hit.get("display_name") or ""):
            continue
        evidences.append(
            TableStructureEvidence(
                document_id=document_id,
                node_id=sec,
                table_id="",
                header_terms=[],
                canonical_concepts=[SCHEDULE],
                time_unit_hits=cr_time,
                schedule_term_hits=[SCHEDULE],
                row_count=0,
                column_count=0,
                section_context=str(hit.get("display_name") or "schedule"),
                score=0.35,
                reason_codes=[
                    "template_schedule_section",
                    "template_only_not_document_grounded",
                    "PATCH_FORBIDDEN_TABLE_OR_SCHEDULE_ONLY",
                ],
                supports_review=False,  # do not promote document REVIEW alone
                supports_patch=False,
                score_components={"template_section": 0.35},
            )
        )

    evidences.sort(key=lambda e: (-e.score, e.document_id, e.node_id))
    # Dedupe by node_id
    seen: set[str] = set()
    uniq: list[TableStructureEvidence] = []
    for e in evidences:
        if e.node_id in seen:
            continue
        seen.add(e.node_id)
        uniq.append(e)
    if len([e for e in uniq if e.supports_review]) > 1:
        for e in uniq:
            if e.supports_review and "SEMANTIC_AMBIGUOUS" not in e.reason_codes:
                e.reason_codes.append("schedule_table_ambiguous")
    return uniq


def evidences_to_review_items(evidences: list[TableStructureEvidence], *, limit: int = 10) -> list[dict[str, Any]]:
    items = []
    for e in evidences:
        if not e.supports_review:
            continue
        items.append(
            {
                "item_id": f"SCH-{e.node_id}",
                "candidate_id": f"SCH-{e.node_id}",
                "document_id": e.document_id,
                "node_id": e.node_id,
                "status": "REVIEW_REQUIRED",
                "overlap": e.score,
                "display_name": e.section_context,
                "reason_codes": e.reason_codes,
                "human_review_required": True,
                "metadata": {
                    "overlap": e.score,
                    "evidence_type": "TABLE_STRUCTURE" if e.table_id else "SCHEDULE_HEADING",
                    "supports_patch": False,
                    "canonical_concepts": e.canonical_concepts,
                },
            }
        )
        if len(items) >= limit:
            break
    return items


def write_general_report_retrieval_artifacts(
    out_dir: Path,
    *,
    change_request: str,
    evidences: list[TableStructureEvidence],
) -> dict[str, str]:
    out_dir.mkdir(parents=True, exist_ok=True)
    concepts = sorted(normalize_concepts(change_request))
    summary = {
        "query_concepts": concepts,
        "time_units": extract_time_unit_hits(change_request),
        "review_candidates": sum(1 for e in evidences if e.supports_review),
        "patch_from_schedule_table": 0,
    }
    validation = {
        "ok": all(not e.supports_patch for e in evidences),
        "issues": [],
        "invariants": [
            "no_table_only_patch",
            "schedule_table_candidate_has_structure_evidence",
            "review_requires_substantive_evidence",
        ],
    }
    files = {
        "general_report_concept_analysis.json": {
            "change_request": change_request,
            "concepts": concepts,
            "tokens": sorted(tokenize(change_request)),
        },
        "general_report_table_candidates.json": [e.to_dict() for e in evidences if e.table_id],
        "general_report_schedule_review.json": evidences_to_review_items(evidences),
        "general_report_retrieval_summary.json": summary,
        "general_report_retrieval_validation.json": validation,
    }
    written = {}
    for name, payload in files.items():
        path = out_dir / name
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        written[name] = str(path).replace("\\", "/")
    return written
