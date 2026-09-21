# -*- coding: utf-8 -*-
"""PR-20: DOCX logical structure → DocumentModel (read-only, no write)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from document_ai.document_parser.structure import (
    DocumentModel,
    ListModel,
    ParagraphModel,
    SectionModel,
    TableModel,
)

_HEADING_STYLE = re.compile(r"heading\s*(\d+)", re.IGNORECASE)


def _slug(text: str, order: int) -> str:
    raw = re.sub(r"[^\w가-힣]+", "_", (text or "").strip().lower())
    raw = re.sub(r"_+", "_", raw).strip("_") or "section"
    return f"sec_{order:03d}_{raw[:48]}"


def _heading_level(paragraph: Any) -> int | None:
    style_name = ""
    try:
        style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    except Exception:  # noqa: BLE001
        style_name = ""
    m = _HEADING_STYLE.match(style_name.strip())
    if m:
        return int(m.group(1))
    # outline level fallback
    try:
        pPr = paragraph._element.pPr  # noqa: SLF001
        if pPr is not None and pPr.outlineLvl is not None:
            return int(pPr.outlineLvl.val) + 1
    except Exception:  # noqa: BLE001
        pass
    return None


def _is_list_paragraph(paragraph: Any) -> tuple[bool, bool]:
    """Return (is_list, ordered)."""
    try:
        style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "").lower()
        if "list bullet" in style_name or style_name.startswith("list paragraph"):
            return True, False
        if "list number" in style_name:
            return True, True
        pPr = paragraph._element.pPr  # noqa: SLF001
        if pPr is None or pPr.numPr is None:
            return False, False
        return True, False
    except Exception:  # noqa: BLE001
        return False, False


def parse_docx_file(
    path: str | Path,
    *,
    document_id: str | None = None,
    document_type: str = "unknown",
    title: str | None = None,
    metadata: dict[str, Any] | None = None,
) -> DocumentModel:
    """Parse DOCX logical structure (headings/paragraphs/tables/lists). Read-only."""
    from docx import Document  # lazy import
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    p = Path(path)
    doc = Document(str(p))
    sections: list[SectionModel] = []
    stack: list[tuple[int, str]] = []
    current: SectionModel | None = None
    order = 0
    para_i = 0
    table_i = 0
    list_i = 0
    doc_title = title or ""
    pending_list: ListModel | None = None

    def _flush_list() -> None:
        nonlocal pending_list
        if pending_list is not None and current is not None:
            current.lists.append(pending_list)
            pending_list = None

    def _ensure_root() -> SectionModel:
        nonlocal current, order
        if current is not None:
            return current
        order += 1
        sid = _slug("document", order)
        current = SectionModel(
            section_id=sid,
            heading=doc_title or p.stem,
            heading_level=0,
            parent=None,
            order=order,
        )
        sections.append(current)
        stack.clear()
        stack.append((0, sid))
        return current

    # Walk body elements in order
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag == "p":
            paragraph = DocxParagraph(child, doc)
            text = (paragraph.text or "").strip()
            level = _heading_level(paragraph)
            if level is not None and text:
                _flush_list()
                if level == 1 and not doc_title:
                    doc_title = text
                order += 1
                sid = _slug(text, order)
                while stack and stack[-1][0] >= level:
                    stack.pop()
                parent = stack[-1][1] if stack else None
                if parent:
                    for s in sections:
                        if s.section_id == parent and sid not in s.children:
                            s.children.append(sid)
                sec = SectionModel(
                    section_id=sid,
                    heading=text,
                    heading_level=level,
                    parent=parent,
                    order=order,
                )
                sections.append(sec)
                stack.append((level, sid))
                current = sec
                continue

            if not text:
                continue
            is_list, ordered = _is_list_paragraph(paragraph)
            cur = _ensure_root()
            if is_list:
                if pending_list is None or pending_list.ordered != ordered:
                    _flush_list()
                    list_i += 1
                    pending_list = ListModel(
                        list_id=f"lst_{list_i:03d}",
                        items=[text],
                        ordered=ordered,
                        order=list_i,
                    )
                else:
                    pending_list.items.append(text)
            else:
                _flush_list()
                para_i += 1
                cur.paragraphs.append(
                    ParagraphModel(
                        paragraph_id=f"p_{para_i:03d}",
                        text=text,
                        order=para_i,
                    )
                )
        elif tag == "tbl":
            _flush_list()
            cur = _ensure_root()
            table = DocxTable(child, doc)
            rows_data: list[list[str]] = []
            for row in table.rows:
                rows_data.append([(c.text or "").strip() for c in row.cells])
            headers = rows_data[0] if rows_data else []
            body_rows = rows_data[1:] if len(rows_data) > 1 else []
            table_i += 1
            cur.tables.append(
                TableModel(
                    table_id=f"tbl_{table_i:03d}",
                    headers=headers,
                    rows=body_rows,
                    order=table_i,
                )
            )

    _flush_list()
    return DocumentModel(
        document_id=document_id or p.stem,
        document_type=document_type,
        title=doc_title or p.stem,
        source_format="docx",
        metadata={**(metadata or {}), "source_path": str(p), "read_only": True},
        sections=sections,
    )
