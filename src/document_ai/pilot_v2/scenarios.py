# -*- coding: utf-8 -*-
"""Real User Document Pilot v2 — 12 fixed scenario definitions (4 EC-SW, 4 GR, 4 BP).

Scenario metadata is the JSON file under
``data/pilot/real_user_document_pilot/scenarios/scenarios.json`` (single source of
truth); this module loads/validates it into :class:`PilotScenario` and lazily builds
small, pilot-owned fixture ``.docx`` files.

Important: this module NEVER reads or writes anything under
``data/eval/document_set_benchmark_v2`` (frozen benchmark data) or ``data/examples``.
The benchmark v2 fixtures directory is regenerated on demand by
``scripts/generate_document_set_benchmark_v2.py`` (which also rewrites the sealed
holdout manifest), so touching it here would risk violating the
"do not overwrite immutable benchmark runs" constraint. Instead this module builds
independent pilot fixtures, structurally similar to the benchmark ones, under the
pilot's own data directory — scenario IDs are pilot-only (``pilot_*``), never
benchmark/holdout case IDs.
"""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from docx import Document

REPO = Path(__file__).resolve().parents[3]
PILOT_ROOT = REPO / "data" / "pilot" / "real_user_document_pilot"
SCENARIOS_DIR = PILOT_ROOT / "scenarios"
SCENARIOS_JSON = SCENARIOS_DIR / "scenarios.json"
FIXTURES_ROOT = SCENARIOS_DIR / "fixtures"

ALLOWED_DOMAINS: frozenset[str] = frozenset({"ec_sw", "general_report", "business_proposal"})


class ScenarioError(ValueError):
    pass


@dataclass
class PilotScenario:
    scenario_id: str
    domain: str
    document_set: str
    title: str
    change_request: str
    fixture_relative_path: str
    fixture_builder: str
    fixture_role: str
    writer_enabled_hint: bool = True
    expected_domain: str = ""
    failure_conditions: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    def fixture_path(self) -> Path:
        return SCENARIOS_DIR / self.fixture_relative_path


def _mdtm_fixture(path: Path) -> None:
    doc = Document()
    doc.add_heading("Traceability Matrix (Pilot)", level=1)
    rows = [
        ("Req. 10", "5.1.1", "TC-10"),
        ("Req. 11", "5.1.2", "TC-11"),
        ("Req. 12", "5.1.3", "TC-12"),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=4)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Requirement", "Design", "Test", "Note"
    for i, (req, des, test) in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = req, des, test, ""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _report_fixture(path: Path, *, schedule_table: bool = False) -> None:
    doc = Document()
    for text in ("요약", "방법론", "결과", "결론"):
        doc.add_heading(text, level=1)
        doc.add_paragraph(f"{text} 본문 내용입니다.")
    if schedule_table:
        doc.add_heading("일정", level=1)
        t = doc.add_table(rows=3, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "단계", "기간"
        t.rows[1].cells[0].text, t.rows[1].cells[1].text = "분석", "2026-Q3"
        t.rows[2].cells[0].text, t.rows[2].cells[1].text = "구현", "2026-Q4"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _proposal_fixture(path: Path) -> None:
    doc = Document()
    doc.add_heading("사업 제안서 (Pilot)", level=0)
    for section in ("실행 일정", "예산", "위험 관리", "수행 조직", "기대 효과"):
        doc.add_heading(section, level=1)
        doc.add_paragraph(f"{section}에 대한 설명입니다.")
        if section == "예산":
            t = doc.add_table(rows=3, cols=2)
            t.rows[0].cells[0].text, t.rows[0].cells[1].text = "항목", "금액"
            t.rows[1].cells[0].text, t.rows[1].cells[1].text = "인건비", "100"
            t.rows[2].cells[0].text, t.rows[2].cells[1].text = "장비", "50"
        if section == "실행 일정":
            doc.add_paragraph("2026-Q3 착수, 2026-Q4 완료")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


FIXTURE_BUILDERS = {
    "mdtm_pilot_base": _mdtm_fixture,
    "report_pilot_base": lambda p: _report_fixture(p, schedule_table=False),
    "report_pilot_schedule": lambda p: _report_fixture(p, schedule_table=True),
    "proposal_pilot_base": _proposal_fixture,
}


def _default_scenarios() -> list[dict[str, Any]]:
    def ec(sid: str, cr: str, title: str, failure: list[str] | None = None) -> dict[str, Any]:
        return {
            "scenario_id": sid,
            "domain": "ec_sw",
            "document_set": "ec_sw",
            "title": title,
            "change_request": cr,
            "fixture_relative_path": "fixtures/ec_sw/mdtm_pilot_base.docx",
            "fixture_builder": "mdtm_pilot_base",
            "fixture_role": "traceability",
            "writer_enabled_hint": not failure,
            "expected_domain": "ec_sw",
            "failure_conditions": failure or [],
        }

    def gr(sid: str, cr: str, title: str, *, schedule: bool = False, failure: list[str] | None = None) -> dict[str, Any]:
        fixture = "report_pilot_schedule" if schedule else "report_pilot_base"
        return {
            "scenario_id": sid,
            "domain": "general_report",
            "document_set": "general_report",
            "title": title,
            "change_request": cr,
            "fixture_relative_path": f"fixtures/general_report/{fixture}.docx",
            "fixture_builder": fixture,
            "fixture_role": "general_report",
            "writer_enabled_hint": not failure,
            "expected_domain": "general_report",
            "failure_conditions": failure or [],
        }

    def bp(sid: str, cr: str, title: str, failure: list[str] | None = None) -> dict[str, Any]:
        return {
            "scenario_id": sid,
            "domain": "business_proposal",
            "document_set": "business_proposal",
            "title": title,
            "change_request": cr,
            "fixture_relative_path": "fixtures/business_proposal/proposal_pilot_base.docx",
            "fixture_builder": "proposal_pilot_base",
            "fixture_role": "custom",
            "writer_enabled_hint": not failure,
            "expected_domain": "business_proposal",
            "failure_conditions": failure or [],
        }

    return [
        ec("pilot_ec_req_single", "Req. 11 추적성 행 갱신 요청", "EC-SW 단일 요구사항 변경"),
        ec("pilot_ec_design_review", "설계 ID 5.1.2 관련 검토 요청", "EC-SW 설계 항목 검토"),
        ec(
            "pilot_ec_no_impact",
            "사내 행사 안내 문구 변경",
            "EC-SW 영향 없음 케이스",
            failure=["no_impact"],
        ),
        ec(
            "pilot_ec_semantic",
            "사용자 인증 정책 강화 필요",
            "EC-SW 의미 기반 검토",
            failure=["semantic_only"],
        ),
        gr("pilot_gr_methodology", "방법론 데이터 출처 보완", "일반보고서 방법론 보완"),
        gr(
            "pilot_gr_schedule_table",
            "일정 표 2026-Q4 수정",
            "일반보고서 일정 표 수정",
            schedule=True,
        ),
        gr(
            "pilot_gr_no_impact",
            "표지 디자인 색상 변경",
            "일반보고서 영향 없음 케이스",
            failure=["no_impact"],
        ),
        gr("pilot_gr_conclusion", "결론 권고사항 수정", "일반보고서 결론 수정"),
        bp("pilot_bp_schedule", "실행 일정 2026-Q4 조정", "사업제안서 일정 조정"),
        bp("pilot_bp_budget", "예산 표 인건비 수정", "사업제안서 예산 수정"),
        bp("pilot_bp_risk", "위험 관리 항목 추가", "사업제안서 위험 관리 추가"),
        bp(
            "pilot_bp_no_impact",
            "표지 로고 위치 변경",
            "사업제안서 영향 없음 케이스",
            failure=["no_impact"],
        ),
    ]


def ensure_scenarios_json(*, force: bool = False) -> Path:
    """Write scenarios.json if missing (never overwrites unless force=True)."""
    if SCENARIOS_JSON.is_file() and not force:
        return SCENARIOS_JSON
    payload = {
        "meta": {
            "name": "real_user_document_pilot_scenarios",
            "version": "1.0",
            "count": 12,
            "domains": sorted(ALLOWED_DOMAINS),
            "note": "Pilot-owned fixtures; independent of frozen benchmark data.",
        },
        "scenarios": _default_scenarios(),
    }
    SCENARIOS_JSON.parent.mkdir(parents=True, exist_ok=True)
    SCENARIOS_JSON.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return SCENARIOS_JSON


def load_scenarios(*, scenarios_json: Path | None = None) -> list[PilotScenario]:
    path = scenarios_json or ensure_scenarios_json()
    if not path.is_file():
        raise ScenarioError(f"missing_scenarios_json:{path}")
    raw = json.loads(path.read_text(encoding="utf-8"))
    rows = raw.get("scenarios") or []
    out: list[PilotScenario] = []
    seen: set[str] = set()
    for row in rows:
        sid = row.get("scenario_id")
        if not sid:
            raise ScenarioError("scenario_missing_id")
        if sid in seen:
            raise ScenarioError(f"duplicate_scenario_id:{sid}")
        seen.add(sid)
        if row.get("domain") not in ALLOWED_DOMAINS:
            raise ScenarioError(f"invalid_domain:{row.get('domain')}")
        fields = {k: row[k] for k in PilotScenario.__dataclass_fields__ if k in row}
        out.append(PilotScenario(**fields))
    return out


def get_scenario(scenario_id: str, *, scenarios_json: Path | None = None) -> PilotScenario:
    for s in load_scenarios(scenarios_json=scenarios_json):
        if s.scenario_id == scenario_id:
            return s
    raise ScenarioError(f"unknown_scenario_id:{scenario_id}")


def ensure_fixture(scenario: PilotScenario) -> Path:
    """Lazily build the scenario's fixture docx (idempotent; never mutates if present)."""
    path = scenario.fixture_path()
    if path.is_file() and path.stat().st_size > 0:
        return path
    builder = FIXTURE_BUILDERS.get(scenario.fixture_builder)
    if builder is None:
        raise ScenarioError(f"unknown_fixture_builder:{scenario.fixture_builder}")
    builder(path)
    return path


def ensure_all_fixtures(*, scenarios_json: Path | None = None) -> dict[str, Path]:
    return {s.scenario_id: ensure_fixture(s) for s in load_scenarios(scenarios_json=scenarios_json)}
