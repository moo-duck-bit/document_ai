from __future__ import annotations

import re
from typing import Any

from docx.document import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from document_ai.learn.docx_io import (
    iter_blocks,
    paragraph_deep_text,
    set_paragraph_text,
    table_matrix,
)

RESIDUAL_PATTERN = re.compile(
    r"Mindrium|mindrium|범불안|의료기기|환자|"
    r"인지\s*행동|인지행동|인지치료|CBT|E06070|Electronic Medical|예시:\s*Mindrium|Dart/Flutter|"
    r"Android Studio,\s*X code|Python 3\.13|Chip\b|걱정의\s*유용성|역기능적|사회\s*기술\s*훈련|"
    r"명상\s*기법|범불안장애",
    re.I,
)

STRUCTURAL_LABELS = frozenset(
    {
        "Req.",
        "설명",
        "목적",
        "기준",
        "제품명",
        "모델명",
        "소프트웨어명",
        "안전성 등급",
        "제조자",
        "Code language",
        "Platform",
        "Editor",
        "결재",
        "성명",
        "직책",
        "서명",
        "일자",
        "승인자",
        "검토자",
        "작성자",
        "개정번호",
        "개정일자",
        "개정내용",
    }
)

OPERATING_PRINCIPLE_PATTERN = re.compile(
    r"범불안|의료기기|환자|인지\s*행동|인지행동|인지치료|CBT|Electronic Medical|"
    r"걱정|역기능적|사회\s*기술|명상|범불안장애",
    re.I,
)

PLACEHOLDER_COLORS = frozenset({"0000FF", "0070C0", "FF0000"})
STANDARD_ANCHOR_RE = re.compile(r"IEC\s*62304|ISO\s*14971", re.I)
PLACEHOLDER_DOC_CODE_RE = re.compile(r"XX[-–]XX", re.I)
COVER_TITLE_MARKERS = ("Software Requirement Specification", "Software Design Specification")
SCOPE_BOILERPLATE_RE = re.compile(r"1\.3에\s*식별된\s*제품에")

MINDRIUM_TEMPLATE_RESIDUAL = re.compile(
    r"Mindrium|mindrium|범불안|인지\s*행동|인지행동|인지치료|CBT|E06070|Electronic Medical|"
    r"예시:\s*Mindrium|Android\s*12|Python\s*3\.7|Google\s*Cloud\s*Platform|quad-core",
    re.I,
)


def _is_hospital_domain(domain: str) -> bool:
    return "hospital" in (domain or "").lower()


def is_template_residual(text: str, *, domain: str = "") -> bool:
    """Detect template/example residual text; hospital domains allow clinical terms."""
    if not text or not text.strip():
        return False
    if _is_hospital_domain(domain):
        return bool(
            MINDRIUM_TEMPLATE_RESIDUAL.search(text)
            or UNIQUE_MINDRIUM_PATTERN.search(text)
        )
    return bool(
        RESIDUAL_PATTERN.search(text)
        or UNIQUE_MINDRIUM_PATTERN.search(text)
        or OPERATING_PRINCIPLE_PATTERN.search(text)
    )

UNIQUE_MINDRIUM_PATTERN = re.compile(
    r"Android\s*12|Python\s*3\.7|Cloud\s*Run|Intel\s*Pentium|훈련\s*데이터|"
    r"환자\s*코드|환자의|의료기기|범불안|인지행동|CBT|대시보드를\s*통해|"
    r"Google\s*Cloud\s*Platform|quad-core\s*2GHz|128GB\s*이상|"
    r"의료기기\s*사이버|보건의료정보|의료기기\s*백업|의료기기\s*복구|"
    r"무선\s*통신\(Wi-Fi\)|입력된\s*모든\s*정보|중요\s*데이터는\s*정기|"
    r"데이터\s*통합을\s*위한|데이터\s*중복성\s*방지|감사\s*로그는\s*보안을",
    re.I,
)

PRODUCT_OVERVIEW_LABELS = frozenset(
    {
        "제품 개요",
        "동작원리",
        "사용목적",
        "구성",
        "컨텍스트 다이어그램",
        "제품특성",
        "사용자 특성",
        "일반적 제한",
        "가정 및 종속성",
        "소프트웨어 시스템 속성",
    }
)

UNIQUE_REQUIREMENTS_LABELS = frozenset(
    {
        "고유 요구사항",
        "개발 요구사항",
        "소프트웨어 실행 환경",
        "모바일 애플리케이션",
        "서버",
        "대시보드",
        "통신 인터페이스",
        "소프트웨어 개발 환경",
        "소프트웨어 성능 요구사항",
        "논리적 데이터베이스 요구사항",
        "데이터 구조 및 모델링",
        "데이터 무결성 및 정확성",
        "보안 및 데이터 암호화",
        "백업 및 재해 복구",
        "상호 운용성 및 통합",
        "감사 추적 및 규제 준수",
        "소프트웨어 기능 요구사항",
        "사이버 보안 필수 요구사항",
    }
)

SECURITY_REQ_TITLES: dict[str, str] = {
    "IA-01": "사용자 식별 및 인증",
    "IA-02": "계정 관리",
    "IA-03": "식별정보 관리",
    "IA-04": "인증정보 관리",
    "IA-05": "비밀번호 강도 설정",
    "IA-06": "인증정보에 대한 피드백",
    "IA-07": "연속적인 로그인 시도 실패 시 로그인 제한",
    "IA-08": "시스템 사용 알림 메시지",
    "UC-01": "권한 부여",
    "UC-02": "모바일 코드 사용 통제",
    "UC-03": "세션 잠금",
    "UC-04": "감사기록 생성",
    "UC-05": "감사 처리 실패 대응",
    "UC-06": "타임스탬프",
    "UC-07": "부인 방지",
    "SI-01": "통신에 대한 무결성 보장",
    "SI-02": "악성코드로부터 보호",
    "SI-03": "보안 기능 검증",
    "SI-04": "입력값 검증",
    "SI-05": "오류 처리",
    "SI-06": "메모리 관리",
    "SI-07": "설정 관리",
    "SI-08": "자원 관리",
    "SI-09": "시스템 무결성 보장",
    "SI-10": "보호 메커니즘 우회 방지",
    "SI-11": "보안 기능 접근 통제",
    "DC-01": "데이터 기밀성",
    "DC-02": "데이터 무결성",
    "DC-03": "데이터 가용성",
    "RA-01": "DoS 공격 방어",
    "RA-02": "자원 관리",
    "RA-03": "시스템 백업",
    "RA-04": "통신 보호",
    "RA-05": "네트워크 보호",
}


def _write_table_cell(table: Table, ri: int, ci: int, value: str) -> None:
    """Write via table.cell() to avoid merged-cell aliasing in row.cells."""
    if not value:
        return
    try:
        table.cell(ri, ci).text = value
    except (IndexError, ValueError):
        if ri < len(table.rows) and ci < len(table.rows[ri].cells):
            table.rows[ri].cells[ci].text = value


def _set_cell(row, col: int, value: str) -> None:
    if col < len(row.cells):
        row.cells[col].text = value


def _run_color_rgb(run) -> str | None:
    if run.font.color and run.font.color.rgb:
        return str(run.font.color.rgb).upper()
    return None


def _paragraph_has_colored_placeholder(paragraph: Paragraph) -> bool:
    text = (paragraph.text or "").strip()
    if not text:
        return False
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        color = _run_color_rgb(run)
        if color in PLACEHOLDER_COLORS and len(text) <= 4:
            return True
        if color in PLACEHOLDER_COLORS and (
            RESIDUAL_PATTERN.search(run.text) or run.text.strip().startswith("예시")
        ):
            return True
    return False


def _cell_has_colored_placeholder(cell) -> bool:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            color = _run_color_rgb(run)
            if color not in PLACEHOLDER_COLORS:
                continue
            if RESIDUAL_PATTERN.search(run.text) or run.text.strip().startswith("예시"):
                return True
            if color == "0000FF" and "Mindrium" in run.text:
                return True
    return False


def _document_paragraphs(doc: Document) -> list[Paragraph]:
    return [b for b in iter_blocks(doc) if isinstance(b, Paragraph)]


def _normalize_document_layout(content: dict[str, Any]) -> dict[str, Any]:
    layout = dict(content.get("document_layout") or {})
    if content.get("cover_product_line"):
        layout.setdefault("cover_product_line", content["cover_product_line"])
    if content.get("standards"):
        layout.setdefault("standards", content["standards"])
    if content.get("document_title"):
        layout.setdefault("document_title", content["document_title"])
    if content.get("document_title_sdp"):
        layout.setdefault("document_title_sdp", content["document_title_sdp"])
    if content.get("document_version_line"):
        layout.setdefault("document_version_line", content["document_version_line"])
    if content.get("narrative"):
        layout.setdefault("narrative", content["narrative"])

    paragraphs = content.get("paragraphs") or []
    if paragraphs:
        layout.setdefault("cover_title", paragraphs[0])
        if len(paragraphs) >= 3 and not layout.get("standards"):
            layout["standards"] = paragraphs[1:3]
        if len(paragraphs) >= 5:
            layout.setdefault("document_title", paragraphs[3])
            layout.setdefault("document_version_line", paragraphs[4])
        if len(paragraphs) >= 8 and not layout.get("narrative"):
            layout["narrative"] = paragraphs[5:8]
    return layout


def apply_cover_product_line(doc: Document, product_line: str) -> bool:
    if not product_line:
        return False
    paragraphs = _document_paragraphs(doc)
    for i, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() not in COVER_TITLE_MARKERS:
            continue
        for j in range(i + 1, min(i + 10, len(paragraphs))):
            if not paragraphs[j].text.strip():
                paragraphs[j].text = product_line
                return True
    return False


def apply_signature_date_table(doc: Document, date: str) -> bool:
    """Fill 5-column cover signature dates (Mindrium filled example pattern)."""
    if not date:
        return False
    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        matrix = table_matrix(block)
        if not matrix or len(matrix[0]) < 5:
            continue
        header = " ".join(matrix[0])
        if any(token in header for token in ("제품명", "결재", "개정번호", "Req.")):
            continue
        applied = False
        for ri in range(1, min(len(block.rows), 4)):
            _write_table_cell(block, ri, 4, date)
            applied = True
        if applied:
            return True
    return False


def apply_cover_signature_people(doc: Document, approval: dict[str, Any]) -> bool:
    """Fill headerless 5-column cover approval rows (name, title, date)."""
    people = approval.get("people") or {}
    date = approval.get("date", "")
    if not people:
        return False

    role_order = ["승인자", "검토자", "작성자"]
    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        matrix = table_matrix(block)
        if not matrix or len(matrix[0]) < 5:
            continue
        header = " ".join(matrix[0])
        if any(token in header for token in ("제품명", "결재", "개정번호", "Req.")):
            continue
        applied = False
        for ri, role in enumerate(role_order, start=1):
            if ri >= len(block.rows):
                break
            info = people.get(role) or {}
            if info.get("name"):
                _write_table_cell(block, ri, 0, role)
                _write_table_cell(block, ri, 1, info["name"])
            if info.get("title"):
                _write_table_cell(block, ri, 2, info["title"])
            if date:
                _write_table_cell(block, ri, 4, date)
            applied = True
        if applied:
            return True
    return False


def apply_document_layout(doc: Document, content: dict[str, Any]) -> dict[str, int]:
    """Place cover, standards, document codes, and narrative like filled Mindrium layout."""
    layout = _normalize_document_layout(content)
    stats = {
        "cover_product_line": 0,
        "standards_filled": 0,
        "document_codes_filled": 0,
        "narrative_filled": 0,
    }

    cover_line = layout.get("cover_product_line") or layout.get("manufacturer") or ""
    if apply_cover_product_line(doc, cover_line):
        stats["cover_product_line"] = 1

    paragraphs = _document_paragraphs(doc)
    standards = layout.get("standards") or []
    for paragraph in paragraphs:
        if STANDARD_ANCHOR_RE.search(paragraph.text) and standards:
            paragraph.text = standards[min(stats["standards_filled"], len(standards) - 1)]
            stats["standards_filled"] += 1

    doc_title = layout.get("document_title", "")
    doc_version = layout.get("document_version_line", "")
    placeholder_paragraphs = [p for p in paragraphs if PLACEHOLDER_DOC_CODE_RE.search(p.text)]
    if doc_title and placeholder_paragraphs:
        placeholder_paragraphs[0].text = doc_title
        stats["document_codes_filled"] += 1
    if doc_version and len(placeholder_paragraphs) > 1:
        placeholder_paragraphs[1].text = doc_version
        stats["document_codes_filled"] += 1

    sdp_title = layout.get("document_title_sdp", "")
    for paragraph in paragraphs:
        text = paragraph.text or ""
        if not (PLACEHOLDER_DOC_CODE_RE.search(text) or re.search(r"\bX{3,}\b", text)):
            continue
        if "개발 계획서" in text and sdp_title:
            paragraph.text = sdp_title
            stats["document_codes_filled"] += 1
        elif ("요구사항" in text or "설계" in text) and doc_title:
            paragraph.text = doc_title
            stats["document_codes_filled"] += 1
        elif doc_version:
            paragraph.text = doc_version
            stats["document_codes_filled"] += 1

    narrative = layout.get("narrative") or []
    narrative_index = 0
    past_intro = False
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if STANDARD_ANCHOR_RE.search(text):
            past_intro = True
            continue
        if doc_title and text == doc_title:
            past_intro = True
            continue
        if doc_version and text == doc_version:
            past_intro = True
            continue
        if PLACEHOLDER_DOC_CODE_RE.search(text):
            past_intro = True
            continue
        if past_intro and not text and narrative_index < len(narrative):
            paragraph.text = narrative[narrative_index]
            narrative_index += 1
            stats["narrative_filled"] += 1
    return stats


def apply_mddr_residual_sections(doc: Document, content: dict[str, Any]) -> dict[str, int]:
    """Replace MDDR Mindrium boilerplate under 목적/사용자/범위/시스템 개요."""
    overview = content.get("overview") or {}
    stats = {"sections": 0}
    paragraphs = _document_paragraphs(doc)
    skip_labels = {
        "목적",
        "범위",
        "사용자",
        "시스템 개요",
        "개요",
        "식별",
        "참조문서",
        "목차",
        "개정이력",
        "소프트웨어 개요",
    }

    def _fill_after(label: str, texts: list[str] | str) -> None:
        nonlocal stats
        if isinstance(texts, str):
            values = [texts] if texts else []
        else:
            values = [text for text in texts if text]
        if not values:
            return
        value_index = 0
        for i, paragraph in enumerate(paragraphs):
            if paragraph_deep_text(paragraph) != label:
                continue
            for j in range(i + 1, min(i + 12, len(paragraphs))):
                candidate = paragraph_deep_text(paragraphs[j])
                if candidate in skip_labels or candidate.startswith("Figure"):
                    break
                if value_index >= len(values):
                    break
                if (
                    not candidate
                    or candidate == "."
                    or RESIDUAL_PATTERN.search(candidate)
                    or OPERATING_PRINCIPLE_PATTERN.search(candidate)
                ):
                    set_paragraph_text(paragraphs[j], values[value_index])
                    value_index += 1
                    stats["sections"] += 1
            break

    _fill_after("목적", overview.get("purpose", ""))
    _fill_after("범위", overview.get("scope", ""))
    _fill_after("사용자", overview.get("users") or [])
    _fill_after("시스템 개요", overview.get("system", ""))

    software_purpose = overview.get("software_purpose") or overview.get("purpose", "")
    if software_purpose:
        for i, paragraph in enumerate(paragraphs):
            if paragraph_deep_text(paragraph) != "소프트웨어 개요":
                continue
            for j in range(i + 1, min(i + 8, len(paragraphs))):
                if paragraph_deep_text(paragraphs[j]) != "목적":
                    continue
                for k in range(j + 1, min(j + 5, len(paragraphs))):
                    text = paragraph_deep_text(paragraphs[k])
                    if (
                        not text
                        or text == "."
                        or RESIDUAL_PATTERN.search(text)
                        or OPERATING_PRINCIPLE_PATTERN.search(text)
                    ):
                        set_paragraph_text(paragraphs[k], software_purpose)
                        stats["sections"] += 1
                        break
            break

    for paragraph in paragraphs:
        text = paragraph_deep_text(paragraph)
        if SCOPE_BOILERPLATE_RE.search(text) and overview.get("scope"):
            set_paragraph_text(paragraph, overview["scope"])
            stats["sections"] += 1

    return stats


def _find_operating_principle_slots(doc: Document) -> list[int]:
    """Block indices of template paragraphs used for section 3.1 동작원리."""
    slots: list[int] = []
    for bi, block in enumerate(iter_blocks(doc)):
        if not isinstance(block, Paragraph):
            continue
        text = block.text or ""
        if OPERATING_PRINCIPLE_PATTERN.search(text) or _paragraph_has_colored_placeholder(block):
            slots.append(bi)
    return slots


def apply_operating_principle(
    doc: Document,
    paragraphs: list[str],
    slot_indices: list[int] | None = None,
) -> int:
    """Fill 동작원리 slots with case-specific narrative; clear colored placeholders."""
    indices = slot_indices if slot_indices is not None else _find_operating_principle_slots(doc)
    if not indices:
        return 0

    blocks = list(iter_blocks(doc))
    filled = 0
    for slot_i, block_index in enumerate(indices):
        if block_index >= len(blocks):
            continue
        block = blocks[block_index]
        if not isinstance(block, Paragraph):
            continue
        value = paragraphs[slot_i] if slot_i < len(paragraphs) else ""
        block.text = value
        if value:
            filled += 1
    return filled


def _is_structural(text: str) -> bool:
    t = text.strip()
    if not t:
        return False
    if re.match(r"^Req\.\s*\d+", t):
        return True
    if re.match(r"^[A-Z]{2}-\d+", t):
        return True
    return t in STRUCTURAL_LABELS or t.startswith("IA-") or t.startswith("UC-")


def sanitize_document(doc: Document, *, domain: str = "") -> dict[str, int]:
    """Remove Mindrium/medical placeholder text; preserve structural labels."""
    stats = {"paragraphs": 0, "table_cells": 0}

    for paragraph in doc.paragraphs:
        text = paragraph_deep_text(paragraph)
        if not text:
            continue
        if text in PRODUCT_OVERVIEW_LABELS or text in UNIQUE_REQUIREMENTS_LABELS:
            continue
        if text.startswith("그림 ") or text.startswith("Figure "):
            continue
        if (
            is_template_residual(text, domain=domain)
            or _paragraph_has_colored_placeholder(paragraph)
        ):
            set_paragraph_text(paragraph, "")
            stats["paragraphs"] += 1

    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        for row in block.rows:
            for cell in row.cells:
                text = cell.text or ""
                if not text.strip():
                    continue
                if _is_structural(text):
                    continue
                if (
                    is_template_residual(text, domain=domain)
                    or text.strip() in {"2", "Mindrium"}
                    or _cell_has_colored_placeholder(cell)
                ):
                    cell.text = ""
                    stats["table_cells"] += 1
                elif not _is_hospital_domain(domain) and (
                    "Mindrium" in text or "의료" in text or "환자" in text
                ):
                    cell.text = ""
                    stats["table_cells"] += 1

    return stats


def apply_paragraphs_block_order(doc: Document, paragraphs: list[str]) -> int:
    """Set body paragraphs by block order (includes empty slots)."""
    para_blocks = [b for b in iter_blocks(doc) if isinstance(b, Paragraph)]
    count = 0
    for i, text in enumerate(paragraphs):
        if i >= len(para_blocks):
            break
        if text:
            para_blocks[i].text = text
            count += 1
    return count


def apply_approval_table(doc: Document, approval: dict[str, Any]) -> bool:
    date = approval.get("date", "")
    people = approval.get("people") or {}
    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        labels = {row.cells[0].text.strip() for row in block.rows if row.cells}
        if not {"승인자", "검토자", "작성자"} & labels:
            continue
        for ri, row in enumerate(block.rows):
            label = row.cells[0].text.strip()
            if label not in people:
                continue
            info = people[label]
            if info.get("name"):
                _write_table_cell(block, ri, 1, info["name"])
            if info.get("title"):
                _write_table_cell(block, ri, 2, info["title"])
            if date:
                _write_table_cell(block, ri, 4, date)
        return True
    return False


def apply_revision_table(doc: Document, revision: dict[str, str]) -> bool:
    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        matrix = table_matrix(block)
        if not matrix:
            continue
        header = " ".join(matrix[0])
        has_revision = "개정번호" in header or (
            len(matrix) > 1 and matrix[1][0].strip() in {"0", "1"}
        )
        if not has_revision:
            continue
        if len(block.rows) > 1:
            _write_table_cell(block, 1, 0, revision.get("number", "0"))
            _write_table_cell(block, 1, 1, revision.get("date", ""))
            _write_table_cell(block, 1, 2, revision.get("summary", "제정"))
            _write_table_cell(block, 1, 3, revision.get("author", ""))
        return True
    return False


def apply_matrix_table(doc: Document, anchor: str, matrix_data: list[list[str]]) -> bool:
    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        m = table_matrix(block)
        if not m or m[0][0] != anchor:
            continue
        for ri, row_data in enumerate(matrix_data):
            if ri >= len(block.rows):
                break
            for ci, val in enumerate(row_data):
                _write_table_cell(block, ri, ci, val)
        return True
    return False


def apply_component_detail_table(doc: Document, rows: list[list[str]]) -> bool:
    seen_product = False
    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        m = table_matrix(block)
        if not m:
            continue
        if m[0][0] == "제품명":
            seen_product = True
            continue
        if not seen_product:
            continue
        labels = {r[0].strip() for r in m if r}
        if "Code language" in labels:
            break
        for ri, row_data in enumerate(rows):
            if ri >= len(block.rows):
                break
            for ci, val in enumerate(row_data):
                _write_table_cell(block, ri, ci, val)
        return True
    return False


def _fill_unused_tech_stack_slots(doc: Document, used_count: int) -> int:
    """Mark unused Code-language tables as N/A so template slots are not left blank."""
    seen = 0
    filled = 0
    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        matrix = table_matrix(block)
        if not matrix:
            continue
        labels = {row[0].strip() for row in matrix if row}
        if "Code language" not in labels:
            continue
        seen += 1
        if seen <= used_count:
            continue
        _write_table_cell(block, 0, 1, "해당 없음")
        for ri, row in enumerate(block.rows):
            label = row.cells[0].text.strip()
            if label == "Code language":
                _write_table_cell(block, ri, 1, "N/A")
            elif label in {"Platform", "Editor"}:
                _write_table_cell(block, ri, 1, "N/A")
        filled += 1
    return filled


def apply_tech_stack_tables(doc: Document, stacks: list[dict[str, str]]) -> int:
    applied = 0
    stack_idx = 0
    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        matrix = table_matrix(block)
        if not matrix:
            continue
        labels = {row[0].strip() for row in matrix if row}
        if "Code language" not in labels:
            continue
        if stack_idx >= len(stacks):
            break
        stack = stacks[stack_idx]
        header = stack.get("header")
        if header:
            _write_table_cell(block, 0, 1, header)
        for ri, row in enumerate(block.rows):
            label = row.cells[0].text.strip()
            if label == "Code language":
                _write_table_cell(block, ri, 1, stack.get("code_language", ""))
            elif label == "Platform":
                _write_table_cell(block, ri, 1, stack.get("platform", ""))
            elif label == "Editor":
                _write_table_cell(block, ri, 1, stack.get("editor", ""))
        stack_idx += 1
        applied += 1
    return applied


def apply_traceability_details(
    doc: Document, traceability: list[dict[str, Any]], *, default_applicability: str = "해당"
) -> int:
    from document_ai.render.requirements import _find_traceability_table_index

    trace_idx = _find_traceability_table_index(doc)
    if trace_idx is None:
        return 0

    table = doc.tables[trace_idx]
    by_id = {t["requirement"]: t for t in traceability}
    count = 0

    for ri, row in enumerate(table.rows):
        key = table.cell(ri, 0).text.strip()
        if not re.match(r"^[A-Z]{2}-\d+", key):
            continue
        entry = by_id.get(key, {})
        title = entry.get("title") or SECURITY_REQ_TITLES.get(key, "")
        if title:
            # Col 1 may be vertically merged in template; put title beside the ID in col 0.
            base_key = key.split()[0]
            _write_table_cell(table, ri, 0, f"{base_key} {title}")
        app = entry.get("applicability")
        if not app:
            app = "비해당" if entry.get("linked_reqs") == "N/A" else default_applicability
        _write_table_cell(table, ri, 2, app)
        linked = entry.get("linked_reqs", "")
        _write_table_cell(table, ri, 3, linked if linked else "N/A")
        count += 1
    return count


def verify_completeness(
    doc: Document,
    requirements: list[dict[str, Any]],
    *,
    domain: str = "",
) -> dict[str, Any]:
    """Review pass: empty req values, residual text, empty traceability."""
    issues: list[str] = []
    residual: list[str] = []

    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            t = paragraph_deep_text(block)
            if is_template_residual(t, domain=domain):
                residual.append(f"paragraph: {t[:80]}")
        elif isinstance(block, Table):
            m = table_matrix(block)
            if m and re.match(r"^Req\.\s*\d+", m[0][0]):
                rid = m[0][0]
                has_content = any(
                    c.strip()
                    for ri, row in enumerate(m)
                    for ci, c in enumerate(row)
                    if ri > 0 and ci > 0
                )
                if not has_content:
                    issues.append(f"{rid}: no description/purpose filled")
            if any(r and r[0].startswith("IA-01") for r in m):
                for row in m:
                    if row and re.match(r"^[A-Z]{2}-\d+", row[0]):
                        if len(row) < 4 or not row[3].strip():
                            issues.append(f"traceability {row[0]}: missing linked reqs")

    for block in iter_blocks(doc):
        if not isinstance(block, Table):
            continue
        m = table_matrix(block)
        if not m or m[0][0] != "제품명":
            continue
        for ri, row in enumerate(m):
            label = row[0].strip() if row else ""
            for ci, c in enumerate(row):
                if ci == 0:
                    continue
                if label == "제조자" and ci > 1:
                    continue
                if not c.strip():
                    issues.append(f"product_table r{ri}c{ci} empty")
                if is_template_residual(c, domain=domain):
                    residual.append(f"product r{ri}c{ci}: {c[:60]}")

    return {
        "issue_count": len(issues),
        "residual_count": len(residual),
        "issues": issues[:50],
        "residual": residual[:20],
        "requirements_expected": len(requirements),
    }


def _find_paragraph_by_label(doc: Document, label: str, *, min_index: int = 0) -> int | None:
    for index, paragraph in enumerate(doc.paragraphs):
        if index < min_index:
            continue
        if paragraph_deep_text(paragraph) == label:
            return index
    return None


def _fill_after_label(
    doc: Document,
    label: str,
    text: str,
    *,
    stop_labels: frozenset[str],
    max_scan: int = 8,
    min_index: int = 0,
    skip_labels: frozenset[str] | None = None,
) -> bool:
    if not text:
        return False
    start = _find_paragraph_by_label(doc, label, min_index=min_index)
    if start is None:
        return False
    for index in range(start + 1, min(start + 1 + max_scan, len(doc.paragraphs))):
        candidate = paragraph_deep_text(doc.paragraphs[index])
        if candidate in stop_labels:
            break
        if skip_labels and candidate in skip_labels:
            continue
        if candidate.startswith("그림 ") or candidate.startswith("Figure "):
            break
        if not candidate or candidate == ".":
            set_paragraph_text(doc.paragraphs[index], text)
            return True
    return False


def _replace_paragraphs_after_label(
    doc: Document,
    label: str,
    texts: list[str],
    *,
    stop_labels: frozenset[str],
    max_scan: int = 20,
    min_index: int = 0,
) -> int:
    if not texts:
        return 0
    start = _find_paragraph_by_label(doc, label, min_index=min_index)
    if start is None:
        return 0
    filled = 0
    text_index = 0
    for index in range(start + 1, min(start + 1 + max_scan, len(doc.paragraphs))):
        candidate = paragraph_deep_text(doc.paragraphs[index])
        if candidate in stop_labels:
            break
        if candidate.startswith("그림 ") or candidate.startswith("Figure "):
            break
        if text_index < len(texts):
            set_paragraph_text(doc.paragraphs[index], texts[text_index])
            text_index += 1
            filled += 1
            continue
        if (
            not candidate
            or RESIDUAL_PATTERN.search(candidate)
            or UNIQUE_MINDRIUM_PATTERN.search(candidate)
            or candidate.endswith("(예시)")
        ):
            set_paragraph_text(doc.paragraphs[index], "")
            filled += 1
    return filled


def _replace_exact_paragraphs(doc: Document, replacements: dict[str, str]) -> int:
    filled = 0
    for paragraph in doc.paragraphs:
        key = paragraph_deep_text(paragraph)
        if key in replacements:
            set_paragraph_text(paragraph, replacements[key])
            filled += 1
    return filled


def apply_product_overview_sections(doc: Document, overview: dict[str, Any]) -> dict[str, int]:
    """Fill section 3 (제품 개요) narrative slots in SDT paragraphs."""
    stats = {"filled": 0}
    if not overview:
        return stats

    stop = PRODUCT_OVERVIEW_LABELS | UNIQUE_REQUIREMENTS_LABELS

    if overview.get("figure_caption"):
        for paragraph in doc.paragraphs:
            text = paragraph_deep_text(paragraph)
            if text.startswith("그림 1."):
                set_paragraph_text(paragraph, overview["figure_caption"])
                stats["filled"] += 1
                break

    if overview.get("product_characteristics"):
        for paragraph in doc.paragraphs:
            if paragraph_deep_text(paragraph) == "제품특성":
                continue
            if paragraph_deep_text(paragraph) == ".":
                set_paragraph_text(paragraph, overview["product_characteristics"])
                stats["filled"] += 1
                break

    field_map = (
        ("operating_principle", "동작원리"),
        ("usage_purpose", "사용목적"),
        ("composition", "컨텍스트 다이어그램"),
        ("user_characteristics", "사용자 특성"),
        ("general_constraints", "일반적 제한"),
        ("assumptions_dependencies", "가정 및 종속성"),
        ("system_attributes", "소프트웨어 시스템 속성"),
    )
    for field, label in field_map:
        if _fill_after_label(doc, label, overview.get(field, ""), stop_labels=stop):
            stats["filled"] += 1
    return stats


def apply_unique_requirements_sections(doc: Document, unique: dict[str, Any]) -> dict[str, int]:
    """Replace section 4 (고유 요구사항) Mindrium boilerplate with case-specific text."""
    stats = {"filled": 0}
    if not unique:
        return stats

    stop = UNIQUE_REQUIREMENTS_LABELS | PRODUCT_OVERVIEW_LABELS

    exec_anchor = _find_paragraph_by_label(doc, "소프트웨어 실행 환경") or 0
    exec_env = unique.get("execution_environment") or {}
    stats["filled"] += _replace_paragraphs_after_label(
        doc,
        "모바일 애플리케이션",
        exec_env.get("jm_web") or [],
        stop_labels=stop,
        max_scan=12,
        min_index=exec_anchor,
    )
    stats["filled"] += _replace_paragraphs_after_label(
        doc,
        "서버",
        exec_env.get("jm_api") or [],
        stop_labels=stop,
        max_scan=12,
        min_index=exec_anchor,
    )
    stats["filled"] += _replace_paragraphs_after_label(
        doc,
        "대시보드",
        exec_env.get("jm_admin") or [],
        stop_labels=stop,
        max_scan=12,
        min_index=exec_anchor,
    )

    if unique.get("communication_interface"):
        stats["filled"] += _replace_paragraphs_after_label(
            doc,
            "통신 인터페이스",
            [unique["communication_interface"]],
            stop_labels=stop,
            max_scan=4,
            min_index=exec_anchor,
        )

    dev_anchor = _find_paragraph_by_label(doc, "소프트웨어 개발 환경") or exec_anchor
    dev_env = unique.get("development_environment") or {}
    for label, key in (
        ("모바일 애플리케이션", "jm_web"),
        ("서버", "jm_api"),
        ("대시보드", "jm_admin"),
    ):
        texts = dev_env.get(key) or []
        if texts:
            stats["filled"] += _replace_paragraphs_after_label(
                doc, label, texts, stop_labels=stop, max_scan=6, min_index=dev_anchor
            )

    stats["filled"] += _replace_paragraphs_after_label(
        doc, "소프트웨어 성능 요구사항", unique.get("performance") or [], stop_labels=stop, max_scan=6
    )

    db_sections = unique.get("database_sections") or {}
    db_order = (
        ("structure", "데이터 구조 및 모델링"),
        ("integrity", "데이터 무결성 및 정확성"),
        ("encryption", "보안 및 데이터 암호화"),
        ("backup", "백업 및 재해 복구"),
        ("integration", "상호 운용성 및 통합"),
        ("audit", "감사 추적 및 규제 준수"),
    )
    for field, label in db_order:
        texts = db_sections.get(field) or []
        if isinstance(texts, str):
            texts = [texts]
        stats["filled"] += _replace_paragraphs_after_label(
            doc, label, texts, stop_labels=stop, max_scan=6
        )

    if unique.get("cybersecurity_intro"):
        stats["filled"] += _replace_paragraphs_after_label(
            doc,
            "사이버 보안 필수 요구사항",
            [unique["cybersecurity_intro"]],
            stop_labels=stop,
            max_scan=4,
        )

    stats["filled"] += _replace_exact_paragraphs(doc, unique.get("exact_replacements") or {})
    return stats


def apply_paragraph_overrides(doc: Document, content: dict[str, Any]) -> int:
    if not content.get("paragraph_overrides"):
        return 0
    parsed = {int(k): v for k, v in content["paragraph_overrides"].items()}
    paras = _document_paragraphs(doc)
    filled = 0
    for idx, text in parsed.items():
        if idx < len(paras) and text:
            paras[idx].text = text
            filled += 1
    return filled


def apply_mdsr_content(
    doc: Document,
    content: dict[str, Any],
    *,
    operating_principle_slots: list[int] | None = None,
) -> dict[str, Any]:
    stats: dict[str, Any] = {}
    domain = str(content.get("domain", ""))
    stats["sanitized"] = sanitize_document(doc, domain=domain)

    if content.get("approval"):
        approval = content["approval"]
        stats["signature_dates"] = apply_signature_date_table(doc, approval.get("date", ""))
        stats["cover_signature"] = apply_cover_signature_people(doc, approval)
        stats["approval"] = apply_approval_table(doc, approval) or stats.get("cover_signature", False)

    layout_stats = apply_document_layout(doc, content)
    stats["layout"] = layout_stats
    stats["paragraphs"] = (
        layout_stats.get("cover_product_line", 0)
        + layout_stats.get("standards_filled", 0)
        + layout_stats.get("document_codes_filled", 0)
        + layout_stats.get("narrative_filled", 0)
    )

    if content.get("paragraph_overrides"):
        stats["paragraph_overrides"] = apply_paragraph_overrides(doc, content)

    if content.get("revision"):
        stats["revision"] = apply_revision_table(doc, content["revision"])
    if content.get("product_matrix"):
        stats["product_matrix"] = apply_matrix_table(doc, "제품명", content["product_matrix"])
    elif content.get("product"):
        stats["product_matrix"] = apply_matrix_table(
            doc, "제품명", _product_dict_to_matrix(content["product"])
        )
    if content.get("component_detail_rows"):
        stats["component_detail"] = apply_component_detail_table(doc, content["component_detail_rows"])
    if content.get("tech_stacks"):
        stats["tech_stacks"] = apply_tech_stack_tables(doc, content["tech_stacks"])
        stats["tech_stacks"] += _fill_unused_tech_stack_slots(doc, len(content["tech_stacks"]))
    if content.get("operating_principle"):
        stats["operating_principle"] = apply_operating_principle(
            doc,
            content["operating_principle"],
            slot_indices=operating_principle_slots,
        )

    if content.get("product_overview"):
        stats["product_overview"] = apply_product_overview_sections(doc, content["product_overview"])
    if content.get("unique_requirements"):
        stats["unique_requirements"] = apply_unique_requirements_sections(
            doc, content["unique_requirements"]
        )

    return stats


def _product_dict_to_matrix(product: dict[str, Any]) -> list[list[str]]:
    names = product.get("product_name") or ["", "", ""]
    models = product.get("model_name") or ["", "", ""]
    comps = product.get("software_components") or ["", "", ""]
    safety = product.get("safety_class") or ["A", "A", "A"]
    mfr = product.get("manufacturer", "")
    return [
        ["제품명", *names[:3]],
        ["모델명", *models[:3]],
        ["소프트웨어명", "고객몰", "API 서버", "운영 백오피스"],
        ["소프트웨어명", *comps[:3]],
        ["안전성 등급", *safety[:3]],
        ["제조자", mfr, mfr, mfr],
    ]


def apply_mdsr_full(
    doc: Document,
    content: dict[str, Any],
    requirements_payload: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Sanitize Mindrium template then fill all MDSR sections."""
    from document_ai.render.requirements import apply_requirements_payload

    operating_principle_slots = (
        _find_operating_principle_slots(doc) if content.get("operating_principle") else []
    )
    domain = str(content.get("domain", ""))
    stats = apply_mdsr_content(doc, content, operating_principle_slots=operating_principle_slots)

    if requirements_payload:
        stats["requirements"] = apply_requirements_payload(
            doc, requirements_payload, overwrite_descriptions=True
        )
        stats["traceability"] = apply_traceability_details(
            doc, requirements_payload.get("traceability", [])
        )
        for _pass in range(1, 6):
            review = verify_completeness(
                doc,
                requirements_payload.get("requirements", []),
                domain=domain,
            )
            stats[f"review_pass_{_pass}"] = review
            if review["issue_count"] == 0 and review["residual_count"] == 0:
                stats["review_passes_completed"] = _pass
                break
            sanitize_document(doc, domain=domain)
            apply_mdsr_content(doc, content, operating_principle_slots=operating_principle_slots)
            apply_requirements_payload(doc, requirements_payload, overwrite_descriptions=True)
            apply_traceability_details(doc, requirements_payload.get("traceability", []))
        else:
            stats["review_passes_completed"] = 5

    return stats
