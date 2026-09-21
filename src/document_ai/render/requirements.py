from __future__ import annotations

import re
from copy import deepcopy
from typing import Any

from docx.document import Document
from docx.table import Table

from document_ai.learn.docx_io import table_matrix


def _req_id_from_table(table: Table) -> str | None:
    if not table.rows:
        return None
    text = table.rows[0].cells[0].text.strip()
    if re.match(r"^Req\.\s*\d+", text):
        return text
    return None


def _find_traceability_table_index(doc: Document) -> int | None:
    for i, table in enumerate(doc.tables):
        matrix = table_matrix(table)
        if any(row and row[0].startswith("IA-01") for row in matrix):
            return i
    return len(doc.tables) - 1 if doc.tables else None


def _find_last_req_table_index(doc: Document, trace_index: int) -> int:
    for i in range(trace_index - 1, -1, -1):
        if _req_id_from_table(doc.tables[i]):
            return i
    raise ValueError("No requirement table prototype found in template")


def _set_description_cell(table: Table, description: str, *, overwrite: bool = False) -> None:
    if description is None:
        return
    for row in table.rows[1:]:
        if len(row.cells) < 2:
            continue
        if row.cells[0].text.strip() == "설명":
            if overwrite or not row.cells[1].text.strip():
                row.cells[1].text = description
            return

    for row_idx in (2, 1, 3):
        if row_idx >= len(table.rows):
            continue
        row = table.rows[row_idx]
        if len(row.cells) < 2:
            continue
        label = row.cells[0].text.strip()
        if label in {"목적", "기준", "설명"}:
            continue
        if overwrite or not row.cells[1].text.strip():
            row.cells[1].text = description
            return
    if table.rows[-1].cells and (overwrite or not table.rows[-1].cells[-1].text.strip()):
        table.rows[-1].cells[-1].text = description


REQ_LABELS = ("설명", "목적", "기준")


def _ensure_req_row_labels(table: Table) -> bool:
    """Ensure data rows use 설명/목적/기준 labels when template shells are blank."""
    if len(table.rows) < 4:
        return False
    changed = False
    for ri, label in enumerate(REQ_LABELS, start=1):
        if ri >= len(table.rows):
            break
        row = table.rows[ri]
        if len(row.cells) < 2:
            continue
        current = row.cells[0].text.strip()
        if not current:
            row.cells[0].text = label
            changed = True
    return changed


def _short_req_title(req: dict[str, Any]) -> str:
    if req.get("title") or req.get("summary"):
        return str(req.get("title") or req.get("summary"))
    description = (req.get("description") or "").strip()
    if not description:
        return ""
    return description.split(".")[0].strip()[:80]


def _default_criteria(req_id: str) -> str:
    m = re.search(r"(\d+)", req_id)
    num = int(m.group(1)) if m else 0
    if num >= 200:
        return "명세된 비기능 목표·측정 기준을 성능/가용성/호환성 시험에서 검증 가능해야 한다."
    if num >= 100:
        return "OWASP ASVS·PCI DSS·개인정보보호법 관련 통제를 보안 시험에서 검증 가능해야 한다."
    return ""


def _fill_req_table_full(table: Table, req: dict[str, Any], *, overwrite: bool = False) -> bool:
    description = req.get("description", "")
    purpose = req.get("purpose", "")
    criteria = req.get("criteria", "") or _default_criteria(_req_id_from_table(table) or "")
    title = _short_req_title(req)
    if not any([description, purpose, criteria, title]):
        return False

    req_id = _req_id_from_table(table) or ""
    _ensure_req_row_labels(table)

    if title and len(table.rows[0].cells) > 1:
        cell = table.rows[0].cells[1]
        if overwrite or not cell.text.strip():
            cell.text = title

    filled = False
    unlabeled: list = []

    for row in table.rows[1:]:
        if len(row.cells) < 2:
            continue
        label = row.cells[0].text.strip()
        cell = row.cells[1]
        if label == "설명":
            if description and (overwrite or not cell.text.strip()):
                cell.text = description
                filled = True
        elif label == "목적":
            if purpose and (overwrite or not cell.text.strip()):
                cell.text = purpose
                filled = True
        elif label == "기준":
            if criteria and (overwrite or not cell.text.strip()):
                cell.text = criteria
                filled = True
        elif not label:
            unlabeled.append(cell)

    if unlabeled and not any(
        row.cells[0].text.strip() in REQ_LABELS for row in table.rows[1:] if len(row.cells) >= 1
    ):
        if len(unlabeled) == 1:
            if description and (overwrite or not unlabeled[0].text.strip()):
                unlabeled[0].text = description
                filled = True
        elif len(unlabeled) >= 2:
            if purpose and (overwrite or not unlabeled[0].text.strip()):
                unlabeled[0].text = purpose
                filled = True
            if description and (overwrite or not unlabeled[-1].text.strip()):
                unlabeled[-1].text = description
                filled = True
        elif description and (overwrite or not unlabeled[0].text.strip()):
            unlabeled[0].text = description
            filled = True

    return filled


def remove_extra_requirement_tables(doc: Document, req_ids: list[str]) -> list[str]:
    """Drop requirement shells present in the template but not in the case payload."""
    wanted = set(req_ids)
    trace_idx = _find_traceability_table_index(doc)
    if trace_idx is None:
        return []

    removed: list[str] = []
    for table in doc.tables[:trace_idx]:
        req_id = _req_id_from_table(table)
        if req_id and req_id not in wanted:
            table._tbl.getparent().remove(table._tbl)
            removed.append(req_id)
    return removed


def ensure_requirement_tables(doc: Document, req_ids: list[str]) -> list[str]:
    trace_idx = _find_traceability_table_index(doc)
    if trace_idx is None:
        return []

    prototype_idx = _find_last_req_table_index(doc, trace_idx)
    prototype = doc.tables[prototype_idx]._tbl
    trace_element = doc.tables[trace_idx]._tbl

    existing: set[str] = set()
    for table in doc.tables:
        rid = _req_id_from_table(table)
        if rid:
            existing.add(rid)

    def sort_key(req_id: str) -> int:
        m = re.search(r"(\d+)", req_id)
        return int(m.group(1)) if m else 0

    added: list[str] = []
    for req_id in sorted(req_ids, key=sort_key):
        if req_id in existing:
            continue
        new_tbl = deepcopy(prototype)
        trace_element.addprevious(new_tbl)
        new_table = Table(new_tbl, doc)
        new_table.rows[0].cells[0].text = req_id
        for ri in range(1, len(new_table.rows)):
            for cell in new_table.rows[ri].cells:
                if cell.text.strip() and not re.match(r"^Req\.", cell.text.strip()):
                    cell.text = ""
        _ensure_req_row_labels(new_table)
        existing.add(req_id)
        added.append(req_id)

    return added


def fill_requirement_tables(
    doc: Document, requirements: list[dict[str, Any]], *, overwrite: bool = False
) -> list[str]:
    by_id = {r["req_id"]: r for r in requirements}
    filled: list[str] = []

    for table in doc.tables:
        req_id = _req_id_from_table(table)
        if not req_id:
            continue
        req = by_id.get(req_id)
        if not req:
            continue
        if _fill_req_table_full(table, req, overwrite=overwrite):
            filled.append(req_id)
    return filled


def fill_traceability_matrix(doc: Document, traceability: list[dict[str, Any]]) -> int:
    if not traceability:
        return 0

    trace_idx = _find_traceability_table_index(doc)
    if trace_idx is None:
        return 0

    table = doc.tables[trace_idx]
    by_req = {t["requirement"]: t.get("linked_reqs", "") for t in traceability}
    count = 0

    for row in table.rows:
        if not row.cells:
            continue
        key = row.cells[0].text.strip()
        if key in by_req and by_req[key]:
            if len(row.cells) >= 4:
                row.cells[3].text = by_req[key]
                count += 1
    return count


def apply_requirements_payload(
    doc: Document, payload: dict[str, Any], *, overwrite_descriptions: bool = False
) -> dict[str, Any]:
    requirements = payload.get("requirements", [])
    traceability = payload.get("traceability", [])
    req_ids = [r["req_id"] for r in requirements]

    removed = remove_extra_requirement_tables(doc, req_ids)
    added = ensure_requirement_tables(doc, req_ids)
    desc_filled = fill_requirement_tables(doc, requirements, overwrite=overwrite_descriptions)
    trace_filled = fill_traceability_matrix(doc, traceability)

    return {
        "req_tables_removed": removed,
        "req_tables_added": added,
        "req_descriptions_filled": desc_filled,
        "traceability_rows_filled": trace_filled,
        "total_req_ids": len(req_ids),
    }
