"""Generated DOCX quality analyzer."""

from __future__ import annotations

import json
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx.document import Document

from document_ai.learn.docx_io import iter_blocks, load_document, paragraph_deep_text, table_matrix
from document_ai.quality.patterns import (
    BROKEN_SENTENCE_PATTERNS,
    DOMAIN_FOREIGN_TERMS,
    FIGURE_TABLE_PATTERNS,
    MINDRIUM_PATTERNS,
    PLACEHOLDER_PATTERNS,
    STRUCTURAL_SKIP,
)
from document_ai.render.mddr import verify_mddr_completeness
from document_ai.render.mdsr import verify_completeness

REQ_ID_RE = re.compile(r"^Req\.\s*\d+", re.I)
TRACE_ID_RE = re.compile(r"^[A-Z]{2}-\d+")


@dataclass
class QualityFinding:
    category: str
    severity: str  # fail | warning | info
    location: str
    message: str
    auto_fixable: bool
    excerpt: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "category": self.category,
            "severity": self.severity,
            "location": self.location,
            "message": self.message,
            "auto_fixable": self.auto_fixable,
            "excerpt": self.excerpt,
        }


@dataclass
class DocumentAnalysis:
    document: str
    template: str
    findings: list[QualityFinding] = field(default_factory=list)
    structure_review: dict[str, Any] = field(default_factory=dict)
    paragraph_count: int = 0
    table_count: int = 0

    def to_dict(self) -> dict[str, Any]:
        return {
            "document": self.document,
            "template": self.template,
            "findings": [f.to_dict() for f in self.findings],
            "structure_review": self.structure_review,
            "paragraph_count": self.paragraph_count,
            "table_count": self.table_count,
        }


def _normalize_paragraph(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip().lower())


def _is_structural(text: str) -> bool:
    stripped = text.strip()
    if not stripped or len(stripped) < 20:
        return True
    if stripped in STRUCTURAL_SKIP:
        return True
    if REQ_ID_RE.match(stripped):
        return True
    if TRACE_ID_RE.match(stripped):
        return True
    return False


def _scan_text(
    text: str,
    location: str,
    *,
    domain: str,
    product_name: str,
    allow_brand: set[str],
) -> list[QualityFinding]:
    findings: list[QualityFinding] = []
    if not text or not text.strip():
        return findings

    for name, pattern in PLACEHOLDER_PATTERNS:
        if pattern.search(text):
            findings.append(
                QualityFinding(
                    category="placeholder",
                    severity="fail",
                    location=location,
                    message=f"Placeholder detected ({name})",
                    auto_fixable=True,
                    excerpt=text[:120],
                )
            )
            break

    for name, pattern in MINDRIUM_PATTERNS:
        if domain == "hospital_reservation" and name == "medical_residual":
            continue
        if pattern.search(text):
            allow = any(brand.lower() in text.lower() for brand in allow_brand if brand)
            if name == "mindrium_brand" and allow:
                continue
            findings.append(
                QualityFinding(
                    category="mindrium",
                    severity="fail",
                    location=location,
                    message=f"Residual reference text ({name})",
                    auto_fixable=True,
                    excerpt=text[:120],
                )
            )
            break

    for label, pattern in DOMAIN_FOREIGN_TERMS.get(domain, []):
        if pattern.search(text):
            if product_name and product_name in text:
                continue
            findings.append(
                QualityFinding(
                    category="domain_mismatch",
                    severity="warning",
                    location=location,
                    message=f"Domain mismatch — foreign term ({label})",
                    auto_fixable=True,
                    excerpt=text[:120],
                )
            )
            break

    for name, pattern in FIGURE_TABLE_PATTERNS:
        if pattern.search(text):
            findings.append(
                QualityFinding(
                    category="figure_table",
                    severity="warning",
                    location=location,
                    message=f"Figure/table placeholder ({name})",
                    auto_fixable=False,
                    excerpt=text[:120],
                )
            )
            break

    if len(text.strip()) >= 40 and not _is_structural(text):
        for name, pattern in BROKEN_SENTENCE_PATTERNS:
            if pattern.search(text.strip()):
                findings.append(
                    QualityFinding(
                        category="broken_sentence",
                        severity="warning",
                        location=location,
                        message=f"Possibly broken sentence ({name})",
                        auto_fixable=False,
                        excerpt=text[:120],
                    )
                )
                break

    return findings


def analyze_document(
    doc_path: Path,
    *,
    template: str,
    domain: str = "",
    product_name: str = "",
    allow_brand: set[str] | None = None,
    requirements: list[dict[str, Any]] | None = None,
    design_items: list[dict[str, Any]] | None = None,
    security_tests: list[dict[str, Any]] | None = None,
) -> DocumentAnalysis:
    doc_path = doc_path.resolve()
    doc = load_document(doc_path)
    allow_brand = allow_brand or set()
    if product_name:
        allow_brand = {*allow_brand, product_name}

    analysis = DocumentAnalysis(document=str(doc_path), template=template)
    analysis.paragraph_count = len(doc.paragraphs)
    analysis.table_count = len(doc.tables)

    paragraphs: list[str] = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph_deep_text(paragraph)
        if text:
            paragraphs.append(text)
        location = f"paragraph[{index}]"
        analysis.findings.extend(
            _scan_text(
                text,
                location,
                domain=domain,
                product_name=product_name,
                allow_brand=allow_brand,
            )
        )

    for t_index, table in enumerate(doc.tables):
        matrix = table_matrix(table)
        for r_index, row in enumerate(matrix):
            for c_index, cell in enumerate(row):
                location = f"table[{t_index}] r{r_index}c{c_index}"
                analysis.findings.extend(
                    _scan_text(
                        cell,
                        location,
                        domain=domain,
                        product_name=product_name,
                        allow_brand=allow_brand,
                    )
                )

    normalized = [_normalize_paragraph(p) for p in paragraphs if not _is_structural(p)]
    counts = Counter(normalized)
    for text, count in counts.items():
        if count < 2 or len(text) < 40:
            continue
        analysis.findings.append(
            QualityFinding(
                category="duplicate_paragraph",
                severity="warning",
                location="document",
                message=f"Duplicate paragraph appears {count} times",
                auto_fixable=False,
                excerpt=text[:120],
            )
        )

    if template == "spec_requirements":
        analysis.structure_review = verify_completeness(
            doc, requirements or [], domain=domain
        )
        for issue in analysis.structure_review.get("issues", [])[:20]:
            analysis.findings.append(
                QualityFinding(
                    category="structure",
                    severity="fail",
                    location="structure",
                    message=issue,
                    auto_fixable=True,
                )
            )
        for residual in analysis.structure_review.get("residual", [])[:10]:
            analysis.findings.append(
                QualityFinding(
                    category="mindrium",
                    severity="fail",
                    location="structure",
                    message=residual,
                    auto_fixable=True,
                    excerpt=residual[:120],
                )
            )
    elif template == "spec_design":
        analysis.structure_review = verify_mddr_completeness(
            doc, design_items or [], domain=domain
        )
        for issue in analysis.structure_review.get("issues", [])[:20]:
            analysis.findings.append(
                QualityFinding(
                    category="structure",
                    severity="fail",
                    location="structure",
                    message=issue,
                    auto_fixable=True,
                )
            )
    elif template == "report_security_verification":
        from document_ai.render.xxcs import verify_xxcs_completeness

        # Mindrium XXCS skeleton retains DC/RA section titles that trip MDSR residual scanners.
        analysis.findings = [
            f
            for f in analysis.findings
            if f.category not in {"mindrium", "domain_mismatch"}
        ]
        analysis.structure_review = verify_xxcs_completeness(doc, security_tests or [])
        for issue in analysis.structure_review.get("issues", [])[:20]:
            severity = "warning" if str(issue).startswith("missing heading") else "fail"
            analysis.findings.append(
                QualityFinding(
                    category="structure",
                    severity=severity,
                    location="xxcs",
                    message=issue,
                    auto_fixable=True,
                )
            )
        if analysis.structure_review.get("tables_with_content", 0) == 0:
            analysis.findings.append(
                QualityFinding(
                    category="structure",
                    severity="fail",
                    location="xxcs",
                    message="Missing security test result tables / empty XXCS content",
                    auto_fixable=True,
                )
            )
        for residual in analysis.structure_review.get("residual", [])[:10]:
            analysis.findings.append(
                QualityFinding(
                    category="placeholder",
                    severity="warning",
                    location="xxcs",
                    message=residual,
                    auto_fixable=True,
                    excerpt=residual[:120],
                )
            )

    return analysis


def analyze_case(case_dir: Path) -> dict[str, Any]:
    case_dir = case_dir.resolve()
    input_path = case_dir / "input.json"
    payload = json.loads(input_path.read_text(encoding="utf-8")) if input_path.exists() else {}
    domain = payload.get("domain", "")
    product_name = payload.get("product_name") or (payload.get("facts") or {}).get("product_name", "")
    facts = payload.get("facts") or {}
    product_code = facts.get("product_code") or payload.get("product_code", "")
    allow_brand = {product_name} if product_name else set()
    if product_code:
        allow_brand.update(
            {
                product_code,
                f"{product_code}-web",
                f"{product_code}-api",
                f"{product_code}-admin",
                f"{product_code}-MALL",
            }
        )
    if payload.get("case_id") in {"jm_collection", "lab_ec_sw"}:
        allow_brand.update({"JM COLLECTION", "JM-web", "JM-api", "JM-admin", "JM-MALL-001"})

    requirements = []
    req_path = case_dir / "requirements.json"
    if req_path.exists():
        requirements = json.loads(req_path.read_text(encoding="utf-8")).get("requirements", [])

    design_items = []
    design_path = case_dir / "design_items.json"
    if design_path.exists():
        design_items = json.loads(design_path.read_text(encoding="utf-8")).get("items", [])

    security_tests = []
    from document_ai.form_fill.security_execution import build_effective_security_payload

    if (case_dir / "security_tests.json").exists():
        security_tests = build_effective_security_payload(case_dir).get("tests", [])

    documents: dict[str, Any] = {}
    mdsr = case_dir / "output_mdsr.docx"
    if mdsr.exists():
        documents["mdsr"] = analyze_document(
            mdsr,
            template="spec_requirements",
            domain=domain,
            product_name=product_name,
            allow_brand=allow_brand,
            requirements=requirements,
        ).to_dict()

    mddr = case_dir / "output_mddr.docx"
    if mddr.exists():
        documents["mddr"] = analyze_document(
            mddr,
            template="spec_design",
            domain=domain,
            product_name=product_name,
            allow_brand=allow_brand,
            design_items=design_items,
        ).to_dict()

    xxcs = case_dir / "output_xxcs.docx"
    if xxcs.exists():
        documents["xxcs"] = analyze_document(
            xxcs,
            template="report_security_verification",
            domain=domain,
            product_name=product_name,
            allow_brand=allow_brand,
            security_tests=security_tests,
        ).to_dict()

    all_findings = []
    for doc_analysis in documents.values():
        all_findings.extend(doc_analysis.get("findings", []))

    return {
        "case": str(case_dir),
        "domain": domain,
        "product_name": product_name,
        "documents": documents,
        "findings": all_findings,
    }
