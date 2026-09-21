"""Embed and read hidden cell node IDs inside DOCX table cells.

IDs are stored as vanished runs: ``[[NODE:<id>]]`` so they survive round-trips
and stay invisible in normal Word display.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.table import _Cell

NODE_TAG_RE = re.compile(r"\[\[NODE:([^\]]+)\]\]")


def _vanish_run(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    r_pr = run._r.get_or_add_rPr()
    vanish = OxmlElement("w:vanish")
    r_pr.append(vanish)
    # Also mark as webHidden for some viewers.
    web_hidden = OxmlElement("w:webHidden")
    r_pr.append(web_hidden)


def set_cell_text_with_node_id(cell: _Cell, text: str, node_id: str) -> None:
    """Replace cell content with visible text + vanished node tag."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.add_run(text)
    _vanish_run(paragraph, f"[[NODE:{node_id}]]")


def cell_visible_text(cell: _Cell) -> str:
    """Return cell text with hidden node tags stripped."""
    raw = cell.text or ""
    return NODE_TAG_RE.sub("", raw).strip()


def cell_node_id(cell: _Cell) -> str | None:
    m = NODE_TAG_RE.search(cell.text or "")
    return m.group(1) if m else None


def read_node_map(docx_path: Path | str) -> dict[str, dict[str, Any]]:
    """Map node_id → {table, row, col, text} for all tagged cells."""
    doc = Document(str(docx_path))
    out: dict[str, dict[str, Any]] = {}
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                nid = cell_node_id(cell)
                if not nid:
                    continue
                out[nid] = {
                    "node_id": nid,
                    "table": ti,
                    "row": ri,
                    "col": ci,
                    "text": cell_visible_text(cell),
                }
    return out


def apply_node_patches(
    docx_path: Path | str,
    patches: list[dict[str, Any]],
    out_path: Path | str,
) -> list[str]:
    """Apply REPLACE_CELL patches keyed by node_id; write to out_path (never in-place unless same)."""
    doc = Document(str(docx_path))
    by_id = {str(p.get("node_id")): p for p in patches if p.get("node_id")}
    applied: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                nid = cell_node_id(cell)
                if not nid or nid not in by_id:
                    continue
                patch = by_id[nid]
                new_val = patch.get("after") or patch.get("new_value") or patch.get("value") or ""
                set_cell_text_with_node_id(cell, str(new_val), nid)
                applied.append(nid)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return applied
