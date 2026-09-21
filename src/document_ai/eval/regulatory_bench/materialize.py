"""Materialize synthetic MDSR/MDDR DOCX packs with hidden cell IDs."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .case_loader import load_case
from .hidden_ids import set_cell_text_with_node_id


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_mdsr_docx(
    out_path: Path,
    *,
    product_name: str,
    requirements: list[dict[str, str]],
    trace_rows: list[dict[str, str]],
) -> Path:
    doc = Document()
    _title(doc, f"MDSR — {product_name} (synthetic bench)")
    doc.add_paragraph("소프트웨어 요구사항명세서 (벤치마크용 합성본)")

    doc.add_paragraph("1. 요구사항")
    table = doc.add_table(rows=1 + len(requirements), cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "요구사항 ID"
    table.rows[0].cells[1].text = "설명"
    for i, req in enumerate(requirements, start=1):
        table.rows[i].cells[0].text = req["req_id"]
        set_cell_text_with_node_id(table.rows[i].cells[1], req["text"], req["node_id"])

    doc.add_paragraph("2. 추적성")
    tr = doc.add_table(rows=1 + len(trace_rows), cols=2)
    tr.style = "Table Grid"
    tr.rows[0].cells[0].text = "요구사항"
    tr.rows[0].cells[1].text = "설계 항목"
    for i, row in enumerate(trace_rows, start=1):
        set_cell_text_with_node_id(tr.rows[i].cells[0], row["req_label"], row["node_id"])
        tr.rows[i].cells[1].text = row["design_label"]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def build_mddr_docx(
    out_path: Path,
    *,
    product_name: str,
    design_items: list[dict[str, Any]],
) -> Path:
    doc = Document()
    _title(doc, f"MDDR — {product_name} (synthetic bench)")
    doc.add_paragraph("소프트웨어 설계명세서 (벤치마크용 합성본)")

    doc.add_paragraph("1. 설계 항목")
    table = doc.add_table(rows=1 + len(design_items), cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "설계 ID"
    table.rows[0].cells[1].text = "파라미터"
    table.rows[0].cells[2].text = "알림"
    for i, item in enumerate(design_items, start=1):
        table.rows[i].cells[0].text = item["design_id"]
        set_cell_text_with_node_id(
            table.rows[i].cells[1], item["param_text"], item["param_node_id"]
        )
        notify_id = item.get("notify_node_id")
        notify_text = item.get("notify_text") or ""
        if notify_id:
            set_cell_text_with_node_id(table.rows[i].cells[2], notify_text, notify_id)
        else:
            table.rows[i].cells[2].text = notify_text

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def _specs_from_gold(case: dict[str, Any]) -> dict[str, Any]:
    meta = case["meta"]
    gold = case["gold"]
    product = (meta.get("profile") or {}).get("product_name") or "Synthetic Product"
    patch_by_id = {
        str(p["node_id"]): p
        for p in (gold.get("expected_patch") or {}).get("patches") or []
    }

    req_defaults: dict[str, tuple[str, str]] = {
        "REQ_001_DESC": ("Req. 1", "사용자는 로그인 후 서비스를 이용할 수 있어야 한다."),
        "REQ_006_DESC": ("Req. 6", "시스템은 사용자 동의 없이 개인정보를 외부로 전송하지 않아야 한다."),
        "REQ_007_DESC": ("Req. 7", "사용자는 수면일기를 매일 입력할 수 있어야 한다."),
    }
    for nid, p in patch_by_id.items():
        if nid.startswith("REQ_") and nid.endswith("_DESC"):
            num = nid.split("_")[1]
            label = str(p.get("req_id") or f"Req. {int(num)}")
            prior = req_defaults.get(nid, (label, ""))[1]
            req_defaults[nid] = (label, str(p.get("before") or prior))

    requirements = [
        {"req_id": label, "text": text, "node_id": nid}
        for nid, (label, text) in req_defaults.items()
    ]

    design_items = [
        {
            "design_id": "DI-1",
            "param_text": "auth_method=password",
            "param_node_id": "DI_001_PARAM",
            "notify_text": "",
            "notify_node_id": None,
        },
        {
            "design_id": "DI-12",
            "param_text": str(
                patch_by_id.get("DI_012_PARAM", {}).get("before")
                or "diary_input_frequency=daily"
            ),
            "param_node_id": "DI_012_PARAM",
            "notify_text": str(
                patch_by_id.get("DI_012_NOTIFY", {}).get("before") or "매일 21:00 입력 알림"
            ),
            "notify_node_id": "DI_012_NOTIFY",
        },
    ]

    trace_rows = [
        {"req_label": "Req. 1 → DI-1", "design_label": "DI-1", "node_id": "TRACE_ROW_REQ_001"},
        {"req_label": "Req. 7 → DI-12", "design_label": "DI-12", "node_id": "TRACE_ROW_REQ_007"},
    ]
    return {
        "product_name": product,
        "requirements": requirements,
        "design_items": design_items,
        "trace_rows": trace_rows,
    }


def write_fingerprints(case_dir: Path, rel_paths: list[str]) -> Path:
    lines = [
        "# sha256 fingerprints of immutable inputs",
        "# FORMAT: <sha256>  <relative_path>",
    ]
    for rel in rel_paths:
        lines.append(f"{sha256_file(case_dir / rel)}  {rel}")
    fp_path = case_dir / "fp" / "originals.sha256"
    fp_path.parent.mkdir(parents=True, exist_ok=True)
    fp_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return fp_path


def materialize_case(case_dir: Path | str, *, force: bool = True) -> dict[str, Any]:
    """Write MDSR/MDDR DOCX + fingerprints for a case pack."""
    case = load_case(case_dir)
    root = Path(case["case_dir"])
    specs = _specs_from_gold(case)

    doc_paths = {
        d["document_id"]: root / d["path"] for d in (case["meta"].get("documents") or [])
    }
    mdsr_path = doc_paths.get("MDSR_v1") or (root / "input" / "MDSR_v1.docx")
    mddr_path = doc_paths.get("MDDR_v1") or (root / "input" / "MDDR_v1.docx")

    if force or not mdsr_path.exists():
        build_mdsr_docx(
            mdsr_path,
            product_name=specs["product_name"],
            requirements=specs["requirements"],
            trace_rows=specs["trace_rows"],
        )
    if force or not mddr_path.exists():
        build_mddr_docx(
            mddr_path,
            product_name=specs["product_name"],
            design_items=specs["design_items"],
        )

    rels = [of["path"] for of in (case["gold"].get("original_files") or []) if of.get("path")]
    if not rels:
        rels = [str(mdsr_path.relative_to(root)), str(mddr_path.relative_to(root))]
    fp_path = write_fingerprints(root, rels)

    manifest = {
        "case_id": case["case_id"],
        "mdsr": str(mdsr_path),
        "mddr": str(mddr_path),
        "fingerprints": str(fp_path),
        "sha256": {rel: sha256_file(root / rel) for rel in rels},
    }
    (root / "input" / "materialize_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return manifest
