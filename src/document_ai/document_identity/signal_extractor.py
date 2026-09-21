# -*- coding: utf-8 -*-
"""Filename / content / structure / identifier signal extraction."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from docx import Document

from document_ai.document_identity.schema import DocumentSignal
from document_ai.domain_packs.ec_sw.mdtm_schema import (
    extract_design_ids,
    extract_requirement_ids,
    extract_test_ids,
)
from document_ai.template.concept_normalization import normalize_concepts, tokenize

_SHORT_ID_PATTERNS: list[tuple[re.Pattern[str], str, str, str]] = [
    # pattern, short_id, document_type, role
    (re.compile(r"mdtm|traceability|matrix|추적", re.I), "MDTM", "MDTM", "traceability"),
    (re.compile(r"mdsr|requirement|요구", re.I), "MDSR", "MDSR", "requirements"),
    (re.compile(r"mddr|design|설계", re.I), "MDDR", "MDDR", "design"),
    (re.compile(r"mdvp|verification|검증계획", re.I), "MDVP", "MDVP", "verification_plan"),
    (re.compile(r"xxcs|security.?test|보안시험", re.I), "XXCS", "XXCS", "security_tests"),
    (re.compile(r"proposal|제안|사업계획", re.I), "PROPOSAL", "BUSINESS_PROPOSAL", "business_proposal"),
    (re.compile(r"report|보고서|성과", re.I), "REPORT", "GENERAL_REPORT", "general_report"),
]

_REPORT_HEADINGS = frozenset(
    {"배경", "목적", "방법론", "결과", "논의", "결론", "참고문헌", "methodology", "results", "conclusion", "background", "objectives"}
)
_PROPOSAL_HEADINGS = frozenset(
    {"문제", "제안", "수행", "일정", "예산", "위험", "기대", "조직", "schedule", "budget", "risk", "execution"}
)


def _sid(prefix: str, value: str) -> str:
    h = hashlib.sha1(f"{prefix}:{value}".encode("utf-8")).hexdigest()[:8]
    return f"{prefix}_{h}"


def extract_filename_signals(*, source_document_id: str, filename: str) -> list[DocumentSignal]:
    name = filename or ""
    signals: list[DocumentSignal] = []
    for rx, short_id, dtype, role in _SHORT_ID_PATTERNS:
        if rx.search(name):
            signals.append(
                DocumentSignal(
                    signal_id=_sid("fn", f"{short_id}:{name}"),
                    document_source_id=source_document_id,
                    signal_type="FILENAME_TOKEN",
                    signal_value=name,
                    normalized_value=short_id,
                    confidence=0.45,
                    source="filename",
                    reason_codes=[f"filename_hint:{short_id}", "filename_not_sufficient_alone"],
                    independent_group="filename",
                )
            )
            signals.append(
                DocumentSignal(
                    signal_id=_sid("fnrole", f"{role}:{name}"),
                    document_source_id=source_document_id,
                    signal_type="DOCUMENT_ROLE_HINT",
                    signal_value=role,
                    normalized_value=role,
                    confidence=0.4,
                    source="filename",
                    reason_codes=["filename_role_hint"],
                    independent_group="filename",
                )
            )
    signals.append(
        DocumentSignal(
            signal_id=_sid("fmt", name),
            document_source_id=source_document_id,
            signal_type="FILE_FORMAT",
            signal_value=Path(name).suffix.lower() or ".docx",
            normalized_value="docx",
            confidence=0.2,
            source="filename",
            reason_codes=["file_format"],
            independent_group="format",
        )
    )
    return signals


def extract_content_structure_signals(
    *,
    source_document_id: str,
    docx_path: Path | None,
    file_bytes: bytes | None = None,
) -> list[DocumentSignal]:
    signals: list[DocumentSignal] = []
    doc: Document | None = None
    try:
        if docx_path and Path(docx_path).is_file():
            doc = Document(str(docx_path))
        elif file_bytes:
            import io

            doc = Document(io.BytesIO(file_bytes))
    except Exception:
        return signals
    if doc is None:
        return signals

    headings: list[str] = []
    for i, p in enumerate(doc.paragraphs[:40]):
        text = (p.text or "").strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        if i == 0 or "heading" in style.lower() or len(text) <= 40:
            headings.append(text)
            signals.append(
                DocumentSignal(
                    signal_id=_sid("hd", f"{i}:{text[:40]}"),
                    document_source_id=source_document_id,
                    signal_type="HEADING_TEXT" if i > 0 else "DOCUMENT_TITLE",
                    signal_value=text[:200],
                    normalized_value=text.lower()[:80],
                    confidence=0.55 if "heading" in style.lower() else 0.5,
                    source="content",
                    source_locator={"paragraph_index": i},
                    reason_codes=["heading_or_title"],
                    independent_group="heading",
                )
            )

    heading_blob = " ".join(headings).lower()
    if any(h.lower() in heading_blob for h in _REPORT_HEADINGS):
        signals.append(
            DocumentSignal(
                signal_id=_sid("concept", "general_report"),
                document_source_id=source_document_id,
                signal_type="CONTENT_CONCEPT",
                signal_value="general_report",
                normalized_value="general_report",
                confidence=0.6,
                source="content",
                reason_codes=["report_heading_concepts"],
                independent_group="template_structure",
            )
        )
    if any(h.lower() in heading_blob for h in _PROPOSAL_HEADINGS):
        signals.append(
            DocumentSignal(
                signal_id=_sid("concept", "business_proposal"),
                document_source_id=source_document_id,
                signal_type="CONTENT_CONCEPT",
                signal_value="business_proposal",
                normalized_value="business_proposal",
                confidence=0.6,
                source="content",
                reason_codes=["proposal_heading_concepts"],
                independent_group="template_structure",
            )
        )

    # Tables
    n_tables = len(doc.tables)
    if n_tables:
        signals.append(
            DocumentSignal(
                signal_id=_sid("tbln", str(n_tables)),
                document_source_id=source_document_id,
                signal_type="TEMPLATE_STRUCTURE",
                signal_value=f"tables:{n_tables}",
                normalized_value=str(n_tables),
                confidence=0.35,
                source="structure",
                reason_codes=["table_count"],
                independent_group="structure",
            )
        )
    req_hits = des_hits = test_hits = 0
    matrixish = False
    for ti, table in enumerate(doc.tables[:8]):
        headers = []
        if table.rows:
            headers = [(c.text or "").strip() for c in table.rows[0].cells]
            header_join = " ".join(headers).lower()
            signals.append(
                DocumentSignal(
                    signal_id=_sid("th", f"{ti}:{header_join[:60]}"),
                    document_source_id=source_document_id,
                    signal_type="TABLE_HEADER",
                    signal_value=header_join[:200],
                    normalized_value=header_join[:80],
                    confidence=0.55,
                    source="structure",
                    source_locator={"table_index": ti},
                    reason_codes=["table_header"],
                    independent_group="table_header",
                )
            )
            if any(k in header_join for k in ("req", "요구", "design", "설계", "test", "시험", "trace")):
                matrixish = True
        for ri, row in enumerate(table.rows[1:40], start=1):
            cells = [(c.text or "").strip() for c in row.cells]
            blob = " | ".join(cells)
            rids = extract_requirement_ids(blob)
            dids = extract_design_ids(blob)
            tids = extract_test_ids(blob)
            req_hits += len(rids)
            des_hits += len(dids)
            test_hits += len(tids)
            if rids and (dids or tids):
                matrixish = True
                signals.append(
                    DocumentSignal(
                        signal_id=_sid("idrow", f"{ti}:{ri}:{rids[0]}"),
                        document_source_id=source_document_id,
                        signal_type="IDENTIFIER_PATTERN",
                        signal_value=",".join(rids + dids[:2] + tids[:2]),
                        normalized_value="trace_row",
                        confidence=0.85,
                        source="identifiers",
                        source_locator={"table_index": ti, "row_index": ri},
                        reason_codes=["req_with_design_or_test_same_row", "mdtm_strong_evidence"],
                        independent_group="identifier_structure",
                    )
                )

    if req_hits:
        signals.append(
            DocumentSignal(
                signal_id=_sid("reqc", str(req_hits)),
                document_source_id=source_document_id,
                signal_type="IDENTIFIER_PATTERN",
                signal_value=f"requirement_count:{req_hits}",
                normalized_value="REQUIREMENT",
                confidence=min(0.9, 0.4 + 0.05 * req_hits),
                source="identifiers",
                reason_codes=["requirement_id_density"],
                independent_group="identifier_density",
            )
        )
    if des_hits and not req_hits:
        signals.append(
            DocumentSignal(
                signal_id=_sid("desc", str(des_hits)),
                document_source_id=source_document_id,
                signal_type="IDENTIFIER_PATTERN",
                signal_value=f"design_count:{des_hits}",
                normalized_value="DESIGN",
                confidence=0.7,
                source="identifiers",
                reason_codes=["design_id_density"],
                independent_group="identifier_density",
            )
        )
    if matrixish:
        signals.append(
            DocumentSignal(
                signal_id=_sid("mtx", "mdtm"),
                document_source_id=source_document_id,
                signal_type="TEMPLATE_STRUCTURE",
                signal_value="traceability_matrix",
                normalized_value="MDTM",
                confidence=0.8,
                source="structure",
                reason_codes=["traceability_matrix_structure"],
                independent_group="structure",
            )
        )
    return signals


def extract_user_hint_signals(
    *,
    source_document_id: str,
    user_hints: dict[str, Any] | None,
) -> list[DocumentSignal]:
    if not user_hints:
        return []
    out: list[DocumentSignal] = []
    for key in ("short_id", "document_role", "domain_pack_id", "document_type"):
        val = user_hints.get(key)
        if not val:
            continue
        out.append(
            DocumentSignal(
                signal_id=_sid("uh", f"{key}:{val}"),
                document_source_id=source_document_id,
                signal_type="EXPLICIT_USER_HINT",
                signal_value=str(val),
                normalized_value=str(val).upper() if key == "short_id" else str(val),
                confidence=0.7,
                source="user",
                reason_codes=["explicit_user_hint", "user_hint_not_sufficient_alone"],
                independent_group="user_hint",
            )
        )
    return out


def extract_all_signals(
    *,
    source_document_id: str,
    filename: str,
    docx_path: Path | None = None,
    file_bytes: bytes | None = None,
    user_hints: dict[str, Any] | None = None,
) -> list[DocumentSignal]:
    signals: list[DocumentSignal] = []
    signals.extend(extract_filename_signals(source_document_id=source_document_id, filename=filename))
    signals.extend(
        extract_content_structure_signals(
            source_document_id=source_document_id,
            docx_path=docx_path,
            file_bytes=file_bytes,
        )
    )
    signals.extend(
        extract_user_hint_signals(source_document_id=source_document_id, user_hints=user_hints)
    )
    return signals
