# -*- coding: utf-8 -*-
"""PR-25: Controlled Writer Activation orchestrator."""

from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any

from document_ai.controlled_writer.approval import make_approval
from document_ai.controlled_writer.copy_workspace import file_sha256
from document_ai.controlled_writer.executor import execute_controlled_write
from document_ai.controlled_writer.schema import (
    ControlledWriterInput,
    DiffResult,
    RollbackPoint,
    WriterPlan,
    WriterResult,
)
from document_ai.controlled_writer.validation import validate_controlled_writer
from document_ai.patch_contract.fingerprint import fingerprint_text


def _write_fixture(path: Path, text: str) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return fingerprint_text(text)["fingerprint"]


def _write_docx_paragraph(path: Path, paragraphs: list[str]) -> str:
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    # clear default empty para content by using first
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        for t in paragraphs[1:]:
            doc.add_paragraph(t)
    else:
        for t in paragraphs:
            doc.add_paragraph(t)
    doc.save(str(path))
    return file_sha256(path)


def _write_docx_table(path: Path, headers: list[str], rows: list[list[str]]) -> str:
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("Table fixture")
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    doc.save(str(path))
    return file_sha256(path)


def prepare_sample_inputs(fixture_dir: Path) -> list[ControlledWriterInput]:
    """Create original fixtures + ControlledWriterInput list (never mutate later)."""
    fixture_dir.mkdir(parents=True, exist_ok=True)

    # 1. Paragraph UPDATE (markdown)
    para_text = (
        "# Report\n\n## Methodology\n\n"
        "내부 artifact 및 샘플 fixture\n\n"
        "## Results\n\n- 핵심 발견 1\n"
    )
    para_path = fixture_dir / "sample_paragraph.md"
    para_fp = _write_fixture(para_path, para_text)
    orig = "내부 artifact 및 샘플 fixture"
    start = para_text.index(orig)
    span = (start, start + len(orig))

    # 2. Table-like markdown UPDATE
    table_text = "# T\n\n| Col | Val |\n| --- | --- |\n| A | 항목A |\n"
    table_path = fixture_dir / "sample_table.md"
    table_fp = _write_fixture(table_path, table_text)
    cell_orig = "항목A"
    cell_start = table_text.index(cell_orig)
    cell_span = (cell_start, cell_start + len(cell_orig))

    # 3. List UPDATE
    list_text = "# R\n\n- 핵심 발견 1\n- 핵심 발견 2\n"
    list_path = fixture_dir / "sample_list.md"
    list_fp = _write_fixture(list_path, list_text)
    list_orig = "핵심 발견 1"
    list_start = list_text.index(list_orig)
    list_span = (list_start, list_start + len(list_orig))

    # 4. ADD
    add_text = "# Schedule\n\n1~4주 단계 수행\n"
    add_path = fixture_dir / "sample_add.md"
    add_fp = _write_fixture(add_path, add_text)
    # ADD uses SOURCE_ABSOLUTE with span on whole file for gate
    add_span = (0, len(add_text))

    # 5. DELETE
    del_text = "# Body\n\n일정 본문\n\nKeep me\n"
    del_path = fixture_dir / "sample_delete.md"
    del_fp = _write_fixture(del_path, del_text)
    del_orig = "일정 본문"
    del_start = del_text.index(del_orig)
    del_span = (del_start, del_start + len(del_orig))

    # 6. LINK
    link_text = "# Refs\n\nref\n"
    link_path = fixture_dir / "sample_link.md"
    link_fp = _write_fixture(link_path, link_text)
    link_orig = "ref"
    link_start = link_text.index(link_orig)
    link_span = (link_start, link_start + len(link_orig))

    # 7. Approval REJECT (would-be update)
    rej_text = "# X\n\nreject me\n"
    rej_path = fixture_dir / "sample_reject.md"
    rej_fp = _write_fixture(rej_path, rej_text)
    rej_orig = "reject me"
    rej_start = rej_text.index(rej_orig)
    rej_span = (rej_start, rej_start + len(rej_orig))

    # 8. Rollback forced failure fixture
    rb_text = "# RB\n\nrollback target\n"
    rb_path = fixture_dir / "sample_rollback.md"
    rb_fp = _write_fixture(rb_path, rb_text)
    rb_orig = "rollback target"
    rb_start = rb_text.index(rb_orig)
    rb_span = (rb_start, rb_start + len(rb_orig))

    # 9. Feature flag off case (same as para but separate file)
    ff_text = "# FF\n\nflag off text\n"
    ff_path = fixture_dir / "sample_flag_off.md"
    ff_fp = _write_fixture(ff_path, ff_text)
    ff_orig = "flag off text"
    ff_start = ff_text.index(ff_orig)
    ff_span = (ff_start, ff_start + len(ff_orig))

    # 10. DOCX paragraph
    docx_p = fixture_dir / "sample_paragraph.docx"
    docx_p_fp = _write_docx_paragraph(docx_p, ["Intro", "DOCX paragraph body", "Outro"])

    # 11. DOCX table cell
    docx_t = fixture_dir / "sample_table.docx"
    docx_t_fp = _write_docx_table(
        docx_t, ["H1", "H2"], [["항목A", "B"], ["C", "D"]]
    )

    # 12. Stale fingerprint
    stale_text = "# Stale\n\nstale body\n"
    stale_path = fixture_dir / "sample_stale.md"
    _write_fixture(stale_path, stale_text)
    stale_orig = "stale body"
    stale_start = stale_text.index(stale_orig)
    stale_span = (stale_start, stale_start + len(stale_orig))

    # 13. Unsupported adapter
    uns_text = "# U\n\nunsupported\n"
    uns_path = fixture_dir / "sample_unsupported.md"
    uns_fp = _write_fixture(uns_path, uns_text)
    uns_orig = "unsupported"
    uns_start = uns_text.index(uns_orig)
    uns_span = (uns_start, uns_start + len(uns_orig))

    inputs: list[ControlledWriterInput] = [
        ControlledWriterInput(
            writer_input_id="CWI-0001",
            patch_contract_id="PCT-0001",
            change_id="CHG-001",
            document_id="sample_paragraph",
            requested_operation="UPDATE",
            proposed_text="내부 artifact, 샘플 fixture 및 공개 데이터셋",
            original_text=orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(para_path.resolve()),
            location_type="PARAGRAPH",
            paragraph_index=0,
            character_span=span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=para_fp,
            approval=make_approval(
                approval_id="APR-0001",
                patch_contract_id="PCT-0001",
                decision="APPROVED",
                reason="paragraph update",
            ),
            metadata={"case": "paragraph_update"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0002",
            patch_contract_id="PCT-0002",
            change_id="CHG-002",
            document_id="sample_table",
            requested_operation="UPDATE",
            proposed_text="항목A-갱신",
            original_text=cell_orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(table_path.resolve()),
            location_type="TABLE_CELL",
            character_span=cell_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=table_fp,
            approval=make_approval(
                approval_id="APR-0002",
                patch_contract_id="PCT-0002",
                decision="AUTO_APPROVED",
                reason="table cell",
            ),
            metadata={"case": "table_update"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0003",
            patch_contract_id="PCT-0003",
            change_id="CHG-003",
            document_id="sample_list",
            requested_operation="UPDATE",
            proposed_text="핵심 발견 1 (보강)",
            original_text=list_orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(list_path.resolve()),
            location_type="LIST",
            list_index=0,
            character_span=list_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=list_fp,
            approval=make_approval(
                approval_id="APR-0003",
                patch_contract_id="PCT-0003",
                decision="APPROVED",
            ),
            metadata={"case": "list_update"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0004",
            patch_contract_id="PCT-0004",
            change_id="CHG-004",
            document_id="sample_add",
            requested_operation="ADD",
            proposed_text="검수 단계 추가",
            original_text=None,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(add_path.resolve()),
            location_type="PARAGRAPH",
            character_span=add_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=add_fp,
            approval=make_approval(
                approval_id="APR-0004",
                patch_contract_id="PCT-0004",
                decision="APPROVED",
            ),
            metadata={"case": "add"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0005",
            patch_contract_id="PCT-0005",
            change_id="CHG-005",
            document_id="sample_delete",
            requested_operation="DELETE",
            proposed_text=None,
            original_text=del_orig,
            contract_status="CONTRACT_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(del_path.resolve()),
            location_type="PARAGRAPH",
            character_span=del_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=del_fp,
            approval=make_approval(
                approval_id="APR-0005",
                patch_contract_id="PCT-0005",
                decision="APPROVED",
                reason="explicit delete approval",
            ),
            metadata={"case": "delete"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0006",
            patch_contract_id="PCT-0006",
            change_id="CHG-006",
            document_id="sample_link",
            requested_operation="LINK",
            proposed_text="공식 문서",
            original_text=link_orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(link_path.resolve()),
            location_type="PARAGRAPH",
            character_span=link_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=link_fp,
            link_metadata={"url": "https://example.com/doc"},
            approval=make_approval(
                approval_id="APR-0006",
                patch_contract_id="PCT-0006",
                decision="APPROVED",
            ),
            metadata={"case": "link"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0007",
            patch_contract_id="PCT-0007",
            change_id="CHG-007",
            document_id="sample_reject",
            requested_operation="UPDATE",
            proposed_text="should not apply",
            original_text=rej_orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(rej_path.resolve()),
            location_type="PARAGRAPH",
            character_span=rej_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=rej_fp,
            approval=make_approval(
                approval_id="APR-0007",
                patch_contract_id="PCT-0007",
                decision="REJECTED",
                reason="rejected by reviewer",
            ),
            metadata={"case": "approval_reject"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0008",
            patch_contract_id="PCT-0008",
            change_id="CHG-008",
            document_id="sample_rollback",
            requested_operation="UPDATE",
            proposed_text="should rollback",
            original_text=rb_orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(rb_path.resolve()),
            location_type="PARAGRAPH",
            character_span=rb_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=rb_fp,
            approval=make_approval(
                approval_id="APR-0008",
                patch_contract_id="PCT-0008",
                decision="APPROVED",
            ),
            metadata={"case": "rollback", "force_fail": True},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0009",
            patch_contract_id="PCT-0009",
            change_id="CHG-009",
            document_id="sample_flag_off",
            requested_operation="UPDATE",
            proposed_text="blocked by flag",
            original_text=ff_orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(ff_path.resolve()),
            location_type="PARAGRAPH",
            character_span=ff_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=ff_fp,
            approval=make_approval(
                approval_id="APR-0009",
                patch_contract_id="PCT-0009",
                decision="APPROVED",
            ),
            metadata={"case": "feature_flag_off", "require_flags_off": True},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0010",
            patch_contract_id="PCT-0010",
            change_id="CHG-010",
            document_id="sample_docx_paragraph",
            requested_operation="UPDATE",
            proposed_text="DOCX paragraph body UPDATED",
            original_text="DOCX paragraph body",
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="DOCX_PARAGRAPH_WRITER",
            source_path=str(docx_p.resolve()),
            location_type="PARAGRAPH",
            paragraph_index=1,
            character_span=(0, len("DOCX paragraph body")),
            span_kind="SOURCE_ABSOLUTE",
            source_format="docx",
            expected_fingerprint=docx_p_fp,
            approval=make_approval(
                approval_id="APR-0010",
                patch_contract_id="PCT-0010",
                decision="APPROVED",
            ),
            metadata={"case": "docx_paragraph_update"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0011",
            patch_contract_id="PCT-0011",
            change_id="CHG-011",
            document_id="sample_docx_table",
            requested_operation="UPDATE",
            proposed_text="항목A-갱신",
            original_text="항목A",
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="DOCX_TABLE_CELL_WRITER",
            source_path=str(docx_t.resolve()),
            location_type="TABLE_CELL",
            table_index=0,
            cell_coordinate=(1, 0),
            character_span=(0, len("항목A")),
            span_kind="SOURCE_ABSOLUTE",
            source_format="docx",
            expected_fingerprint=docx_t_fp,
            approval=make_approval(
                approval_id="APR-0011",
                patch_contract_id="PCT-0011",
                decision="APPROVED",
            ),
            metadata={"case": "docx_table_update"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0012",
            patch_contract_id="PCT-0012",
            change_id="CHG-012",
            document_id="sample_stale",
            requested_operation="UPDATE",
            proposed_text="new",
            original_text=stale_orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="MARKDOWN_BLOCK_WRITER",
            source_path=str(stale_path.resolve()),
            location_type="PARAGRAPH",
            character_span=stale_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint="0" * 64,  # stale / wrong
            approval=make_approval(
                approval_id="APR-0012",
                patch_contract_id="PCT-0012",
                decision="APPROVED",
            ),
            metadata={"case": "stale_fingerprint"},
        ),
        ControlledWriterInput(
            writer_input_id="CWI-0013",
            patch_contract_id="PCT-0013",
            change_id="CHG-013",
            document_id="sample_unsupported",
            requested_operation="UPDATE",
            proposed_text="x",
            original_text=uns_orig,
            contract_status="CONTRACT_READY_FOR_REVIEW",
            writer_adapter="UNSUPPORTED_WRITER",
            source_path=str(uns_path.resolve()),
            location_type="PARAGRAPH",
            character_span=uns_span,
            span_kind="SOURCE_ABSOLUTE",
            source_format="markdown",
            expected_fingerprint=uns_fp,
            approval=make_approval(
                approval_id="APR-0013",
                patch_contract_id="PCT-0013",
                decision="APPROVED",
            ),
            metadata={"case": "unsupported_writer"},
        ),
    ]
    return inputs


def build_summary(
    inputs: list[ControlledWriterInput],
    plans: list[WriterPlan],
    results: list[WriterResult],
    diffs: list[DiffResult],
    rollbacks: list[RollbackPoint],
    *,
    env: dict[str, str] | None = None,
) -> dict[str, Any]:
    status_counts = Counter(r.result_status for r in results)
    op_counts = Counter(r.operation for r in results)
    return {
        "stage": "controlled_writer_summary",
        "input_count": len(inputs),
        "plan_count": len(plans),
        "result_count": len(results),
        "diff_count": len(diffs),
        "rollback_count": len(rollbacks),
        "result_status_counts": dict(sorted(status_counts.items())),
        "operation_counts": dict(sorted(op_counts.items())),
        "patch_success_count": status_counts.get("APPLIED", 0),
        "failure_count": status_counts.get("FAILED", 0)
        + status_counts.get("ROLLED_BACK", 0),
        "rejected_count": status_counts.get("REJECTED", 0),
        "blocked_count": status_counts.get("BLOCKED", 0),
        "original_changed_count": sum(
            1 for r in results if r.actual_original_changed or not r.original_unchanged
        ),
        "copy_modified_count": sum(1 for r in results if r.copy_modified),
        "actual_writer_called_count": sum(
            1 for r in results if r.actual_writer_called
        ),
        "activation_allowed_plan_count": sum(
            1 for p in plans if p.activation_allowed
        ),
        "env_docx_activation": (env or {}).get("DOCX_ACTIVATION_ENABLED", "false"),
        "env_controlled_writer": (env or {}).get(
            "CONTROLLED_WRITER_ENABLED", "false"
        ),
        "note": (
            "PR-25 Controlled Writer. Originals never modified; "
            "copies only; approval + dual flags required."
        ),
    }


def run_controlled_writer_engine(
    *,
    inputs: list[ControlledWriterInput] | None = None,
    work_dir: str | Path | None = None,
    env: dict[str, str] | None = None,
    fixture_dir: str | Path | None = None,
) -> dict[str, Any]:
    """Controlled writer activation engine.

    Default env enables both flags for sample success cases.
    Pass env with flags off to exercise blocking.
    """
    base = Path(work_dir) if work_dir else Path("output") / "writer" / "_engine"
    base.mkdir(parents=True, exist_ok=True)
    fx = Path(fixture_dir) if fixture_dir else base / "fixtures"

    default_env = {
        "DOCX_ACTIVATION_ENABLED": "true",
        "CONTROLLED_WRITER_ENABLED": "true",
    }
    if env is None:
        env = dict(default_env)
    else:
        # merge: explicit env wins; missing keys keep defaults for sample runs
        merged = dict(default_env)
        merged.update(env)
        env = merged

    if inputs is None:
        inputs = prepare_sample_inputs(fx)

    plans: list[WriterPlan] = []
    results: list[WriterResult] = []
    diffs: list[DiffResult] = []
    rollbacks: list[RollbackPoint] = []
    approvals: list[dict[str, Any]] = []
    source_fps: dict[str, str] = {}

    for seq, inp in enumerate(
        sorted(inputs, key=lambda x: x.writer_input_id), start=1
    ):
        case_env = dict(env)
        if (inp.metadata or {}).get("require_flags_off"):
            case_env["DOCX_ACTIVATION_ENABLED"] = "false"
            case_env["CONTROLLED_WRITER_ENABLED"] = "false"

        force_fail = bool((inp.metadata or {}).get("force_fail"))
        src = Path(inp.source_path)
        if src.exists():
            source_fps[str(src.resolve())] = file_sha256(src)
            # also key by given path
            source_fps[inp.source_path] = source_fps[str(src.resolve())]

        plan, result, diff, rb, meta = execute_controlled_write(
            inp,
            seq=seq,
            work_dir=base / "work",
            env=case_env,
            force_fail=force_fail,
        )
        plans.append(plan)
        results.append(result)
        if diff is not None:
            diffs.append(diff)
        if rb is not None:
            rollbacks.append(rb)
        if inp.approval is not None:
            approvals.append(inp.approval.to_dict())
        elif meta.get("approval"):
            approvals.append(meta["approval"].get("approval") or meta["approval"])

    summary = build_summary(inputs, plans, results, diffs, rollbacks, env=env)
    validation = validate_controlled_writer(
        plans=plans,
        results=results,
        diffs=diffs,
        rollbacks=rollbacks,
        summary=summary,
        source_fingerprints=source_fps,
    )

    return {
        "stage": "controlled_writer_engine",
        "schema_version": "controlled_writer_v1",
        "approvals": approvals,
        "inputs": [i.to_dict() for i in sorted(inputs, key=lambda x: x.writer_input_id)],
        "writer_plans": [p.to_dict() for p in plans],
        "writer_results": [r.to_dict() for r in results],
        "diffs": [d.to_dict() for d in diffs],
        "rollbacks": [r.to_dict() for r in rollbacks],
        "summary": summary,
        "validation": validation,
        "work_dir": str(base.resolve()),
        "note": (
            "PR-25 Controlled Writer Activation. "
            "Original files are never modified; only copies under work/copies."
        ),
    }
