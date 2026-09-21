# -*- coding: utf-8 -*-
"""DOCX format preservation metrics (copy writer target; N/A when no write)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document


def _doc_stats(path: Path) -> dict[str, Any] | None:
    if not path.is_file():
        return None
    try:
        d = Document(str(path))
    except Exception:
        return None
    headings = []
    styles = []
    bold = italic = 0
    for p in d.paragraphs:
        style = (p.style.name if p.style is not None else "") or ""
        styles.append(style)
        if "heading" in style.lower():
            headings.append((p.text or "").strip())
        for run in p.runs:
            if run.bold:
                bold += 1
            if run.italic:
                italic += 1
    tables = d.tables
    merged = 0
    rows = cols = 0
    for t in tables:
        rows += len(t.rows)
        cols += len(t.columns) if t.rows else 0
        # python-docx doesn't expose merge count cheaply — approximate via unique cell ids
        try:
            seen = set()
            for row in t.rows:
                for cell in row.cells:
                    seen.add(id(cell._tc))  # noqa: SLF001
            merged += max(0, sum(len(r.cells) for r in t.rows) - len(seen))
        except Exception:
            pass
    return {
        "paragraph_count": len(d.paragraphs),
        "table_count": len(tables),
        "section_count": len(d.sections),
        "heading_texts": headings,
        "style_names": styles,
        "bold_run_count": bold,
        "italic_run_count": italic,
        "table_row_count": rows,
        "table_column_count": cols,
        "merged_cell_estimate": merged,
        "image_count": len(d.inline_shapes),
    }


def compare_format(before: Path, after: Path) -> dict[str, Any]:
    a = _doc_stats(before)
    b = _doc_stats(after)
    if a is None or b is None:
        return {"status": "N/A", "reason": "missing_or_unreadable"}
    checks = {
        "paragraph_count": a["paragraph_count"] == b["paragraph_count"],
        "table_count": a["table_count"] == b["table_count"],
        "section_count": a["section_count"] == b["section_count"],
        "heading_texts": a["heading_texts"] == b["heading_texts"],
        "style_names": a["style_names"] == b["style_names"],
        "bold_run_count": a["bold_run_count"] == b["bold_run_count"],
        "italic_run_count": a["italic_run_count"] == b["italic_run_count"],
        "table_row_count": a["table_row_count"] == b["table_row_count"],
        "table_column_count": a["table_column_count"] == b["table_column_count"],
    }
    structure_keys = ["paragraph_count", "table_count", "section_count", "heading_texts"]
    style_keys = ["style_names", "bold_run_count", "italic_run_count"]
    table_keys = ["table_count", "table_row_count", "table_column_count"]
    def _rate(keys: list[str]) -> float:
        return sum(float(checks[k]) for k in keys) / len(keys)

    non_target = sum(1 for k, v in checks.items() if not v)
    return {
        "status": "OK",
        "structure_preservation_rate": _rate(structure_keys),
        "style_preservation_rate": _rate(style_keys),
        "table_preservation_rate": _rate(table_keys),
        "non_target_change_count": non_target,
        "checks": checks,
        "before": {k: a[k] for k in a if k != "style_names"},
        "after": {k: b[k] for k in b if k != "style_names"},
    }


def aggregate_format_metrics(rows: list[dict[str, Any]]) -> dict[str, Any]:
    usable = [r for r in rows if r.get("status") == "OK"]
    if not usable:
        return {
            "structure_preservation_rate": None,
            "style_preservation_rate": None,
            "table_preservation_rate": None,
            "non_target_change_count": 0,
            "n_applicable": 0,
            "n_na": len(rows),
        }
    n = len(usable)
    return {
        "structure_preservation_rate": sum(r["structure_preservation_rate"] for r in usable) / n,
        "style_preservation_rate": sum(r["style_preservation_rate"] for r in usable) / n,
        "table_preservation_rate": sum(r["table_preservation_rate"] for r in usable) / n,
        "non_target_change_count": sum(r["non_target_change_count"] for r in usable),
        "n_applicable": n,
        "n_na": len(rows) - n,
    }
