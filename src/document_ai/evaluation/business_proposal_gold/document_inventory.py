# -*- coding: utf-8 -*-
"""Structure-only DOCX inventory for Business Proposal gold labeling.

Scans headings / paragraphs / tables with stable indices
(``heading_{i:04d}``, ``paragraph_{i:04d}``, ``table_{ti:02d}``) matching the
runtime node-id convention used by
``document_ai.document_set.proposal_table_retrieval`` / ``workflow.analysis``.

This module reads ONLY the DOCX file (python-docx) and static concept
dictionaries. It never reads prediction artifacts (ranking results,
alignments, structural match matrices, node_ranking outputs, etc.).
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from document_ai.domain_packs.business_proposal.concepts import normalize_proposal_concepts
from document_ai.domain_packs.business_proposal.structural_roles import (
    BUDGET_TABLE,
    KPI_TABLE,
    ORGANIZATION_SECTION,
    SCHEDULE_TABLE,
    classify_table_header_role,
)
from document_ai.template.concept_normalization import normalize_concepts

_BOILERPLATE_RE = re.compile(r"^(?P<heading>.+)에 대한 설명입니다\.$")
_DIGIT_RE = re.compile(r"\d")

_TABLE_ROLE_TO_CONCEPT = {
    SCHEDULE_TABLE: "SCHEDULE",
    BUDGET_TABLE: "BUDGET",
    ORGANIZATION_SECTION: "ORGANIZATION",
    KPI_TABLE: "KPI",
}


@dataclass
class InventoryNode:
    node_id: str
    node_type: str  # HEADING | PARAGRAPH | TABLE
    text: str
    concepts: list[str] = field(default_factory=list)
    level: int | None = None
    table_role: str | None = None
    row_count: int | None = None
    column_count: int | None = None
    is_boilerplate: bool = False
    has_digit: bool = False

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class InventorySection:
    heading_node_id: str | None
    heading_text: str
    concepts: list[str]
    member_node_ids: list[str]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _heading_level(style_name: str) -> int | None:
    name = (style_name or "").strip().lower()
    if name == "title":
        return 0
    m = re.match(r"heading\s*(\d+)", name)
    if m:
        return int(m.group(1))
    return None


def _is_heading_style(style_name: str) -> bool:
    name = (style_name or "").strip().lower()
    return name == "title" or name.startswith("heading")


def _table_headers(table) -> list[str]:
    if not table.rows:
        return []
    return [c.text.strip() for c in table.rows[0].cells]


def _table_blob(table) -> str:
    parts = []
    for row in table.rows:
        for cell in row.cells:
            t = (cell.text or "").strip()
            if t:
                parts.append(t)
    return " ".join(parts)


def _iter_block_items(document: Document):
    """Yield ('paragraph', Paragraph) / ('table', Table) in true document order."""
    body = document.element.body
    para_iter = iter(document.paragraphs)
    table_iter = iter(document.tables)
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", next(para_iter)
        elif child.tag == qn("w:tbl"):
            yield "table", next(table_iter)


def _node_concepts(text: str) -> set[str]:
    return normalize_proposal_concepts(text) | normalize_concepts(text)


def _table_concepts(headers: list[str], blob: str, table_role: str | None) -> set[str]:
    header_join = " ".join(headers)
    concepts = normalize_proposal_concepts(header_join + "\n" + blob[:800]) | normalize_concepts(
        header_join + "\n" + blob[:800]
    )
    mapped = _TABLE_ROLE_TO_CONCEPT.get(table_role or "")
    if mapped:
        concepts.add(mapped)
    return concepts


def inventory_document(docx_path: str | Path, *, document_id: str) -> dict[str, Any]:
    """Build a structure-only inventory of a Business Proposal DOCX.

    Returns a dict with ``document_id``, ``nodes`` (flat, doc order),
    ``sections`` (heading-bounded groups incl. tables), all derived purely
    from the document's own text/structure.
    """
    path = Path(docx_path)
    document = Document(str(path))

    nodes: list[InventoryNode] = []
    sections: list[InventorySection] = []
    current_section_members: list[str] = []
    current_heading_text = ""
    current_heading_node_id: str | None = None
    current_section_concepts: set[str] = set()

    para_idx = 0
    table_idx = 0

    def _flush_section() -> None:
        if current_section_members:
            sections.append(
                InventorySection(
                    heading_node_id=current_heading_node_id,
                    heading_text=current_heading_text,
                    concepts=sorted(current_section_concepts),
                    member_node_ids=list(current_section_members),
                )
            )

    for kind, item in _iter_block_items(document):
        if kind == "paragraph":
            text = (item.text or "").strip()
            i = para_idx
            para_idx += 1
            if not text:
                continue
            style_name = (item.style.name if item.style is not None else "") or ""
            is_heading = _is_heading_style(style_name)
            node_id = f"heading_{i:04d}" if is_heading else f"paragraph_{i:04d}"
            own_concepts = _node_concepts(text)
            m = _BOILERPLATE_RE.match(text)
            is_boilerplate = bool(m and m.group("heading") == current_heading_text)
            node = InventoryNode(
                node_id=node_id,
                node_type="HEADING" if is_heading else "PARAGRAPH",
                text=text,
                concepts=sorted(own_concepts),
                level=_heading_level(style_name) if is_heading else None,
                is_boilerplate=is_boilerplate,
                has_digit=bool(_DIGIT_RE.search(text)),
            )
            nodes.append(node)
            if is_heading:
                _flush_section()
                current_section_members = [node_id]
                current_heading_text = text
                current_heading_node_id = node_id
                current_section_concepts = set(own_concepts)
            else:
                current_section_members.append(node_id)
                current_section_concepts |= own_concepts
        else:  # table
            ti = table_idx
            table_idx += 1
            headers = _table_headers(item)
            blob = _table_blob(item)
            role = classify_table_header_role(" ".join(headers), blob=blob[:400])
            concepts = _table_concepts(headers, blob, role)
            node_id = f"table_{ti:02d}"
            node = InventoryNode(
                node_id=node_id,
                node_type="TABLE",
                text=" ".join(headers) or blob[:120],
                concepts=sorted(concepts),
                table_role=role,
                row_count=len(item.rows),
                column_count=len(item.columns) if item.rows else 0,
            )
            nodes.append(node)
            current_section_members.append(node_id)
            current_section_concepts |= concepts

    _flush_section()

    return {
        "document_id": document_id,
        "source_path": str(path).replace("\\", "/"),
        "nodes": [n.to_dict() for n in nodes],
        "sections": [s.to_dict() for s in sections],
    }


def find_matching_sections(inventory: dict[str, Any], target_concept: str) -> list[dict[str, Any]]:
    if not target_concept:
        return []
    return [s for s in inventory.get("sections") or [] if target_concept in (s.get("concepts") or [])]


def nodes_by_id(inventory: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {n["node_id"]: n for n in inventory.get("nodes") or []}
