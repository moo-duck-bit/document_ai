# -*- coding: utf-8 -*-
"""Minimal DOCX patch-in-place for change_update (preserve original document)."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from docx.document import Document
from docx.table import Table

from document_ai.learn.docx_io import load_document
from document_ai.learn.req_ids import normalize_requirement_id as normalize_req_id
from document_ai.render.requirements import _req_id_from_table

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_DRAWING = f"{{{W_NS}}}drawing"
W_PICT = f"{{{W_NS}}}pict"
W_SECTPR = f"{{{W_NS}}}sectPr"
W_TBL = f"{{{W_NS}}}tbl"
W_P = f"{{{W_NS}}}p"
W_T = f"{{{W_NS}}}t"

# Tokens that indicate JM/lab semantic contamination when newly introduced vs reference.
CONTAMINATION_TOKENS = [
    "JM COLLECTION",
    "JM-web",
    "JM-api",
    "JM-admin",
    "쇼핑몰",
    "상품",
    "주문",
    "결제",
    " PG",
    "PCI DSS",
    "PostgreSQL",
    "React",
    "Node.js",
]


@dataclass
class TextUnit:
    kind: str
    location: str
    text: str


@dataclass
class PreservePatchResult:
    ok: bool
    mdsr_path: str | None = None
    mddr_path: str | None = None
    mdsr_patched_req_ids: list[str] = field(default_factory=list)
    mddr_strategy: str = "no_automatic_design_patch"
    mddr_sha256_matches_reference: bool | None = None
    validation: dict[str, Any] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def patch_mdsr_description_only(
    doc: Document,
    req_id: str,
    description: str,
    *,
    overwrite: bool = True,
) -> bool:
    """Patch only the description cell of one Req table. No purpose/criteria/title defaults."""
    target = normalize_req_id(req_id) or req_id
    description = (description or "").strip()
    if not target or not description:
        return False

    for table in doc.tables:
        found = _req_id_from_table(table)
        if not found or (normalize_req_id(found) or found) != target:
            continue
        return _write_description_cell(table, description, overwrite=overwrite)
    return False


def _write_description_cell(table: Table, description: str, *, overwrite: bool) -> bool:
    if len(table.rows) < 2:
        return False

    # Prefer explicit 설명 row value cell (do not invent labels).
    for row in table.rows[1:]:
        if len(row.cells) < 2:
            continue
        if row.cells[0].text.strip() == "설명":
            if overwrite or not row.cells[1].text.strip():
                row.cells[1].text = description
                return True
            return False

    # Mindrium shells often have blank labels: write ONLY the first data value cell.
    # Do not insert "설명" into the label cell — that would be an extra textual change.
    row = table.rows[1]
    if len(row.cells) < 2:
        return False
    value_cell = row.cells[1]
    if overwrite or not value_cell.text.strip():
        value_cell.text = description
        return True
    return False


def save_document(doc: Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def extract_text_units(docx_path: Path) -> list[TextUnit]:
    doc = load_document(docx_path)
    units: list[TextUnit] = []
    for pi, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if text:
            units.append(TextUnit("paragraph", f"p:{pi}", text))
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                text = (cell.text or "").strip()
                if text:
                    units.append(TextUnit("cell", f"t:{ti}/r:{ri}/c:{ci}", text))
    return units


def text_diff(before: list[TextUnit], after: list[TextUnit]) -> list[dict[str, str]]:
    before_map = {u.location: u.text for u in before}
    after_map = {u.location: u.text for u in after}
    diffs: list[dict[str, str]] = []
    for loc in sorted(set(before_map) | set(after_map)):
        b = before_map.get(loc, "")
        a = after_map.get(loc, "")
        if b != a:
            diffs.append({"location": loc, "before": b, "after": a})
    return diffs


def ooxml_counts(docx_path: Path) -> dict[str, int]:
    with zipfile.ZipFile(docx_path) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    drawings = len(list(root.iter(W_DRAWING))) + len(list(root.iter(W_PICT)))
    sections = len(list(root.iter(W_SECTPR)))
    tables = len(list(root.iter(W_TBL)))
    paragraphs = len(list(root.iter(W_P)))
    return {
        "drawing_like": drawings,
        "sectPr": sections,
        "tbl": tables,
        "p": paragraphs,
    }


def style_count(docx_path: Path) -> int:
    doc = load_document(docx_path)
    try:
        return len(doc.styles)
    except Exception:  # noqa: BLE001
        return -1


def contamination_new_tokens(reference: Path, candidate: Path) -> list[dict[str, Any]]:
    def blob(path: Path) -> str:
        with zipfile.ZipFile(path) as zf:
            root = ET.fromstring(zf.read("word/document.xml"))
        return "".join((t.text or "") for t in root.iter(W_T))

    ref_text = blob(reference)
    cand_text = blob(candidate)
    hits: list[dict[str, Any]] = []
    for token in CONTAMINATION_TOKENS:
        ref_n = ref_text.count(token)
        cand_n = cand_text.count(token)
        if cand_n > ref_n:
            hits.append({"token": token, "reference_count": ref_n, "candidate_count": cand_n})
    return hits


def build_patch_preserving_outputs(
    *,
    reference_mdsr: Path,
    reference_mddr: Path,
    req_id: str,
    description: str,
    out_dir: Path,
    mddr_strategy: str = "no_automatic_design_patch",
) -> PreservePatchResult:
    """Byte-copy reference docs, patch MDSR description only, leave MDDR unchanged by default."""
    errors: list[str] = []
    warnings: list[str] = []
    out_dir.mkdir(parents=True, exist_ok=True)

    out_mdsr = out_dir / "output_mdsr.docx"
    out_mddr = out_dir / "output_mddr.docx"

    if not reference_mdsr.exists():
        return PreservePatchResult(ok=False, errors=[f"missing reference MDSR: {reference_mdsr}"])
    if not reference_mddr.exists():
        return PreservePatchResult(ok=False, errors=[f"missing reference MDDR: {reference_mddr}"])

    shutil.copy2(reference_mdsr, out_mdsr)
    shutil.copy2(reference_mddr, out_mddr)

    # Snapshot reference text units from the copied file before patch (identical to reference).
    before_units = extract_text_units(out_mdsr)
    before_counts = ooxml_counts(out_mdsr)
    before_styles = style_count(out_mdsr)
    before_size = out_mdsr.stat().st_size

    doc = load_document(out_mdsr)
    patched = patch_mdsr_description_only(doc, req_id, description, overwrite=True)
    if not patched:
        errors.append(f"failed to patch description for {req_id}")
        return PreservePatchResult(ok=False, errors=errors)
    save_document(doc, out_mdsr)

    after_units = extract_text_units(out_mdsr)
    diffs = text_diff(before_units, after_units)
    after_counts = ooxml_counts(out_mdsr)
    after_styles = style_count(out_mdsr)

    desc = description.strip()

    def _allowed(diff: dict[str, str]) -> bool:
        before = (diff.get("before") or "").strip()
        after = (diff.get("after") or "").strip()
        # Only the description value cell may change (exact CR text).
        if after == desc and before != desc:
            return True
        return False

    unexpected = [d for d in diffs if not _allowed(d)]

    mddr_hash_ref = sha256_file(reference_mddr)
    mddr_hash_out = sha256_file(out_mddr)
    mddr_match = mddr_hash_ref == mddr_hash_out

    if mddr_strategy == "no_automatic_design_patch" and not mddr_match:
        errors.append("MDDR hash mismatch under no-patch strategy")

    contam_mdsr = contamination_new_tokens(reference_mdsr, out_mdsr)
    contam_mddr = contamination_new_tokens(reference_mddr, out_mddr)
    if contam_mdsr or contam_mddr:
        errors.append("new contamination tokens vs reference")

    structure_ok = (
        before_counts == after_counts
        and before_styles == after_styles
        and abs(out_mdsr.stat().st_size - before_size) / max(before_size, 1) < 0.05
    )
    if not structure_ok:
        warnings.append("structure metrics changed beyond expected small edit")

    # Size should not collapse like skeleton regeneration
    if out_mdsr.stat().st_size < reference_mdsr.stat().st_size * 0.5:
        errors.append("MDSR size collapsed vs reference (>50% reduction)")
    if out_mddr.stat().st_size < reference_mddr.stat().st_size * 0.5:
        errors.append("MDDR size collapsed vs reference (>50% reduction)")

    # XXCS must not be written into out_dir
    if (out_dir / "output_xxcs.docx").exists():
        errors.append("XXCS present in patch-preserving output dir (out of scope)")

    validation = {
        "mdsr_text_diffs": diffs,
        "mdsr_unexpected_diffs": unexpected,
        "mdsr_unexpected_diff_count": len(unexpected),
        "mdsr_ooxml_before": before_counts,
        "mdsr_ooxml_after": after_counts,
        "mdsr_styles_before": before_styles,
        "mdsr_styles_after": after_styles,
        "mdsr_size_before_copy": before_size,
        "mdsr_size_after_patch": out_mdsr.stat().st_size,
        "mdsr_reference_size": reference_mdsr.stat().st_size,
        "mddr_reference_size": reference_mddr.stat().st_size,
        "mddr_output_size": out_mddr.stat().st_size,
        "mdsr_sha256": sha256_file(out_mdsr),
        "mddr_sha256": mddr_hash_out,
        "mddr_reference_sha256": mddr_hash_ref,
        "contamination_mdsr": contam_mdsr,
        "contamination_mddr": contam_mddr,
        "xxcs_in_output_dir": False,
    }

    ok = not errors and len(unexpected) == 0 and mddr_match and not contam_mdsr and not contam_mddr
    return PreservePatchResult(
        ok=ok,
        mdsr_path=str(out_mdsr),
        mddr_path=str(out_mddr),
        mdsr_patched_req_ids=[normalize_req_id(req_id) or req_id],
        mddr_strategy=mddr_strategy,
        mddr_sha256_matches_reference=mddr_match,
        validation=validation,
        errors=errors,
        warnings=warnings,
    )


def find_req6_design_candidates(mddr_path: Path) -> list[dict[str, str]]:
    """Report-only: locate Mentions of Req. 6 in MDDR without patching."""
    units = extract_text_units(mddr_path)
    hits: list[dict[str, str]] = []
    for u in units:
        if re.search(r"Req\.?\s*6\b", u.text, re.I):
            hits.append({"location": u.location, "text": u.text[:240]})
    return hits
